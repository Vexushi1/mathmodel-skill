"""Captured-source adapter for the existing analysis delivery/receipt/runtime gates.

The pure comparison checker never qualifies its own candidate or writes project state.
This adapter observes bytes; existing coordinators retain numerical acceptance authority.
"""
from __future__ import annotations

import hashlib
import importlib.util
import io
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Mapping

import openpyxl
import safe_yaml
import analysis_comparison as COMPARISON
import conformance_gate as CONFORMANCE
import model_code_conformance as BOUNDED
import run_config_parser
import stage_code
import state_transitions
from claim_workbook import Workbook
from execution_protocol import COMPARISON_FIELDS, comparison_config_issues, comparison_receipt_issues, declared_input_paths
from stage_inputs import input_files

ROOT = Path(__file__).resolve().parent.parent
SHEETS = {"多模型检验", "同模型多算法检验"}
AUTHORITIES = (
    "core/user_execution_contract.yaml", "core/model_approval_contract.yaml",
    "core/project_state.schema.yaml", "core/workbook_schema.yaml",
    "core/claim_evidence_contract.yaml", "modules/03_result_analysis.md",
    "scripts/analysis_comparison.py", "scripts/analysis_comparison_gate.py",
    "scripts/claim_workbook.py", "scripts/claim_values.py", "scripts/execution_protocol.py",
    "scripts/validate_model_approval.py", "scripts/semantic_identity.py", "scripts/run_config_parser.py",
    "scripts/stage_code.py", "scripts/stage_inputs.py", safe_yaml.SOURCE_RELATIVE_PATH,
    "scripts/state_transitions.py", "core/state_transition_contract.yaml",
    "templates/code/hsk_pipeline/workbook_validation.py",
)
AUTHORIZATION_FIELDS = (
    "semantic_identity_schema_version", "semantic_revision", "validated_semantic_revision", "approved_semantic_revision",
    "semantic_identity_hash", "validated_semantic_identity_hash", "approved_semantic_identity_hash", "semantic_text_hash",
    "semantic_hash", "approved_semantic_hash", "model_challenge_status", "human_model_approval_status",
    "problem_contract_status", "semantic_closure_status", "complexity_sanity_status",
    "analysis_comparison", "analysis_methods", "result_analysis_requirement_reason", "analysis_evidence_dispositions",
    "implementation_conformance_policy", "implementation_conformance", "depends_on",
)


def authorization_context(state: Mapping[str, Any], question: str) -> dict:
    """Freeze permissions and declarations while allowing receipt-batch execution updates."""
    if not isinstance(state, Mapping) or not isinstance(state.get("subproblems"), Mapping):
        raise ValueError("comparison authorization State/subproblems must be mappings")
    entry = state["subproblems"].get(question)
    if not isinstance(entry, Mapping):
        raise ValueError("comparison authorization question must be a mapping")
    result = {"question": {field: entry[field] for field in AUTHORIZATION_FIELDS if field in entry}}
    for field in ("execution", "review_receipt_policy", "review_receipts"):
        if field in state:
            result[field] = state[field]
    framework = state.get("paper_framework") or {}
    if not isinstance(framework, Mapping):
        raise ValueError("comparison authorization paper_framework must be a mapping")
    result["paper_framework"] = {field: framework[field] for field in ("claim_consumption_policy", "claim_evidence")
                                  if field in framework}
    if "paper_fragments" in framework:
        fragments = framework["paper_fragments"]
        if not isinstance(fragments, list) or any(not isinstance(row, Mapping) for row in fragments):
            raise ValueError("comparison authorization paper_fragments must be mappings")
        result["paper_framework"]["paper_fragments"] = [
            {field: value for field, value in row.items() if field != "status"} for row in fragments]
    # current_phase, framework sync/hash, execution statuses and artifact/validated
    # hash registries are coordinator outputs; a preceding receipt may update them.
    return result


def assert_authorization_context(state: Mapping[str, Any], captured: Mapping[str, Any], question: str) -> None:
    if authorization_context(state, question) != authorization_context(captured, question):
        raise ValueError("comparison caller authorization context differs from captured project_state bytes")


def workbook_structure_issues(raw: bytes, contract: Mapping[str, Any], entry: Mapping[str, Any]) -> list[str]:
    """Use the existing workbook checker on the captured candidate, never on a reread path."""
    name = "hsk_comparison_workbook_validation"
    module = sys.modules.get(name)
    if module is None:
        spec = importlib.util.spec_from_file_location(name, ROOT / "templates/code/hsk_pipeline/workbook_validation.py")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
    try:
        methods = entry.get("analysis_methods", ())
        methods = methods if isinstance(methods, (list, tuple)) else ()
        plan = entry.get("analysis_comparison")
        module.validate_tables(module.read_workbook_tables(io.BytesIO(raw)), "result_analysis", schema=contract,
                               analysis_methods=methods,
                               comparison_plan=plan if isinstance(plan, Mapping) else None)
    except (ValueError, TypeError, KeyError) as exc:
        return [f"comparison workbook structure: {exc}"]
    return []


def present(entry: Mapping[str, Any], config=None, receipt=None) -> bool:
    return COMPARISON.enabled(entry if isinstance(entry, Mapping) else {}, config=config, receipt=receipt)


def receipt_from_bytes(raw: bytes) -> tuple[dict, list[str]]:
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if "运行配置" not in book.sheetnames:
            return {}, ["comparison workbook is missing 运行配置"]
        rows = list(book["运行配置"].iter_rows(values_only=True))
        if not rows or tuple(rows[0][:2]) != ("项目", "值"):
            return {}, ["comparison 运行配置 must have 项目/值 headers"]
        result, issues = {}, []
        for row in rows[1:]:
            if not row or row[0] in (None, ""):
                continue
            key = str(row[0]).strip()
            if key in result:
                issues.append(f"duplicate comparison receipt field: {key}")
            result[key] = row[1] if len(row) > 1 else None
        return result, issues
    finally:
        book.close()


def _metadata_present(raw: bytes) -> bool:
    """Discover markers without handing an unidentified candidate to openpyxl."""
    rules = safe_yaml.safe_load((ROOT / "core/claim_evidence_contract.yaml").read_bytes())
    limits = rules["limits"]
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        infos = archive.infolist()
        if len(infos) > limits["zip_members"]:
            return False
        info = archive.getinfo("xl/workbook.xml")
        if info.file_size > limits["framework_bytes"]:
            return False
        metadata = archive.read(info)
        if b"<!DOCTYPE" in metadata or b"<!ENTITY" in metadata:
            return False
        tree = ET.fromstring(metadata)
        if any(row.get("name") in SHEETS for row in tree.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet")):
            return True
        scanned = 0
        for item in infos:
            if item.filename != "xl/sharedStrings.xml" and not item.filename.startswith("xl/worksheets/"):
                continue
            if item.file_size > limits["source_file_bytes"]:
                continue
            scanned += item.file_size
            if scanned > limits["total_source_bytes"]:
                break
            part = archive.read(item)
            if any(field.encode("utf-8") in part for field in COMPARISON_FIELDS):
                return True
    return False


def workbook_present(path: Path) -> bool:
    """Discover strict activation before a writer takes its immutable State snapshot."""
    try:
        if path.stat().st_size > 32 * 1024 * 1024:
            return False
        return _metadata_present(path.read_bytes())
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, ET.ParseError):
        return False  # The original receipt reader still rejects malformed workbooks.


def capture_workbook(root: Path, path: Path, observed: dict) -> bytes:
    relative = path.resolve().relative_to(root.resolve()).as_posix()
    raw = BOUNDED._read(root, relative, 32 * 1024 * 1024, observed.setdefault("project", {}))
    candidate_preflight(raw, observed)
    return raw


def candidate_preflight(raw: bytes, observed: dict) -> None:
    rules_raw = BOUNDED._read(ROOT, "core/claim_evidence_contract.yaml", 4 * 1024 * 1024,
                              observed.setdefault("skill", {}))
    Workbook(raw, safe_yaml.safe_load(rules_raw))


def required_ids(entry: Mapping[str, Any]) -> list[str]:
    record = entry.get("analysis_comparison")
    checks = record.get("checks", []) if isinstance(record, Mapping) else []
    if not isinstance(checks, list):
        return []
    return [str(row.get("id")) for row in checks if isinstance(row, Mapping)
            and row.get("requirement") == "required"]


def rejection_events(entry: Mapping[str, Any]) -> list[dict]:
    """Only bound current comparison dispositions authorize the new caller scope."""
    record = entry.get("analysis_comparison")
    if not isinstance(record, Mapping) or record.get("protocol_version") != "1.0.0":
        return []
    checks = record.get("checks", [])
    if not isinstance(checks, list):
        return []
    references = {row["disposition_ref"] for row in checks if isinstance(row, Mapping)
                  and isinstance(row.get("disposition_ref"), str)}
    result = []
    for row in entry.get("analysis_evidence_dispositions", []) or []:
        if (not isinstance(row, Mapping) or row.get("id") not in references
                or row.get("status", "current") != "current" or row.get("disposition") != "reject"):
            continue
        scope = row.get("impact_scope")
        event = {"core_answer": "core_answer_rejected", "model_validity": "model_validity_rejected"}.get(scope)
        expected = {"core_answer": "solve_validate", "model_validity": "model_design"}.get(scope)
        if event and row.get("return_stage") == expected:
            result.append({"event": event, "disposition_id": row["id"], "impact_scope": scope})
    return result


def current_actions(entry: Mapping[str, Any]) -> list[dict]:
    record = entry.get("analysis_comparison")
    checks = record.get("checks", []) if isinstance(record, Mapping) else []
    references = {row["disposition_ref"] for row in checks if isinstance(row, Mapping)
                  and isinstance(row.get("disposition_ref"), str)} if isinstance(checks, list) else set()
    return [row for row in entry.get("analysis_evidence_dispositions", []) or []
            if isinstance(row, dict) and row.get("id") in references and row.get("status", "current") == "current"
            and (row.get("disposition") == "modify"
                 or (row.get("disposition") == "reject" and row.get("impact_scope") == "auxiliary_wording"))]


def action_fragment_closure(state: Mapping[str, Any], entry: Mapping[str, Any]) -> dict:
    """Locate existing claim dependencies or exact anchors, without inventing a B1 claim."""
    result = {"fragment_ids": [], "unbound_ids": [], "issues": []}
    actions = current_actions(entry)
    fragments = (state.get("paper_framework") or {}).get("paper_fragments", [])
    if not actions:
        return result
    if not isinstance(fragments, list):
        result["issues"].append("comparison action paper_fragments must be a list")
        return result
    try:
        claim_ids, anchors = set(), set()
        for action in actions:
            target = action.get("target_claim")
            claim_match = [row for row in fragments if isinstance(row, Mapping)
                           and isinstance(row.get("depends_on"), list) and f"claim:{target}" in row["depends_on"]]
            anchor = action.get("paper_or_figure_anchor")
            anchor_matches = [row["id"] for row in fragments if isinstance(row, Mapping) and anchor
                              and (anchor == row.get("id") or anchor == row.get("anchor")
                                   or anchor == f"{row.get('source_file')}:{row.get('anchor')}")]
            if claim_match:
                claim_ids.add(target)
            elif len(anchor_matches) == 1:
                anchors.add(anchor_matches[0])
            else:
                result["unbound_ids"].append(action["id"])
        result["fragment_ids"] = state_transitions.claim_fragment_stale_closure(
            fragments, claim_ids, fragment_ids=anchors)
    except (ValueError, TypeError, KeyError) as exc:
        result["issues"].append(f"comparison action fragment binding: {exc}")
    return result


def inspect_gate(root: Path, state: Mapping[str, Any], question: str, *, boundary="delivery",
                 config=None, receipt=None, code_path: Path | None = None,
                 workbook: Path | None = None, analysis_bytes: bytes | None = None) -> dict:
    root = Path(root).resolve()
    entry = (state.get("subproblems") or {}).get(question, {})
    result = {"enabled": present(entry, config, receipt), "issues": [], "plan_issues": [], "plan": None,
              "observed_sources": {"project": {}, "skill": {}}, "disposition_impacts": []}
    observed = result["observed_sources"]
    try:
        if not result["enabled"]:
            framework_path = root / "模型论文框架.md"
            if framework_path.is_file():
                raw = framework_path.read_bytes()
                marker = f"HSK_ANALYSIS_COMPARISON_BEGIN {question}".encode("utf-8")
                if marker in raw:
                    result["enabled"] = True
                    BOUNDED._read(root, "模型论文框架.md", 2 * 1024 * 1024, observed["project"])
        if not result["enabled"] and config is None and boundary != "plan" and entry.get("result_analysis_code"):
            try:
                _, probe = stage_code.parse_stage_config(stage_code._relative_path(root, entry["result_analysis_code"]))
                result["enabled"] = present(entry, probe, receipt)
            except (OSError, ValueError, SyntaxError):
                pass  # Historical non-config sources retain their original source-binding diagnostics.
        # New sheet presence is also a strict activation source, never a legacy bypass.
        if not result["enabled"] and analysis_bytes is None:
            candidate_path = workbook
            if candidate_path is None and boundary == "current" and entry.get("result_analysis_workbook"):
                # Discovery preserves the existing artifact reader's native-path
                # compatibility. Activated comparisons use strict paths below.
                candidate_path = (root / Path(entry["result_analysis_workbook"]).expanduser()).resolve()
                if not candidate_path.is_relative_to(root):
                    candidate_path = None
            if candidate_path is not None:
                result["enabled"] = workbook_present(Path(candidate_path))
        if analysis_bytes is None and workbook is not None and result["enabled"]:
            relative = Path(workbook).resolve().relative_to(root).as_posix()
            analysis_bytes = BOUNDED._read(root, relative, 32 * 1024 * 1024, observed["project"])
        if analysis_bytes is not None:
            result["enabled"] = result["enabled"] or _metadata_present(analysis_bytes)
            if result["enabled"]:
                candidate_preflight(analysis_bytes, observed)
        if not result["enabled"]:
            return result
        state_bytes = BOUNDED._read(root, "state/project_state.yaml", 2 * 1024 * 1024, observed["project"])
        captured_state = safe_yaml.safe_load(state_bytes)
        assert_authorization_context(state, captured_state, question)
        captured_entry = captured_state["subproblems"][question]
        framework = BOUNDED._read(root, "模型论文框架.md", 2 * 1024 * 1024, observed["project"])
        authorities = {path: BOUNDED._read(ROOT, path, 4 * 1024 * 1024, observed["skill"])
                       for path in AUTHORITIES}
        from validate_model_approval import validate_comparison_approval
        review = None
        if "review_receipt_policy" in state:
            from review_receipt_consumption import evaluate_gate
            review = evaluate_gate(root, gate="model_challenge", questions=[question], comparison_scope=True)
            CONFORMANCE.merge_read_sets(observed, review["observed_sources"])
        approval_issues = validate_comparison_approval(
            root, state, entry, question, specs=framework.decode("utf-8"), review_report=review)
        result["issues"].extend(approval_issues)
        delivered = config
        # Capture both stage source bundles and inputs before callers read/check them again.
        for stage in (() if boundary == "plan" else ("primary", "analysis")):
            relative = entry.get("code" if stage == "primary" else "result_analysis_code")
            if stage == "analysis" and code_path is not None:
                relative = Path(code_path).resolve().relative_to(root).as_posix()
            if not relative:
                if stage == "analysis" and boundary == "delivery" and delivered is None:
                    continue
                raise ValueError(f"comparison {stage} source is missing")
            raw = BOUNDED._read(root, relative, 2 * 1024 * 1024, observed["project"])
            identity = stage_code.script_identity(root / relative)
            _, actual_config = run_config_parser.parse_embedded_config(
                raw.decode("utf-8-sig"), messages=run_config_parser.RETURNED_EXECUTION_MESSAGES,
                backend=identity.backend)
            if stage == "analysis":
                if delivered is not None and delivered != actual_config:
                    raise ValueError("comparison captured source RUN_CONFIG differs from checked configuration")
                delivered = actual_config
            dependencies = actual_config.get("code_dependencies", [])
            if not isinstance(dependencies, list) or len(dependencies) > 63:
                raise ValueError("comparison source dependency count exceeds budget")
            for dependency in dependencies:
                dependency_raw = BOUNDED._read(root, dependency["path"], 2 * 1024 * 1024, observed["project"])
                if hashlib.sha256(dependency_raw).hexdigest() != str(dependency.get("sha256", "")).lower():
                    raise ValueError("comparison source dependency identity differs from delivered configuration")
            for path in input_files(root, declared_input_paths(actual_config)):
                BOUNDED._read(root, path.relative_to(root).as_posix(), 512 * 1024 * 1024, observed["project"])
        if delivered is None:
            # Prerequisite/design checks need the current plan but not a delivered analysis file.
            plan_result = COMPARISON.inspect_plan(entry, question=question, specs=framework.decode("utf-8"))
        else:
            result["issues"].extend(comparison_config_issues(delivered, required=True))
            plan_result = COMPARISON.inspect_plan(entry, question=question, specs=framework.decode("utf-8"),
                                                   config=delivered, receipt=receipt)
        if not plan_result["enabled"]:
            plan_result["issues"].append("comparison activation is missing its complete registry/method/protocol binding")
        result.update({key: value for key, value in plan_result.items() if key not in {"issues", "enabled"}})
        result["plan_issues"] = [*approval_issues, *plan_result["issues"]]
        result["issues"].extend(plan_result["issues"])
        primary_bytes = None
        if boundary in {"delivery", "receipt", "current"}:
            primary = entry.get("solution_workbook")
            primary_bytes = BOUNDED._read(root, primary, 32 * 1024 * 1024, observed["project"])
            accepted = (entry.get("validated_artifact_hashes") or {}).get("solution_workbook")
            if boundary in {"delivery", "current"}:
                captured_accepted = (captured_entry.get("validated_artifact_hashes") or {}).get("solution_workbook")
                if str(accepted or "").lower() != str(captured_accepted or "").lower():
                    result["issues"].append("comparison caller primary accepted hash differs from captured State")
                accepted = captured_accepted
            if not accepted or hashlib.sha256(primary_bytes).hexdigest() != str(accepted).lower():
                result["issues"].append("comparison baseline bytes differ from the accepted primary workbook")
        if boundary in {"receipt", "current"}:
            if analysis_bytes is None:
                relative = entry.get("result_analysis_workbook")
                analysis_bytes = BOUNDED._read(root, relative, 32 * 1024 * 1024, observed["project"])
                candidate_preflight(analysis_bytes, observed)
            if boundary == "current":
                accepted_analysis = (captured_entry.get("validated_artifact_hashes") or {}).get("result_analysis_workbook")
                caller_analysis = (entry.get("validated_artifact_hashes") or {}).get("result_analysis_workbook")
                if str(caller_analysis or "").lower() != str(accepted_analysis or "").lower():
                    result["issues"].append("comparison caller analysis accepted hash differs from captured State")
                if not accepted_analysis or hashlib.sha256(analysis_bytes).hexdigest() != str(accepted_analysis).lower():
                    result["issues"].append("comparison analysis bytes differ from the accepted analysis workbook")
            actual_receipt, receipt_issues = receipt_from_bytes(analysis_bytes)
            result["issues"].extend(receipt_issues)
            if receipt is not None and receipt != actual_receipt:
                result["issues"].append("comparison receipt differs from captured candidate workbook bytes")
            result["issues"].extend(comparison_receipt_issues(actual_receipt, delivered or {}, required=True))
            result["issues"].extend(workbook_structure_issues(
                analysis_bytes, safe_yaml.safe_load(authorities["core/workbook_schema.yaml"]), entry))
            if not result["issues"]:
                evidence = COMPARISON.inspect_evidence(
                    result["plan"], primary_bytes=primary_bytes, analysis_bytes=analysis_bytes,
                    dispositions=entry.get("analysis_evidence_dispositions", []),
                    workbook_contract=safe_yaml.safe_load(authorities["core/workbook_schema.yaml"]),
                    comparison_rules=safe_yaml.safe_load(authorities["core/claim_evidence_contract.yaml"]))
                result.update({key: value for key, value in evidence.items() if key != "issues"})
                result["issues"].extend(evidence["issues"])
                if boundary == "current" and current_actions(entry):
                    closure = action_fragment_closure(state, entry)
                    result["issues"].extend(closure["issues"])
                    result["issues"].append("comparison claim actions remain current; revise the claim and mark its existing disposition resolved before formal consumption")
                    if closure["unbound_ids"]:
                        result["issues"].append("comparison actions have no exact existing claim/fragment anchor: " + ", ".join(closure["unbound_ids"]))
        assert_authorization_context(state, captured_state, question)
        CONFORMANCE.assert_observed(root, observed)
    except (OSError, ValueError, TypeError, KeyError, SyntaxError, zipfile.BadZipFile) as exc:
        result["issues"].append(f"analysis comparison: {exc}")
    result["issues"] = list(dict.fromkeys(result["issues"]))
    return result
