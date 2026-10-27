"""Shared read-only rules for accepted execution provenance and result consumption.

This module does not run solver code, change project state, approve a result, or
infer that edited source is mathematically equivalent. It verifies the original
accepted workbook and its recorded run identity independently of today's source.
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path
from typing import Any, Mapping

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from execution_protocol import is_source_receipt
import artifact_identity as ARTIFACT_IDENTITY
import artifact_fingerprint as ARTIFACT_FINGERPRINT

READ_ONLY_RESULT_SCOPES = frozenset({"figures", "docx", "latex"})
CLEANUP_PHASES = frozenset({
    "figure_evidence", "writing_docx", "writing_latex", "ai_cleanup",
    "latex_compile_quality",
})
SOLVED_STATUSES = frozenset({"solved", "analyzed", "validated", "written", "completed"})
ANALYZED_STATUSES = frozenset({"analyzed", "validated", "written", "completed"})
SOURCE_LAYERS = frozenset({"primary_code", "analysis_code"})
SHA256_RE = re.compile(r"^[0-9a-fA-F]{64}$")


def is_read_only_result_scope(scope: str | None) -> bool:
    return isinstance(scope, str) and scope.strip() in READ_ONLY_RESULT_SCOPES


def is_cleanup_phase(phase: str | None) -> bool:
    return isinstance(phase, str) and phase.strip() in CLEANUP_PHASES


def semantic_identity_is_current(entry: Mapping[str, Any]) -> bool:
    """Reject known semantic drift; missing legacy comparison fields remain read-only compatible."""
    identity = entry.get("semantic_identity_hash")
    validated_identity = entry.get("validated_semantic_identity_hash")
    approved_identity = entry.get("approved_semantic_identity_hash")
    if identity is not None and (not isinstance(identity, str) or not SHA256_RE.fullmatch(identity)):
        return False
    if validated_identity is not None and (
        not isinstance(validated_identity, str) or not SHA256_RE.fullmatch(validated_identity)
    ):
        return False
    if identity is not None and validated_identity is not None and identity.lower() != validated_identity.lower():
        return False
    if identity is not None and approved_identity is not None and (
        not isinstance(approved_identity, str) or identity.lower() != approved_identity.lower()
    ):
        return False
    revision = entry.get("semantic_revision")
    validated_revision = entry.get("validated_semantic_revision")
    approved_revision = entry.get("approved_semantic_revision")
    if revision is not None and validated_revision is not None and validated_revision != revision:
        return False
    if revision is not None and approved_revision is not None and approved_revision != revision:
        return False
    return True


def semantic_identity_is_fully_verified(entry: Mapping[str, Any]) -> bool:
    """Cleanup-only code delivery needs a complete validated semantic hash, not legacy absence."""
    identity = entry.get("semantic_identity_hash")
    validated_identity = entry.get("validated_semantic_identity_hash")
    if (
        not isinstance(identity, str) or not SHA256_RE.fullmatch(identity)
        or not isinstance(validated_identity, str) or not SHA256_RE.fullmatch(validated_identity)
        or identity.lower() != validated_identity.lower()
    ):
        return False
    approved_identity = entry.get("approved_semantic_identity_hash")
    if approved_identity is not None and (
        not isinstance(approved_identity, str) or identity.lower() != approved_identity.lower()
    ):
        return False
    return True


def _stage_fields(stage: str) -> tuple[str, str, str, str, str]:
    if stage == "primary":
        return ("primary_execution_status", "result_quality_status", "solution_workbook", "solution_workbook", "primary_code")
    if stage == "analysis":
        return ("analysis_execution_status", "result_analysis_status", "result_analysis_workbook", "result_analysis_workbook", "analysis_code")
    raise ValueError("stage must be primary or analysis")


def current_source_differs_from_accepted(
    entry: Mapping[str, Any], stage: str, *,
    observed_bundle: Any = None, observed_entry_sha256: Any = None,
) -> bool:
    """Compare the current/delivered identity with the last accepted execution identity."""
    try:
        _, _, _, _, code_layer = _stage_fields(stage)
    except ValueError:
        return False
    selections = entry.get("solver_execution") or {}
    selection = selections.get(stage, {}) if isinstance(selections, Mapping) else {}
    if not isinstance(selection, Mapping):
        selection = {}
    current_bundle = observed_bundle or selection.get("bundle_sha256")
    accepted_bundle = selection.get("validated_bundle_sha256")
    if current_bundle and accepted_bundle and str(current_bundle).lower() != str(accepted_bundle).lower():
        return True
    current_entry = observed_entry_sha256
    if not current_entry:
        current_entry = entry.get("primary_code_sha256" if stage == "primary" else "analysis_code_sha256")
    validated = entry.get("validated_artifact_hashes") or {}
    accepted_entry = validated.get(code_layer) if isinstance(validated, Mapping) else None
    return bool(current_entry and accepted_entry and str(current_entry).lower() != str(accepted_entry).lower())


def can_consume_accepted_result(entry: Mapping[str, Any], stage: str) -> bool:
    """Cheap state-only eligibility predicate; callers must run accepted_result_issues too."""
    try:
        status_field, quality_field, workbook_field, layer, _ = _stage_fields(stage)
    except ValueError:
        return False
    if entry.get(status_field) not in {"accepted", "awaiting_user_execution"} or entry.get(quality_field) != "passed":
        return False
    allowed = SOLVED_STATUSES if stage == "primary" else ANALYZED_STATUSES
    if entry.get("status") not in allowed:
        return False
    validated = entry.get("validated_artifact_hashes") or {}
    if not isinstance(validated, Mapping) or not SHA256_RE.fullmatch(str(validated.get(layer) or "")):
        return False
    if not entry.get(workbook_field) or not semantic_identity_is_current(entry):
        return False
    try:
        stale = set(ARTIFACT_IDENTITY.normalize_stale_layers(entry.get("stale_layers", []) or []))
    except (ValueError, TypeError):
        return False
    relevant = {"data", "primary_code", "solution_workbook", "framework"}
    if stage == "analysis":
        relevant |= {"analysis_code", "result_analysis_workbook"}
    if stale.intersection(relevant):
        return False
    return not (entry.get("artifacts_stale") is True and not stale)


def _workbook_config(path: Path) -> tuple[dict[str, Any], list[str]]:
    if not path.is_file():
        return {}, ["accepted workbook is missing"]
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "运行配置" not in book.sheetnames:
            return {}, ["accepted workbook has no RUN_CONFIG/RUN_RECEIPT sheet"]
        rows = book["运行配置"].iter_rows(values_only=True)
        header = next(rows, None)
        if not header or tuple(header[:2]) != ("项目", "值"):
            return {}, ["accepted workbook RUN_CONFIG/RUN_RECEIPT headers are invalid"]
        mapping: dict[str, Any] = {}
        issues: list[str] = []
        for row in rows:
            if not row or row[0] in (None, ""):
                continue
            key = str(row[0]).strip()
            if key in mapping:
                issues.append(f"accepted workbook repeats run-configuration field: {key}")
            mapping[key] = row[1] if len(row) > 1 else None
        return mapping, issues
    finally:
        book.close()


def accepted_workbook_data_identity(
    root: Path, state: Mapping[str, Any], entry: Mapping[str, Any], stage: str,
    workbook_path: Path | None = None,
) -> tuple[str | None, list[str], list[str]]:
    """Return a currently rechecked modern receipt data identity, if one is available."""
    _, _, workbook_field, _, _ = _stage_fields(stage)
    raw_path = entry.get(workbook_field)
    if not isinstance(raw_path, str) or not raw_path.strip():
        return None, [f"{stage} accepted workbook path is missing"], []
    root = Path(root).resolve()
    path = workbook_path.resolve() if workbook_path is not None else (root / raw_path).resolve()
    if not path.is_relative_to(root):
        return None, ["accepted workbook path escapes project root"], []
    try:
        config, config_issues = _workbook_config(path)
    except (OSError, ValueError, TypeError, KeyError, IndexError, zipfile.BadZipFile, InvalidFileException) as exc:
        return None, [f"accepted workbook receipt cannot be read: {exc}"], []
    if config_issues:
        if "no RUN_CONFIG/RUN_RECEIPT sheet" in "; ".join(config_issues):
            return None, [], [f"{stage} legacy workbook has no machine-readable run configuration"]
        return None, config_issues, []
    if not config:
        return None, [], [f"{stage} legacy workbook has no machine-readable run configuration"]
    if not is_source_receipt(config.get("run_receipt_protocol_version")):
        return None, [], [f"{stage} legacy receipt cannot independently recheck declared inputs"]
    try:
        from stage_inputs import observe_inputs
        observation = observe_inputs(root, config, state)
        issues = list(observation.get("issues", []))
        digest = str(config.get("data_sha256") or "").lower()
        if not SHA256_RE.fullmatch(digest):
            issues.append(f"{stage} accepted receipt has no valid data_sha256")
        if str(config.get("stage") or "") != stage:
            issues.append(f"{stage} accepted receipt records a different stage")
        return (digest if not issues else None), issues, []
    except (OSError, ValueError, TypeError, KeyError, SyntaxError) as exc:
        return None, [f"{stage} accepted receipt input identity cannot be rechecked: {exc}"], []


def accepted_result_issues(
    root: Path, state: Mapping[str, Any], entry: Mapping[str, Any], stage: str,
    workbook_path: Path | None = None,
) -> tuple[list[str], list[str]]:
    """Validate the original accepted result without requiring current source bytes to match it."""
    issues: list[str] = []
    warnings: list[str] = []
    try:
        status_field, quality_field, workbook_field, layer, code_layer = _stage_fields(stage)
    except ValueError as exc:
        return [str(exc)], []
    if entry.get(status_field) not in {"accepted", "awaiting_user_execution"}:
        issues.append(f"{stage} has no accepted result lifecycle record")
    if entry.get(quality_field) != "passed":
        issues.append(f"{stage} accepted result quality is not passed")
    allowed_project_status = SOLVED_STATUSES if stage == "primary" else ANALYZED_STATUSES
    if entry.get("status") not in allowed_project_status:
        issues.append(f"project status does not retain an accepted {stage} result")
    if not semantic_identity_is_current(entry):
        issues.append("current model semantic identity differs from its validated/approved identity")
    elif not semantic_identity_is_fully_verified(entry):
        warnings.append("legacy accepted result has no complete validated semantic hash; no current-source equivalence is inferred")
    raw_path = entry.get(workbook_field)
    if not isinstance(raw_path, str) or not raw_path.strip():
        return [*issues, f"{stage} accepted workbook path is missing"], warnings
    root = Path(root).resolve()
    path = workbook_path.resolve() if workbook_path is not None else (root / raw_path).resolve()
    if not path.is_relative_to(root):
        return [*issues, "accepted workbook path escapes project root"], warnings
    validated = entry.get("validated_artifact_hashes") or {}
    expected_workbook = validated.get(layer) if isinstance(validated, Mapping) else None
    if not isinstance(expected_workbook, str) or not SHA256_RE.fullmatch(expected_workbook):
        issues.append(f"{stage} accepted workbook lacks a validated raw SHA-256")
    else:
        try:
            actual_workbook = hashlib.sha256(path.read_bytes()).hexdigest()
        except OSError as exc:
            actual_workbook = None
            issues.append(f"{stage} accepted workbook cannot be read: {exc}")
        if actual_workbook and actual_workbook != expected_workbook.lower():
            issues.append(f"{stage} accepted workbook bytes differ from validated raw SHA-256")
    try:
        stale = set(ARTIFACT_IDENTITY.normalize_stale_layers(entry.get("stale_layers", []) or []))
    except (ValueError, TypeError) as exc:
        return [*issues, f"accepted result stale-layer shape is invalid: {exc}"], warnings
    relevant = {"data", "primary_code", "solution_workbook", "framework"}
    if stage == "analysis":
        relevant |= {"analysis_code", "result_analysis_workbook"}
    affected = sorted(stale.intersection(relevant))
    if affected:
        issues.append(f"{stage} accepted result has relevant stale layers: {affected}")
    if entry.get("artifacts_stale") is True and not stale:
        issues.append(f"{stage} accepted result has an unexplained artifacts_stale flag")
    data_hash = str(entry.get("data_hash") or "").lower()
    validated_data_hash = str(
        entry.get("validated_data_hash")
        or (validated.get("data") if isinstance(validated, Mapping) else "")
        or ""
    ).lower()
    if data_hash and validated_data_hash and data_hash != validated_data_hash:
        issues.append("current data identity differs from the accepted result data identity")
    if not validated_data_hash:
        warnings.append("accepted result has no separately stored validated_data_hash")
    expected_framework = validated.get("framework") if isinstance(validated, Mapping) else None
    framework_path = root / "模型论文框架.md"
    framework_anchor = str(entry.get("framework_section") or "")
    if expected_framework and framework_anchor:
        actual_framework = ARTIFACT_FINGERPRINT.framework_section_hash(framework_path, framework_anchor)
        if not actual_framework:
            issues.append("accepted result framework section is missing or cannot be hashed")
        elif str(actual_framework).lower() != str(expected_framework).lower():
            issues.append("current paper-framework section differs from the accepted result framework identity")
    elif expected_framework:
        warnings.append("accepted result has a framework hash but no section anchor for independent read-only verification")
    else:
        warnings.append("accepted result has no validated framework-section identity")
    selections = entry.get("solver_execution") or {}
    selection = selections.get(stage, {}) if isinstance(selections, Mapping) else {}
    selection = selection if isinstance(selection, Mapping) else {}
    accepted_bundle = str(selection.get("validated_bundle_sha256") or "").lower()
    current_bundle = str(selection.get("bundle_sha256") or "").lower()
    try:
        config, config_issues = _workbook_config(path)
    except (OSError, ValueError, TypeError, zipfile.BadZipFile) as exc:
        issues.append(f"{stage} accepted workbook receipt cannot be read: {exc}")
        return list(dict.fromkeys(issues)), list(dict.fromkeys(warnings))
    if config_issues:
        if "no RUN_CONFIG/RUN_RECEIPT sheet" in "; ".join(config_issues):
            warnings.append(f"{stage} legacy workbook has no machine-readable run configuration")
        else:
            issues.extend(config_issues)
    if not config:
        warnings.append(f"{stage} legacy workbook has no machine-readable run configuration")
    if config:
        declared_stage = str(config.get("stage") or "").strip()
        if declared_stage and declared_stage != stage:
            issues.append(f"{stage} accepted workbook records a different execution stage")
        receipt_version = config.get("run_receipt_protocol_version")
        if is_source_receipt(receipt_version):
            if str(config.get("run_receipt_version") or "") != str(receipt_version):
                issues.append(f"{stage} accepted workbook RUN_RECEIPT version does not match RUN_CONFIG")
            if not SHA256_RE.fullmatch(accepted_bundle):
                issues.append(f"{stage} modern accepted result has no validated source bundle identity")
            elif str(config.get("code_bundle_sha256") or "").lower() != accepted_bundle:
                issues.append(f"{stage} accepted workbook bundle differs from its validated execution bundle")
            accepted_code_sha = str(
                (validated.get(code_layer) if isinstance(validated, Mapping) else "") or ""
            ).lower()
            receipt_code_sha = str(config.get("code_sha256") or "").lower()
            if accepted_code_sha and receipt_code_sha and accepted_code_sha != receipt_code_sha:
                issues.append(f"{stage} accepted workbook entry SHA differs from validated source identity")
            elif not accepted_code_sha:
                warnings.append(f"{stage} accepted entry SHA is retained only through the bundle receipt")
            receipt_data = str(config.get("data_sha256") or "").lower()
            expected_data_identity = validated_data_hash or data_hash
            if expected_data_identity and receipt_data != expected_data_identity:
                issues.append(f"{stage} accepted workbook data hash differs from validated data identity")
            try:
                from stage_inputs import observe_inputs
                observation = observe_inputs(root, config, state)
                issues.extend(f"{stage} accepted receipt input: {item}" for item in observation.get("issues", []))
            except (OSError, ValueError, TypeError, KeyError, SyntaxError) as exc:
                issues.append(f"{stage} accepted receipt input identity cannot be rechecked: {exc}")
        elif accepted_bundle:
            warnings.append(f"{stage} legacy receipt cannot independently echo the accepted source bundle")
        if stage == "analysis":
            primary_validated = validated.get("solution_workbook") if isinstance(validated, Mapping) else None
            upstream = str(config.get("primary_workbook_sha256") or "").lower()
            if is_source_receipt(receipt_version) and not SHA256_RE.fullmatch(upstream):
                issues.append("analysis accepted receipt lacks a valid primary_workbook_sha256")
            elif upstream and primary_validated and upstream != str(primary_validated).lower():
                issues.append("analysis receipt is bound to a different primary workbook")
    if current_bundle and accepted_bundle and current_bundle != accepted_bundle:
        warnings.append(
            f"{stage} current source bundle differs from the accepted execution bundle; "
            "only the original accepted workbook is being consumed"
        )
    return list(dict.fromkeys(issues)), list(dict.fromkeys(warnings))
