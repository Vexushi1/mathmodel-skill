"""Instantiate in a question folder; no model approval or acceptance is invented.

The maintenance fixture also instantiates the primary branch to produce its QR
baseline. An analysis instance reads that accepted workbook without calling it.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
from pathlib import Path, PurePosixPath
import platform
import random
import tempfile
import time

import numpy as np
import openpyxl

from comparison_numerics import fit_polynomial

RUN_CONFIG = {}


def bounded(root, relative, *, existing=True):
    if not isinstance(relative, str) or not relative or ":" in relative or "\\" in relative:
        raise ValueError("Use a project-relative POSIX path")
    parts = relative.split("/")
    if PurePosixPath(relative).is_absolute() or any(p in ("", ".", "..") for p in parts):
        raise ValueError("Invalid relative path")
    path = root
    for part in parts:
        path = path / part
        if path.is_symlink():
            raise ValueError("Path aliases are not supported")
    if not path.resolve().is_relative_to(root) or (existing and not path.is_file()):
        raise ValueError("Missing file or path outside the project")
    return path


def file_sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def files_digest(root, paths):
    if len(set(p.casefold() for p in paths)) != len(paths):
        raise ValueError("Duplicate or case-aliased paths")
    digest = hashlib.sha256()
    for relative in sorted(paths, key=lambda value: value.encode("utf-8")):
        digest.update(relative.encode("utf-8") + b"\0")
        digest.update(bytes.fromhex(file_sha(bounded(root, relative))))
    return digest.hexdigest()


def observe(root, code, config):
    if config["run_receipt_protocol_version"] != "1.1.0" or config["solver_backend"] != "python":
        raise ValueError("This combined-input example uses Python receipt 1.1")
    if files_digest(root, config["data_paths"]) != config["data_sha256"].lower():
        raise ValueError("Input identity changed")
    sources = [code.relative_to(root).as_posix()]
    for dependency in config["code_dependencies"]:
        if file_sha(bounded(root, dependency["path"])) != dependency["sha256"].lower():
            raise ValueError("Declared helper identity changed")
        sources.append(dependency["path"])
    helper = bounded(root, (code.parent / "comparison_numerics.py").relative_to(root).as_posix())
    if Path(fit_polynomial.__code__.co_filename).resolve() != helper:
        raise ValueError("The declared numerical helper was not actually imported")
    identity = {"data_sha256": config["data_sha256"], "code_sha256": file_sha(code),
                "code_bundle_sha256": files_digest(root, sources)}
    if config["stage"] == "analysis":
        primary = bounded(root, config["primary_workbook"])
        if file_sha(primary) != config["primary_workbook_sha256"].lower():
            raise ValueError("Accepted primary workbook changed")
        identity["primary_workbook_sha256"] = file_sha(primary)
    return identity


def make_receipt(config, identity, elapsed):
    result = {key: config[key] for key in ("stage", "problem_name", "solver_backend", "solver", "random_seed",
                                          "tolerance", "iteration_or_time_limit")}
    result.update(identity)
    result.update(run_receipt_version="1.1.0", execution_owner="user", execution_profile="full_fidelity",
                  solver_version=np.__version__, actual_stop_reason="all_declared_decompositions_completed",
                  repetitions_or_scenarios=1 if config["stage"] == "primary" else 2,
                  grid_or_time_range="all declared training and evaluation points", fallback_used=False,
                  platform=platform.platform(), elapsed_seconds=elapsed)
    for field in ("allow_reduced_data", "allow_coarser_grid", "allow_shorter_horizon", "allow_fewer_repetitions",
                  "allow_relaxed_tolerance", "allow_silent_solver_fallback"):
        result[field] = False
    if config["stage"] == "primary":
        result["primary_quality_protocol_version"] = config["primary_quality_protocol_version"]
    else:
        for field in ("analysis_comparison_protocol_version", "analysis_comparison_plan_sha256"):
            result[field] = config[field]
    return result


def primary_tables(config, payload):
    coefficients, predictions, residual = fit_polynomial(payload["train_t"], payload["train_y"], 2, "QR", payload["evaluation_t"])
    rows = [[f"prediction-{i}", "prediction", float(value), "dimensionless", scenario, "prediction"]
            for i, (scenario, value) in enumerate(zip(payload["scenarios"], predictions))]
    rows.extend([[f"coefficient-{i}", "coefficient", float(value), "dimensionless", f"coefficient-{i}", "coefficient"]
                 for i, value in enumerate(coefficients)])
    passed = residual <= config["tolerance"]
    return {
        "核心指标": [["指标", "数值"], ["留出预测", predictions[-1]], ["正规方程残差", residual]],
        "数据审计": [["等级", "检查项", "信息", "处理方式"], ["info", "有限成对样本", "完整原始训练和评价点", "只读检查"]],
        "主结果质量门": [["Verification ID", "检查项", "是否通过", "证据", "判定关系", "阈值或容差", "实际值", "证据工作表", "阈值来源"],
                        ["PQ-Q1-01", "最小二乘正规方程残差", passed, "从完整训练设计矩阵重算", "abs<=", config["tolerance"], residual, "均衡残差", "solver_tolerance"]],
        "均衡残差": [["主体或均衡", "残差", "容差", "是否满足"], ["X.T*(X*c-y)=0", residual, config["tolerance"], passed]],
        "状态明细": [["记录键", "状态", "数值", "单位", "实例或场景", "指标"], *rows],
    }


def analysis_tables(root, config, payload):
    book = openpyxl.load_workbook(bounded(root, config["primary_workbook"]), read_only=True, data_only=True)
    try:
        records = {row["记录键"]: row for row in _records(book["状态明细"])}
        baseline = [float(records[f"prediction-{i}"]["数值"]) for i in range(len(payload["evaluation_t"]))]
    finally:
        book.close()
    _, linear, linear_residual = fit_polynomial(payload["train_t"], payload["train_y"], 1, "QR", payload["evaluation_t"])
    _, svd, svd_residual = fit_polynomial(payload["train_t"], payload["train_y"], 2, "SVD", payload["evaluation_t"])
    model_rows, algorithm_rows = [], []
    for i, scenario in enumerate(payload["scenarios"]):
        model_difference, algorithm_difference = linear[i] - baseline[i], svd[i] - baseline[i]
        model_rows.append(["CMP-Q1-01", f"model-{i}", "MODEL-Q1-01", "MODEL-Q1-02", "EVAL-Q1-01", scenario,
                           "prediction", "dimensionless", baseline[i], linear[i], "difference", model_difference,
                           "CRIT-CMP-Q1-01", abs(model_difference) >= payload["model_difference_threshold"], linear_residual, f"QR/NumPy {np.__version__}", "decomposition_completed"])
        algorithm_rows.append(["CMP-Q1-02", f"algorithm-{i}", "MODEL-Q1-01", "ALGO-Q1-01", "ALGO-Q1-02", "EVAL-Q1-02", scenario,
                               1, "prediction", "dimensionless", baseline[i], svd[i], "difference", algorithm_difference,
                               "CRIT-CMP-Q1-02", abs(algorithm_difference) <= config["tolerance"], svd_residual, f"SVD/NumPy {np.__version__}", "decomposition_completed"])
    model_kept = bool(model_rows[-1][13])
    algorithm_kept = all(row[15] for row in algorithm_rows)
    return {
        "分析设计": [["风险来源", "分析问题", "方法", "指标", "通过标准", "检验ID", "判据ID"],
                     ["模型结构", "二次项是否影响共同留出预测", "多模型检验", "prediction", "预设留出差异下界", "CMP-Q1-01", "CRIT-CMP-Q1-01"],
                     ["求解方法", "同一二次模型QR与SVD是否一致", "同模型多算法检验", "prediction", "预设数值容差", "CMP-Q1-02", "CRIT-CMP-Q1-02"]],
        "多模型检验": [["检验ID", "记录键", "主模型ID", "对照模型ID", "评价协议ID", "实例或场景", "指标", "单位", "主模型数值", "对照模型数值", "差异类型", "差异", "判据ID", "判定", "残差", "求解器及版本", "停止原因"], *model_rows],
        "同模型多算法检验": [["检验ID", "记录键", "模型ID", "基准算法ID", "对照算法ID", "评价协议ID", "实例或场景", "重复编号", "指标", "单位", "基准数值", "对照数值", "差异类型", "差异", "判据ID", "判定", "残差", "求解器版本", "停止原因"], *algorithm_rows],
        "结论稳定性汇总": [["核心结论", "分析方法", "稳定范围", "是否保持", "检验ID", "证据工作表"],
                           ["此案例二次项影响留出预测", "多模型检验", "声明的共同留出点", model_kept, "CMP-Q1-01", "多模型检验"],
                           ["同模型预测对QR/SVD一致", "同模型多算法检验", "全部声明评价点", algorithm_kept, "CMP-Q1-02", "同模型多算法检验"]],
    }


def _records(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def write_workbook(path, sheets):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    book.remove(book.active)
    temporary = None
    try:
        for name, rows in sheets.items():
            if len(rows) < 2 or len(set(rows[0])) != len(rows[0]):
                raise ValueError("Nonempty sheets with unique headers are required")
            sheet = book.create_sheet(name)
            for row in rows:
                if any(isinstance(value, float) and not math.isfinite(value) for value in row):
                    raise ValueError("Nonfinite workbook value")
                sheet.append(row)
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        book.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()


def main():
    config, code = RUN_CONFIG, Path(__file__).resolve()
    root = code.parents[1]
    if config["solver"] != ("explicit_qr" if config["stage"] == "primary" else "explicit_qr_svd_comparison") or config["iteration_or_time_limit"] != "direct":
        raise ValueError("Instantiate the declared complete direct-decomposition policy")
    identity = observe(root, code, config)
    random.seed(config["random_seed"])
    started = time.perf_counter()
    payload = json.loads(bounded(root, config["data_paths"][0]).read_text(encoding="utf-8"))
    if len(payload["scenarios"]) != len(payload["evaluation_t"]) or not payload["scenarios"] or len(set(payload["scenarios"])) != len(payload["scenarios"]):
        raise ValueError("Every declared evaluation point needs one unique scenario")
    sheets = primary_tables(config, payload) if config["stage"] == "primary" else analysis_tables(root, config, payload)
    receipt = make_receipt(config, identity, time.perf_counter() - started)
    sheets["运行配置"] = [["项目", "值"], *[list(item) for item in receipt.items()]]
    if observe(root, code, config) != identity:
        raise ValueError("Source, input or primary workbook changed during execution")
    workbook = bounded(root, config["expected_workbook"], existing=False)
    write_workbook(workbook, sheets)
    if config["stage"] == "primary" and not sheets["主结果质量门"][1][2]:
        raise RuntimeError("The actual primary workbook records a failed numerical check")
    return workbook


if __name__ == "__main__":
    print(main())
