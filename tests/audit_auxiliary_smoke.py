"""Native synthetic receipt-1.2 smoke, also used by unit tests and MATLAB CI."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import os
from pathlib import Path
import subprocess
import sys

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
for path in (ROOT, ROOT / "scripts"):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))
import analysis_prerequisites as prerequisites
import stage_code
import validate_code_delivery as delivery
import validate_user_execution as receipts
from tests.test_solver_backend_end_to_end import instantiate, file_hash, reference_digest, save_state
from tests.test_solver_backend_preprocessing import prepare_native_preprocessing


def instantiate_stage(root, backend, stage, state):
    entry = state["subproblems"]["Q1"]
    config = {"stage": stage, "problem_name": "问题一", "solver_backend": backend,
              "run_receipt_protocol_version": "1.2.0", "data_identity_mode": "preprocessing_workbook",
              "data_paths": [state["preprocessing"]["workbook"]],
              "data_sha256": state["preprocessing"]["workbook_sha256"],
              "auxiliary_data_paths": ["constraints.json"],
              "auxiliary_data_sha256": reference_digest(root, ["constraints.json"]),
              "solver": "mldivide" if backend == "matlab" else "direct", "random_seed": 2026,
              "tolerance": 1e-10, "iteration_or_time_limit": "direct", "code_dependencies": [],
              "expected_workbook": f"问题一求解/问题一{'求解结果' if stage == 'primary' else '结果深化分析'}.xlsx"}
    if stage == "primary":
        config["primary_quality_protocol_version"] = "1.0.0"
    else:
        config["primary_workbook_sha256"] = file_hash(root / entry["solution_workbook"])
        config["solver"] = "analytic_parameter_sweep" if backend == "matlab" else "direct"
        config["iteration_or_time_limit"] = 3 if backend == "matlab" else "direct"
        entry.update(result_analysis_status="pending", analysis_methods=["参数敏感性"],
                     result_analysis_requirement_reason="Synthetic coefficient stress test")
    if backend == "matlab":
        code = instantiate(root, stage, **{key: value for key, value in config.items() if key != "stage"})
        text = code.read_text(encoding="utf-8")
        # Model-specific synthetic consumer: the attachment materially changes the answer 3 -> 4.
        marker = "input = read_model_input(root, RUN_CONFIG);"
        assert marker in text
        text = text.replace(marker, marker + "\nauxiliary = jsondecode(fileread(project_path(root, RUN_CONFIG.auxiliary_data_paths{1})));\ninput.right_hand_side = input.right_hand_side + auxiliary.right_hand_side_offset;")
        code.write_text(text, encoding="utf-8")
    else:
        code = root / "问题一求解" / ("问题一求解.py" if stage == "primary" else "问题一结果深化分析.py")
        text = (ROOT / "tests/fixtures/solver_backends/python_auxiliary.py").read_text(encoding="utf-8")
        code.write_text(text.replace("RUN_CONFIG = {}", f"RUN_CONFIG = {config!r}"), encoding="utf-8")
    save_state(root, state)
    return code, config


def prepare_project(root, backend):
    state, _ = prepare_native_preprocessing(root)
    (root / "constraints.json").write_text('{"right_hand_side_offset":2}', encoding="utf-8")
    state["execution"] = {"solver_backend": backend, "solver_backend_selection_reason": "Synthetic whole-project auxiliary fixture"}
    state["subproblems"]["Q1"].update(selected_model="a*x=b+offset", result_quality_status="pending",
                                        result_analysis_status="pending", capabilities={"requires_equilibrium_residual": True})
    save_state(root, state)
    return state


def deliver_stage(root, backend, stage, state, matlab_command=None):
    code, config = instantiate_stage(root, backend, stage, state)
    native = {}
    errors, parsed = delivery.validate_script(root, code, stage, matlab_command=matlab_command,
                                              require_native=backend == "matlab" and bool(matlab_command), native_report=native)
    assert errors == [], errors
    delivery.update_state(root, parsed, code)
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    return state, code, config, native


def run_smoke(root, backend, matlab_command=None):
    state = prepare_project(root, backend)
    preprocessing_bytes = (root / state["preprocessing"]["workbook"]).read_bytes()
    outcomes = []
    primary_bytes = None
    for stage in ("primary", "analysis"):
        state, code, config, native = deliver_stage(root, backend, stage, state, matlab_command)
        if backend == "python":
            command = [sys.executable, "-B", str(code)]
        else:
            assert matlab_command, "Native MATLAB command required"
            directory = str(code.parent).replace("'", "''")
            command = [matlab_command, "-batch", f"cd('{directory}'); {code.stem}();"]
        result = subprocess.run(command, capture_output=True, text=True, env={**os.environ, "PYTHONUTF8": "1"})
        (root / f"{stage}_native.log").write_text(result.stdout + result.stderr, encoding="utf-8")
        assert result.returncode == 0, result.stdout + result.stderr
        workbook = root / config["expected_workbook"]
        errors = receipts.validate_one(root, workbook, state, True)
        assert errors == [], errors
        metadata, errors = receipts.configuration_map(workbook)
        assert errors == [] and metadata["run_receipt_version"] == "1.2.0"
        assert metadata["auxiliary_data_sha256"] == reference_digest(root, ["constraints.json"])
        book = openpyxl.load_workbook(workbook, data_only=True)
        try:
            if stage == "primary":
                values = dict(book["核心指标"].iter_rows(min_row=2, values_only=True))
                assert values["解"] == 4, values  # Auxiliary input changes the actual result.
                primary_bytes = workbook.read_bytes()
            else:
                values = [row[3] for row in book["参数敏感性"].iter_rows(min_row=2, values_only=True)]
                assert all(abs(a-b) < 1e-12 for a,b in zip(values, [8/1.5, 4, 8/3]))
                assert (root / state["subproblems"]["Q1"]["solution_workbook"]).read_bytes() == primary_bytes
        finally:
            book.close()
        save_state(root, state)
        outcomes.append({"stage": stage, "native_exit": result.returncode, "code_analyzer": native, "receipt_accepted": True})
    assert (root / state["preprocessing"]["workbook"]).read_bytes() == preprocessing_bytes
    before = (root / "constraints.json").read_bytes()
    (root / "constraints.json").write_text('{"right_hand_side_offset":99}', encoding="utf-8")
    try:
        assert prerequisites.primary_issues(root, state, state["subproblems"]["Q1"], require_project_policy=True)
        assert receipts.validate_one(root, root / state["subproblems"]["Q1"]["solution_workbook"], deepcopy(state), False)
    finally:
        (root / "constraints.json").write_bytes(before)
    result = {"backend": backend, "synthetic_only": True, "auxiliary_changes_answer": True,
              "preprocessing_preserved": True, "primary_preserved": True, "drift_rejected": True, "stages": outcomes}
    (root / "auxiliary_smoke_report.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--backend", choices=("python", "matlab"), required=True)
    parser.add_argument("--matlab-command")
    args = parser.parse_args()
    print(json.dumps(run_smoke(args.project.resolve(), args.backend, args.matlab_command), ensure_ascii=False))
