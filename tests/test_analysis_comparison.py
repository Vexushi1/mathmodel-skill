"""Captured-evidence comparison boundaries, not task-code execution or real user approval."""
from __future__ import annotations

from copy import deepcopy
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import yaml
from jsonschema import Draft202012Validator

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import analysis_comparison as COMPARISON
import semantic_identity as SEMANTIC
import validate_model_approval as APPROVAL


def model_identity(degree=2):
    return {
        "schema_version": "1.0.0", "question": "Q1", "research_object": "synthetic polynomial prediction",
        "data_scope": [{"id": "D1", "source": "synthetic declared training", "role": "training"}],
        "variables": [{"id": "V1", "symbol": "beta", "domain": f"R^{degree + 1}"}],
        "parameters": [{"id": "P1", "symbol": "degree", "value": degree}], "assumptions": [],
        "objective": {"expression": "sum squared training residual", "sense": "minimize"},
        "constraints": [{"id": "C1", "expression": f"prediction=sum(beta[k]*t^k,k=0..{degree})"}],
        "preprocessing_decision": "not_needed", "dependencies": [],
        "algorithm_semantics": {"model_family": "polynomial_least_squares", "solver": "QR"},
    }


def comparison_scope():
    main = model_identity()
    return {
        "protocol_version": "1.0.0", "question": "Q1",
        "baseline_semantic_identity_hash": SEMANTIC.semantic_identity_hash(main),
        "models": [{"id": "MODEL-Q1-01", "identity": main, "pure_algorithm_keys": ["solver"]},
                   {"id": "MODEL-Q1-02", "identity": model_identity(1), "pure_algorithm_keys": ["solver"]}],
        "algorithms": [
            {"id": "ALGO-Q1-01", "model_ref": "MODEL-Q1-01", "definition": {
                "family": "QR", "update_rule": "orthogonal triangular factorization", "stop_rule": "one factorization"},
             "implementation_anchor": "问题一求解/q1_solver.py:fit_qr"},
            {"id": "ALGO-Q1-02", "model_ref": "MODEL-Q1-01", "definition": {
                "family": "SVD", "update_rule": "singular-vector factorization", "stop_rule": "one factorization"},
             "implementation_anchor": "问题一求解/q1_analysis.py:fit_svd"},
        ],
        "questions": [
            {"id": "CMP-Q1-01", "kind": "model_comparison", "baseline_ref": "MODEL-Q1-01",
             "candidate_refs": ["MODEL-Q1-02"], "question": "Does removing quadratic structure change prediction?",
             "target_claim": "C-model", "evaluation": {"id": "EVAL-Q1-01", "metric": "prediction", "unit": "dimensionless",
                 "direction": "higher", "fixed_axes": ["scenario"], "allowed_changes": ["degree"]}},
            {"id": "CMP-Q1-02", "kind": "algorithm_comparison", "baseline_ref": "ALGO-Q1-01",
             "candidate_refs": ["ALGO-Q1-02"], "question": "Does SVD preserve the same quadratic prediction?",
             "target_claim": "C-algorithm", "evaluation": {"id": "EVAL-Q1-02", "metric": "prediction", "unit": "dimensionless",
                 "direction": "higher", "fixed_axes": ["scenario"], "allowed_changes": ["factorization"]}},
        ],
    }


def selector(sheet, key, value, axis=None, axis_column=None):
    identities = {"metric": "指标", "scenario": "实例或场景"}
    if axis:
        identities[axis] = axis_column
    return {"sheet": sheet, "header_row": 1, "row_key": key, "expected_cardinality": 1,
            "value_type": "scalar", "value_column": value, "identity_columns": identities,
            "unit": {"kind": "column", "column": "单位"}}


def comparison_entry(scope=None):
    scope = scope or comparison_scope()
    digest = COMPARISON.scope_sha256(scope)
    baseline = scope["baseline_semantic_identity_hash"]
    checks = []
    for index, question in enumerate(scope["questions"], 1):
        is_model = question["kind"] == "model_comparison"
        axis = "model" if is_model else "algorithm"
        sheet = COMPARISON.KINDS[question["kind"]]
        record_id = f"REC{index}"
        criterion_id = "CRIT-" + question["id"]
        row_key = {"检验ID": question["id"], "记录键": record_id,
                   "评价协议ID": question["evaluation"]["id"], "判据ID": criterion_id, "差异类型": "difference"}
        if is_model:
            row_key.update({"主模型ID": question["baseline_ref"], "对照模型ID": question["candidate_refs"][0]})
            candidate_col, baseline_col, identity_col = "对照模型数值", "主模型数值", "对照模型ID"
        else:
            row_key.update({"模型ID": "MODEL-Q1-01", "基准算法ID": question["baseline_ref"], "对照算法ID": question["candidate_refs"][0]})
            candidate_col, baseline_col, identity_col = "对照数值", "基准数值", "对照算法ID"
        evidence = {"id": record_id, "candidate_ref": question["candidate_refs"][0],
            "baseline": {"source": "primary", "selector": selector("核心指标", {"指标": "prediction", "实例或场景": "holdout"}, "数值")},
            "candidate": {"source": "analysis", "selector": selector(sheet, deepcopy(row_key), candidate_col, axis, identity_col)},
            "reported_baseline": {"source": "analysis", "selector": selector(sheet, deepcopy(row_key), baseline_col)},
            "reported_difference": {"source": "analysis", "selector": selector(sheet, deepcopy(row_key), "差异")},
            "operation": {"op": "difference", "comparison_axis": axis}}
        checks.append({key: deepcopy(question[key]) for key in ("id", "kind", "question", "target_claim", "baseline_ref", "candidate_refs")} | {
            "requirement": "required", "requirement_source": "claim:" + question["target_claim"],
            "criterion": {"id": criterion_id, "relation": "abs_ge" if is_model else "abs_le",
                          "threshold": "1" if is_model else "1e-10", "unit": "dimensionless", "source": "declared fixture criterion"},
            "evidence_refs": [evidence], "disposition_ref": f"E{index}"})
    approval = {"status": "approved", "owner": "user", "statement": "Approve this synthetic scope for fixture testing.",
                "source_ref": "synthetic:test-only", "approved_scope_sha256": digest,
                "reviews": [{"review_id": f"synthetic-review-{index}", "role": role, "scope_sha256": digest,
                             "verdict": "passed", "method": "separate_pass", "conclusion": "Synthetic definition/boundary/independence fixture.",
                             "blocking_items": [], "unresolved_items": []}
                            for index, role in enumerate(sorted(COMPARISON.REVIEW_ROLES), 1)]}
    return {"semantic_revision": 1, "approved_semantic_revision": 1,
            "semantic_identity_schema_version": "1.0.0", "semantic_identity_hash": baseline,
            "validated_semantic_identity_hash": baseline, "approved_semantic_identity_hash": baseline,
            "model_challenge_status": "passed", "human_model_approval_status": "approved",
            "analysis_methods": ["多模型检验", "同模型多算法检验"], "result_analysis_status": "pending",
            "analysis_comparison": {"protocol_version": "1.0.0", "scope_ref": "Q1", "scope_sha256": digest,
                "baseline_semantic_identity_hash": baseline, "approval_binding": approval, "checks": checks},
            "analysis_evidence_dispositions": [
                {"id": "E1", "target_claim": "C-model", "method_or_source": "declared synthetic model records",
                 "disposition": "support", "key_finding": "linear candidate changes prediction", "required_action": "limit claim to synthetic holdout"},
                {"id": "E2", "target_claim": "C-algorithm", "method_or_source": "declared synthetic algorithm records",
                 "disposition": "support", "key_finding": "same quadratic prediction", "required_action": "limit claim to checked methods"},
            ]}


def comparison_framework(scope=None):
    scope = scope or comparison_scope()
    main = scope["models"][0]["identity"]
    return ("### Q1：Synthetic fixture\n#### 当前模型口径\n<!-- HSK_SEMANTIC_IDENTITY_BEGIN Q1 -->\n```yaml\n"
            + yaml.safe_dump(main, allow_unicode=True, sort_keys=False)
            + "```\n<!-- HSK_SEMANTIC_IDENTITY_END Q1 -->\n#### 结果摘要\nCurrent synthetic baseline.\n"
            + "#### 03B比较计划与证据\n<!-- HSK_ANALYSIS_COMPARISON_BEGIN Q1 -->\n```yaml\n"
            + yaml.safe_dump(scope, allow_unicode=True, sort_keys=False)
            + "```\n<!-- HSK_ANALYSIS_COMPARISON_END Q1 -->\n")


def workbook_bytes(sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def captured_books(*, algorithm_prediction=16, algorithm_difference=0, algorithm_decision=True, model_value=11, mutation=None):
    primary = workbook_bytes({"核心指标": [["指标", "实例或场景", "单位", "数值"], ["prediction", "holdout", "dimensionless", 16]]})
    model_headers = ["检验ID", "记录键", "主模型ID", "对照模型ID", "评价协议ID", "实例或场景", "指标", "单位",
                     "主模型数值", "对照模型数值", "差异类型", "差异", "判据ID", "判定"]
    algorithm_headers = ["检验ID", "记录键", "模型ID", "基准算法ID", "对照算法ID", "评价协议ID", "实例或场景", "重复编号",
                         "指标", "单位", "基准数值", "对照数值", "差异类型", "差异", "判据ID", "判定"]
    sheets = {
        "多模型检验": [model_headers, ["CMP-Q1-01", "REC1", "MODEL-Q1-01", "MODEL-Q1-02", "EVAL-Q1-01", "holdout",
                       "prediction", "dimensionless", 16, model_value, "difference", model_value - 16, "CRIT-CMP-Q1-01", True]],
        "同模型多算法检验": [algorithm_headers, ["CMP-Q1-02", "REC2", "MODEL-Q1-01", "ALGO-Q1-01", "ALGO-Q1-02", "EVAL-Q1-02",
                            "holdout", 1, "prediction", "dimensionless", 16, algorithm_prediction, "difference", algorithm_difference,
                            "CRIT-CMP-Q1-02", algorithm_decision]],
    }
    if mutation:
        mutation(sheets)
    return primary, workbook_bytes(sheets)


class AnalysisComparisonTests(unittest.TestCase):
    def setUp(self):
        self.scope = comparison_scope()
        self.entry = comparison_entry(self.scope)
        self.rules = yaml.safe_load((ROOT / "core/claim_evidence_contract.yaml").read_text(encoding="utf-8"))

    def plan(self):
        result = COMPARISON.inspect_plan(self.entry, question="Q1", specs=self.scope)
        self.assertEqual(result["issues"], [])
        return result

    def evidence(self, **kwargs):
        primary, analysis = captured_books(**kwargs)
        return COMPARISON.inspect_evidence(self.plan()["plan"], primary_bytes=primary, analysis_bytes=analysis,
                                          dispositions=self.entry["analysis_evidence_dispositions"], comparison_rules=self.rules)

    def dimensioned_relative_evidence(self, operation="relative_change", *, unit="ratio", reported_value=.5,
                                     explicit=True, mutation=None):
        scope = deepcopy(self.scope)
        for question in scope["questions"]:
            question["evaluation"].update(unit="m", direction="lower")
        entry = comparison_entry(scope)
        entry["analysis_comparison"]["checks"][0]["criterion"]["unit"] = "m"
        check = entry["analysis_comparison"]["checks"][1]
        check["criterion"].update(relation="ge", threshold="0.4", unit="ratio")
        record = check["evidence_refs"][0]
        record["operation"]["op"] = operation
        if operation == "improvement":
            record["operation"]["direction"] = "lower"
        for field in ("candidate", "reported_baseline", "reported_difference"):
            record[field]["selector"]["row_key"]["差异类型"] = operation
        if explicit:
            record["reported_difference"]["selector"]["unit"] = {"kind": "column", "column": "差异单位"}
        def rows(sheets):
            sheets["多模型检验"][1][7] = "m"
            sheets["同模型多算法检验"][1][9] = "m"
            sheets["同模型多算法检验"][1][12] = operation
            sheets["同模型多算法检验"][0].append("差异单位")
            sheets["同模型多算法检验"][1].append(unit)
            if mutation:
                mutation(sheets)
        _, analysis = captured_books(algorithm_prediction=8 if operation == "improvement" else 24,
                                    algorithm_difference=reported_value, mutation=rows)
        primary = workbook_bytes({"核心指标": [["指标", "实例或场景", "单位", "数值"],
                                             ["prediction", "holdout", "m", 16]]})
        report = COMPARISON.inspect_plan(entry, question="Q1", specs=scope)
        self.assertEqual(report["issues"], [])
        return COMPARISON.inspect_evidence(report["plan"], primary_bytes=primary, analysis_bytes=analysis,
                                          dispositions=entry["analysis_evidence_dispositions"], comparison_rules=self.rules)

    def test_two_types_cover_required_items_and_recompute_real_literal_sources(self):
        result = self.evidence()
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["covered_ids"], ["CMP-Q1-01", "CMP-Q1-02"])
        self.assertEqual(result["computed_metrics"][0]["value"], "-5")

    def test_registered_required_model_cannot_be_replaced_by_sensitivity(self):
        result = self.evidence(mutation=lambda sheets: sheets.pop("多模型检验"))
        self.assertIn("CMP-Q1-01", result["uncovered_required_ids"])
        self.assertTrue(result["issues"])

    def test_unexecuted_exploratory_does_not_require_result_but_reported_rows_do(self):
        optional = self.entry["analysis_comparison"]["checks"][1]
        optional["requirement"] = "exploratory"
        del optional["disposition_ref"]
        result = self.evidence(mutation=lambda sheets: sheets.pop("同模型多算法检验"))
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["covered_ids"], ["CMP-Q1-01"])
        self.assertEqual(result["uncovered_required_ids"], [])
        self.assertTrue(self.evidence()["issues"])
        self.assertTrue(self.evidence(mutation=lambda sheets: sheets["同模型多算法检验"][1].__setitem__(0, "CMP-Q1-99"))["issues"])
        optional["disposition_ref"] = "E2"
        self.assertTrue(self.evidence(mutation=lambda sheets: sheets.pop("同模型多算法检验"))["issues"])

    def test_partial_exploratory_evidence_cannot_be_silently_skipped(self):
        optional = self.entry["analysis_comparison"]["checks"][1]
        optional["requirement"] = "exploratory"
        second = deepcopy(optional["evidence_refs"][0])
        second["id"] = "REC3"
        for field in ("candidate", "reported_baseline", "reported_difference"):
            second[field]["selector"]["row_key"]["记录键"] = "REC3"
        optional["evidence_refs"].append(second)
        self.assertTrue(self.evidence()["issues"])
        optional["evidence_refs"] = []
        self.assertTrue(self.evidence()["issues"])

    def test_wrong_difference_and_wrong_table_decision_are_technical_failures(self):
        self.assertTrue(self.evidence(algorithm_difference=1)["issues"])
        self.assertTrue(self.evidence(algorithm_decision=False)["issues"])

    def test_reported_values_cannot_rebind_the_header_to_another_physical_row(self):
        for field in ("reported_baseline", "reported_difference"):
            entry = deepcopy(self.entry)
            entry["analysis_comparison"]["checks"][0]["evidence_refs"][0][field]["selector"]["header_row"] = 2
            report = COMPARISON.inspect_plan(entry, question="Q1", specs=self.scope)
            self.assertTrue(any("same exact comparison row" in issue for issue in report["issues"]))

    def test_optional_columns_cannot_impersonate_registered_comparison_identities(self):
        for field, axis, column in (("candidate", "metric", "模型族"),
                                    ("candidate", "model", "模型族"),
                                    ("candidate", "scenario", "数据划分"),
                                    ("reported_baseline", "metric", "模型族"),
                                    ("reported_difference", "metric", "模型族"),
                                    ("reported_baseline", "model", "对照模型ID")):
            entry = deepcopy(self.entry)
            entry["analysis_comparison"]["checks"][0]["evidence_refs"][0][field]["selector"]["identity_columns"][axis] = column
            self.assertTrue(COMPARISON.inspect_plan(entry, question="Q1", specs=self.scope)["issues"])
        # This used to verify a displayed wrong metric through an optional text
        # column containing the approved metric, although the numbers agreed.
        check = self.entry["analysis_comparison"]["checks"][0]
        for field in ("candidate", "reported_baseline", "reported_difference"):
            check["evidence_refs"][0][field]["selector"]["identity_columns"]["metric"] = "模型族"
        primary, analysis = captured_books(mutation=lambda sheets: (
            sheets["多模型检验"][0].append("模型族"), sheets["多模型检验"][1].append("prediction"),
            sheets["多模型检验"][1].__setitem__(6, "different_actual_metric")))
        plan = deepcopy(self.entry["analysis_comparison"]) | {"scope": self.scope}
        result = COMPARISON.inspect_evidence(plan, primary_bytes=primary, analysis_bytes=analysis,
                                            dispositions=self.entry["analysis_evidence_dispositions"], comparison_rules=self.rules)
        self.assertTrue(result["issues"])

    def test_declared_reported_baseline_axis_binds_the_primary_object_column(self):
        for check in self.entry["analysis_comparison"]["checks"]:
            axis, column = (("model", "主模型ID") if check["kind"] == "model_comparison" else ("algorithm", "基准算法ID"))
            check["evidence_refs"][0]["reported_baseline"]["selector"]["identity_columns"][axis] = column
        self.assertEqual(self.evidence()["issues"], [])

    def test_dimensioned_relative_change_and_improvement_use_explicit_output_units(self):
        for operation in ("relative_change", "improvement"):
            result = self.dimensioned_relative_evidence(operation)
            self.assertEqual(result["issues"], [])
            self.assertIn("CMP-Q1-02", result["covered_ids"])
            percent = self.dimensioned_relative_evidence(operation, unit="%", reported_value=50)
            self.assertEqual(percent["issues"], [])

    def test_difference_unit_never_infers_missing_conflicting_or_formula_metadata(self):
        for unit in (None, "", "s", "百分点", "=\"ratio\""):
            self.assertTrue(self.dimensioned_relative_evidence(unit=unit)["issues"])
        self.assertTrue(self.dimensioned_relative_evidence(mutation=lambda sheets: (
            sheets["同模型多算法检验"][0].pop(), sheets["同模型多算法检验"][1].pop()))["issues"])
        self.assertTrue(self.dimensioned_relative_evidence(explicit=False)["issues"])
        for value in ("0.5", "=0.5", True):
            self.assertTrue(self.dimensioned_relative_evidence(mutation=lambda sheets, bad=value:
                sheets["同模型多算法检验"][1].__setitem__(13, bad))["issues"])
        for field in ("baseline", "candidate", "reported_baseline"):
            entry = deepcopy(self.entry)
            entry["analysis_comparison"]["checks"][1]["evidence_refs"][0][field]["selector"]["unit"] = {"kind": "column", "column": "差异单位"}
            self.assertTrue(COMPARISON.inspect_plan(entry, question="Q1", specs=self.scope)["issues"])

    def test_nonempty_difference_unit_cannot_be_ignored_by_an_ordinary_selector(self):
        self.assertTrue(self.evidence(mutation=lambda sheets: (
            sheets["同模型多算法检验"][0].append("差异单位"), sheets["同模型多算法检验"][1].append("ratio")))["issues"])
        self.assertEqual(self.evidence(mutation=lambda sheets: (
            sheets["同模型多算法检验"][0].append("差异单位"), sheets["同模型多算法检验"][1].append(None)))["issues"], [])

    def test_difference_unit_origin_is_captured_and_merged_unit_is_refused(self):
        raw = workbook_bytes({"同模型多算法检验": [["指标", "单位", "差异", "差异单位"],
                                                  ["prediction", "m", .5, "ratio"]]})
        source = {"source": "analysis", "selector": {
            "sheet": "同模型多算法检验", "header_row": 1, "row_key": {"指标": "prediction"},
            "expected_cardinality": 1, "value_type": "scalar", "value_column": "差异",
            "identity_columns": {"metric": "指标"}, "unit": {"kind": "column", "column": "差异单位"}}}
        value, location = COMPARISON._select_reported_difference(source, COMPARISON.Workbook(raw, self.rules))
        self.assertEqual(value.unit, "ratio")
        self.assertEqual(location["row"], 2)
        self.assertTrue(any(origin.endswith("|D2") for origin in value.origins))
        book = openpyxl.load_workbook(io.BytesIO(raw))
        try:
            book["同模型多算法检验"].merge_cells("D2:E2")
            buffer = io.BytesIO()
            book.save(buffer)
        finally:
            book.close()
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON._select_reported_difference(source, COMPARISON.Workbook(buffer.getvalue(), self.rules))

    def test_negative_result_with_modify_can_complete_without_core_redo(self):
        self.entry["analysis_evidence_dispositions"][1].update({"disposition": "modify", "impact_scope": "auxiliary_wording",
                                                              "required_action": "narrow algorithm claim"})
        result = self.evidence(algorithm_prediction=17, algorithm_difference=1, algorithm_decision=False)
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["disposition_impacts"][0]["impact_scope"], "auxiliary_wording")
        self.assertIsNone(result["disposition_impacts"][0]["return_stage"])

    def test_core_rejection_emits_typed_upstream_impact(self):
        self.entry["analysis_evidence_dispositions"][1].update({"disposition": "reject", "impact_scope": "core_answer",
                                                              "return_stage": "solve_validate", "required_action": "resolve current core answer"})
        result = self.evidence(algorithm_prediction=17, algorithm_difference=1, algorithm_decision=False)
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["disposition_impacts"][0]["event"], "core_answer_rejected")

    def test_resolved_negative_auxiliary_action_preserves_evidence_qualification(self):
        disposition = self.entry["analysis_evidence_dispositions"][1]
        for action in ("modify", "reject"):
            disposition.update({"disposition": action, "impact_scope": "auxiliary_wording", "status": "resolved"})
            self.assertEqual(self.evidence(algorithm_prediction=17, algorithm_difference=1, algorithm_decision=False)["issues"], [])
        disposition.update({"disposition": "reject", "impact_scope": "core_answer", "return_stage": "solve_validate"})
        self.assertTrue(self.evidence(algorithm_prediction=17, algorithm_difference=1, algorithm_decision=False)["issues"])
        disposition.clear()
        disposition.update({"id": "E2", "target_claim": "C-algorithm", "disposition": "support", "status": "resolved",
                            "method_or_source": "literal records", "key_finding": "agreement", "required_action": "keep scope"})
        self.assertTrue(self.evidence()["issues"])

    def test_missing_disposition_binding_cannot_fall_back_to_same_claim_text(self):
        del self.entry["analysis_comparison"]["checks"][1]["disposition_ref"]
        result = self.evidence()
        self.assertIn("CMP-Q1-02", result["uncovered_required_ids"])

    def test_same_algorithm_renaming_seed_or_stop_only_does_not_count(self):
        original = self.scope["algorithms"][0]["definition"]
        self.scope["algorithms"][1]["definition"] = deepcopy(original) | {"family": "different label", "stop_rule": "larger tolerance"}
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.scope_sha256(self.scope)

    def test_same_model_different_solver_is_not_a_second_mathematical_model(self):
        self.scope["models"][1]["identity"] = deepcopy(self.scope["models"][0]["identity"])
        self.scope["models"][1]["identity"]["algorithm_semantics"]["solver"] = "SVD"
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.scope_sha256(self.scope)

    def test_different_projection_partitions_cannot_manufacture_a_second_model(self):
        self.scope["models"][1]["identity"] = deepcopy(self.scope["models"][0]["identity"])
        self.scope["models"][1]["identity"]["algorithm_semantics"]["solver"] = "SVD"
        self.scope["models"][1]["pure_algorithm_keys"] = []
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.scope_sha256(self.scope)

    def test_unknown_or_mathematical_algorithm_fields_cannot_be_projected_away(self):
        identity = model_identity()
        identity["algorithm_semantics"]["domain_reduction"] = "only t=0"
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.mathematical_projection(identity, ["domain_reduction"])
        before = COMPARISON.mathematical_projection(identity, ["solver"])
        identity["algorithm_semantics"]["domain_reduction"] = "all declared times"
        after = COMPARISON.mathematical_projection(identity, ["solver"])
        self.assertNotEqual(before, after)
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.mathematical_projection(identity, ["stop_rule"])

    def test_algorithm_cannot_use_a_different_model_or_wrong_table_model_id(self):
        self.scope["algorithms"][1]["model_ref"] = "MODEL-Q1-02"
        with self.assertRaises(COMPARISON.ComparisonError):
            COMPARISON.scope_sha256(self.scope)
        self.scope = comparison_scope()
        result = self.evidence(mutation=lambda sheets: sheets["同模型多算法检验"][1].__setitem__(2, "MODEL-Q1-02"))
        self.assertTrue(result["issues"])

    def test_scope_changes_require_new_two_pass_review_and_explicit_user_binding(self):
        for field, value in (("statement", ""), ("approved_scope_sha256", "f" * 64), ("status", "pending")):
            entry = deepcopy(self.entry)
            entry["analysis_comparison"]["approval_binding"][field] = value
            self.assertTrue(COMPARISON.inspect_plan(entry, question="Q1", specs=self.scope)["issues"])
        for field, value in (("review_id", "synthetic-review-1"), ("blocking_items", ["unclosed mathematical defect"]), ("scope_sha256", "f" * 64)):
            entry = deepcopy(self.entry)
            entry["analysis_comparison"]["approval_binding"]["reviews"][1][field] = value
            self.assertTrue(COMPARISON.inspect_plan(entry, question="Q1", specs=self.scope)["issues"])

    def test_frozen_digest_binds_criteria_and_object_coverage_but_not_result_actions(self):
        report = self.plan()
        digest = report["plan_sha256"]
        changed = deepcopy(self.entry["analysis_comparison"])
        changed["checks"][0]["disposition_ref"] = "E99"
        self.assertEqual(COMPARISON.plan_sha256(changed), digest)
        changed["checks"][0]["criterion"]["threshold"] = "2"
        self.assertNotEqual(COMPARISON.plan_sha256(changed), digest)
        config = {"stage": "analysis", COMPARISON.PROTOCOL_FIELD: "1.0.0", COMPARISON.PLAN_FIELD: digest}
        self.assertEqual(COMPARISON.inspect_plan(self.entry, question="Q1", specs=self.scope, config=config)["issues"], [])
        config.pop(COMPARISON.PROTOCOL_FIELD)
        self.assertTrue(COMPARISON.inspect_plan(self.entry, question="Q1", specs=self.scope, config=config)["issues"])

    def test_new_scope_never_changes_primary_sib_identity_or_lock(self):
        framework = comparison_framework(self.scope)
        before = SEMANTIC.inspect_question_semantics(SEMANTIC.question_sections(framework)["Q1"], "Q1")
        changed_scope = deepcopy(self.scope)
        changed_scope["questions"][0]["evaluation"]["allowed_changes"].append("exploratory mechanism")
        after = SEMANTIC.inspect_question_semantics(SEMANTIC.question_sections(comparison_framework(changed_scope))["Q1"], "Q1")
        self.assertEqual(before, after)
        unchanged = deepcopy(self.entry)
        COMPARISON.inspect_plan(self.entry, question="Q1", specs=changed_scope)
        self.assertEqual(self.entry, unchanged)

    def test_framework_rejects_alias_duplicate_partial_and_primary_scope_markers(self):
        frame = comparison_framework(self.scope)
        self.assertEqual(COMPARISON.extract_specs(frame, "Q1"), self.scope)
        for malformed in (
            frame.replace("<!-- HSK_ANALYSIS_COMPARISON_END Q1 -->", ""),
            frame.replace("#### 结果摘要", "#### different heading"),
            frame.replace("protocol_version: 1.0.0", "protocol_version: 1.0.0\nprotocol_version: 1.0.0"),
            frame.replace("protocol_version: 1.0.0", "protocol_version: &v 1.0.0\nunused: *v"),
        ):
            with self.assertRaises(COMPARISON.ComparisonError):
                COMPARISON.extract_specs(malformed, "Q1")

    def test_required_not_required_conflict_and_unknown_protocol_fail_closed(self):
        self.entry["result_analysis_status"] = "not_required"
        self.assertTrue(COMPARISON.inspect_plan(self.entry, question="Q1", specs=self.scope)["issues"])
        self.entry["result_analysis_status"] = "pending"
        self.entry["analysis_comparison"]["protocol_version"] = "9.0.0"
        self.assertTrue(COMPARISON.inspect_plan(self.entry, question="Q1", specs=self.scope)["issues"])
        self.assertFalse(COMPARISON.enabled({"analysis_methods": ["参数敏感性"]}))

    def test_formulas_numeric_strings_wrong_axes_and_wrong_criteria_are_not_verified(self):
        mutations = [
            lambda sheets: sheets["同模型多算法检验"][1].__setitem__(11, "16"),
            lambda sheets: sheets["同模型多算法检验"][1].__setitem__(11, "=16"),
            lambda sheets: sheets["同模型多算法检验"][1].__setitem__(6, "training"),
            lambda sheets: sheets["同模型多算法检验"][1].__setitem__(14, "different-criterion"),
        ]
        for mutation in mutations:
            self.assertTrue(self.evidence(mutation=mutation)["issues"])

    def test_relative_change_zero_baseline_is_not_invented(self):
        check = self.entry["analysis_comparison"]["checks"][1]
        check["evidence_refs"][0]["operation"]["op"] = "relative_change"
        check["criterion"]["unit"] = "ratio"
        for source in ("candidate", "reported_baseline", "reported_difference"):
            check["evidence_refs"][0][source]["selector"]["row_key"]["差异类型"] = "relative_change"
        primary = workbook_bytes({"核心指标": [["指标", "实例或场景", "单位", "数值"], ["prediction", "holdout", "dimensionless", 0]]})
        _, analysis = captured_books(mutation=lambda sheets: (sheets["同模型多算法检验"][1].__setitem__(10, 0),
                                                             sheets["同模型多算法检验"][1].__setitem__(12, "relative_change")))
        result = COMPARISON.inspect_evidence(self.plan()["plan"], primary_bytes=primary, analysis_bytes=analysis,
                                            dispositions=self.entry["analysis_evidence_dispositions"], comparison_rules=self.rules)
        self.assertTrue(any("baseline" in issue for issue in result["issues"]))

    def test_optional_schema_is_strict_without_rewriting_existing_state_fields(self):
        schema = yaml.safe_load((ROOT / "core/project_state.schema.yaml").read_text(encoding="utf-8"))
        validator = Draft202012Validator({"$ref": "#/$defs/analysis_comparison", "$defs": schema["$defs"]})
        self.assertEqual(list(validator.iter_errors(self.entry["analysis_comparison"])), [])
        record = deepcopy(self.entry["analysis_comparison"])
        record["self_reported_pass"] = True
        self.assertTrue(list(validator.iter_errors(record)))

    def test_approval_wrapper_checks_actual_primary_sib_and_opt_in_scoped_receipt(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "模型论文框架.md").write_text(comparison_framework(self.scope), encoding="utf-8")
            state = {"subproblems": {"Q1": self.entry}}
            self.assertEqual(APPROVAL.validate_comparison_approval(root, state, self.entry, "Q1"), [])
            state["review_receipt_policy"] = {"synthetic": "only_for_mock"}
            with patch("review_receipt_consumption.evaluate_gate") as evaluate:
                evaluate.return_value = {"status": "blocked", "issues": ["main-only receipt cannot cover comparison scope"]}
                self.assertTrue(APPROVAL.validate_comparison_approval(root, state, self.entry, "Q1"))
                self.assertTrue(evaluate.call_args.kwargs["comparison_scope"])
            changed = comparison_framework(self.scope).replace("domain: R^3", "domain: R^4", 1)
            self.assertTrue(APPROVAL.validate_comparison_approval(root, {"subproblems": {"Q1": self.entry}}, self.entry,
                                                                "Q1", specs=changed))


if __name__ == "__main__":
    unittest.main()
