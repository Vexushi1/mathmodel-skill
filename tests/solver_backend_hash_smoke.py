"""Fresh-process native regression for equivalent SHA casing in solver templates.

Only creates repository maintenance fixtures in a new output directory.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import subprocess
import sys

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from tests import test_solver_backend_end_to_end as fixture
import stage_code
import validate_code_delivery as delivery
import validate_user_execution as acceptance


def prepare_stage(root: Path, stage: str) -> Path:
    if stage == "primary":
        fixture.prepare(root)
    else:
        fixture.prepare_analysis(root)
    helper = root / fixture.QUESTION / "hsk_hash_helper.m"
    helper.write_text("function y = hsk_hash_helper(x)\ny = x;\nend\n", encoding="utf-8")
    updates = {
        "data_sha256": fixture.reference_digest(root, ["input.json"]).upper(),
        "code_dependencies": [{"path": helper.relative_to(root).as_posix(), "sha256": fixture.file_hash(helper).upper()}],
    }
    if stage == "analysis":
        updates["primary_workbook_sha256"] = fixture.file_hash(root / fixture.QUESTION / "问题一求解结果.xlsx").upper()
    code = fixture.instantiate(root, stage, **updates)
    text = code.read_text(encoding="utf-8")
    marker = "input = read_model_input(root, RUN_CONFIG);"
    assert marker in text
    code.write_text(text.replace(marker, marker + "\ninput.coefficient = hsk_hash_helper(input.coefficient);", 1), encoding="utf-8")
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    fingerprint = stage_code.stage_code_fingerprint(root, code, updates["code_dependencies"])
    entry["code" if stage == "primary" else "result_analysis_code"] = code.relative_to(root).as_posix()
    entry["primary_code_sha256" if stage == "primary" else "analysis_code_sha256"] = fingerprint["entry_sha256"]
    entry["solver_execution"][stage]["bundle_sha256"] = fingerprint["bundle_sha256"]
    fixture.save_state(root, state)
    return code


def run(root: Path, matlab: str) -> None:
    if root.exists() and any(root.iterdir()):
        raise ValueError("Use a fresh synthetic output directory")
    reports = []
    primary = root / fixture.QUESTION / "问题一求解结果.xlsx"
    for stage in ("primary", "analysis"):
        code = prepare_stage(root, stage)
        diagnostics = {}
        issues, config = delivery.validate_script(root, code, stage, matlab_command=matlab,
                                                  require_native=True, native_report=diagnostics)
        assert not issues, issues
        assert diagnostics["status"] == "passed"
        delivery.update_state(root, config, code)
        before_primary = primary.read_bytes() if stage == "analysis" else None
        release_path = root / (stage + "_release.txt")
        quote = lambda path: str(path).replace("'", "''")
        expression = (f"addpath('{quote(code.parent)}'); {code.stem}(); "
                      f"fid=fopen('{quote(release_path)}','w'); assert(fid>=0); "
                      "fprintf(fid,'%s',version('-release')); fclose(fid);")
        result = subprocess.run([matlab, "-batch", expression], capture_output=True, encoding="utf-8",
                                errors="replace", timeout=300, env={**os.environ, "PYTHONUTF8": "1"})
        (root / f"{stage}_native.log").write_text(result.stdout + result.stderr, encoding="utf-8")
        assert result.returncode == 0, f"Native {stage} failed; see {root / (stage + '_native.log')}"
        workbook = root / fixture.QUESTION / ("问题一求解结果.xlsx" if stage == "primary" else "问题一结果深化分析.xlsx")
        state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
        issues = acceptance.validate_one(root, workbook, state, True)
        assert not issues, issues
        assert state["subproblems"]["Q1"][stage + "_execution_status"] == "accepted"
        book = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
        try:
            if stage == "primary":
                values = dict(book["核心指标"].iter_rows(min_row=2, values_only=True))
                assert values["解"] == 3
            else:
                assert [row[3] for row in book["参数敏感性"].iter_rows(min_row=2, values_only=True)] == [4, 3, 2]
                assert primary.read_bytes() == before_primary
        finally:
            book.close()
        fixture.save_state(root, state)
        reports.append({"stage": stage, "actual_matlab_execution": True,
                        "matlab_release": release_path.read_text(encoding="utf-8"),
                        "native_code_analyzer": diagnostics["status"],
                        "uppercase_data_and_helper_sha": True,
                        "uppercase_primary_sha": stage == "analysis", "status": "passed"})
    (root / "hash_casing_report.json").write_text(json.dumps(reports, indent=2), encoding="utf-8")
    print(json.dumps(reports))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--matlab-command", required=True)
    args = parser.parse_args()
    run(args.project.resolve(), args.matlab_command)
