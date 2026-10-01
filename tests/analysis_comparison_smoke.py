"""Real, isolated QR/SVD comparison smoke; never runs a user's modeling project."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
import analysis_comparison as comparison
import semantic_identity
import validate_code_delivery as delivery
import validate_user_execution as receipts
from templates.code.hsk_pipeline import result_io
from tests.test_solver_backend_end_to_end import file_hash, reference_digest, save_state

QUESTION = "问题一求解"
PAYLOAD = {"train_t": [0, 1, 2, 3], "train_y": [0, 1, 4, 9],
           "evaluation_t": [0, 1, 2, 3, 4], "scenarios": ["train-0", "train-1", "train-2", "train-3", "holdout"],
           "model_difference_threshold": 1.0}


def model_identity(degree=2):
    expression = "c0+c1*t+c2*t^2" if degree == 2 else "c0+c1*t"
    return {"schema_version": "1.0.0", "question": "Q1", "research_object": "Synthetic polynomial prediction",
            "data_scope": [{"id": "D1", "source": "input.json", "role": "complete training/evaluation declaration"}],
            "variables": [{"id": "V1", "symbol": "c", "role": "decision", "domain": f"real^{degree+1}"}],
            "parameters": [{"id": "P1", "symbol": "t,y", "unit": "dimensionless"}],
            "assumptions": [{"id": "A1", "statement": "Only declared training points fit the coefficients"}],
            "objective": {"sense": "minimize", "expression": f"sum_train ({expression}-y)^2"},
            "constraints": [], "preprocessing_decision": "not_needed",
            "algorithm_semantics": {"model_family": "polynomial_least_squares", "solver": "QR"}, "dependencies": []}


def framework(identity, specs=None):
    text = ("# 模型论文框架\n## 各问模型与结果\n### Q1：第一问\n#### 当前模型口径\n"
            "<!-- HSK_SEMANTIC_IDENTITY_BEGIN Q1 -->\n```yaml\n"
            + yaml.safe_dump(identity, allow_unicode=True, sort_keys=False)
            + "```\n<!-- HSK_SEMANTIC_IDENTITY_END Q1 -->\n#### 结果摘要\n维护合成案例。\n")
    if specs is not None:
        text += ("#### 03B比较计划与证据\n<!-- HSK_ANALYSIS_COMPARISON_BEGIN Q1 -->\n```yaml\n"
                 + yaml.safe_dump(specs, allow_unicode=True, sort_keys=False)
                 + "```\n<!-- HSK_ANALYSIS_COMPARISON_END Q1 -->\n")
    return text


def prepare_project(root, backend):
    root.mkdir(parents=True, exist_ok=True)
    if (root / "state/project_state.yaml").exists():
        raise ValueError("Use a fresh isolated repository-fixture directory")
    (root / QUESTION).mkdir()
    (root / "input.json").write_text(json.dumps(PAYLOAD), encoding="utf-8")
    identity = model_identity()
    digest = semantic_identity.semantic_identity_hash(identity)
    text = framework(identity)
    (root / "模型论文框架.md").write_text(text, encoding="utf-8")
    state = {"project": {"competition": "test", "problem": "synthetic_comparison", "current_phase": "solve_validate"},
             "execution": {"solver_backend": backend, "solver_backend_selection_reason": "Isolated whole-project comparison fixture"},
             "preprocessing": {"decision": "not_needed", "status": "not_applicable", "quality_status": "not_applicable"},
             "paper_framework": {"path": "模型论文框架.md", "sync_status": "current", "paper_fragments": []},
             "subproblems": {"Q1": {"status": "designed", "selected_model": "quadratic least squares", "capabilities": {"requires_equilibrium_residual": True},
                 "data_hash": reference_digest(root, ["input.json"]), "result_quality_status": "pending", "result_analysis_status": "pending",
                 "semantic_identity_schema_version": "1.0.0", "semantic_revision": 1, "validated_semantic_revision": 1, "approved_semantic_revision": 1,
                 "semantic_identity_hash": digest, "validated_semantic_identity_hash": digest, "approved_semantic_identity_hash": digest,
                 "semantic_text_hash": semantic_identity.sha256_text(semantic_identity.semantic_scope(semantic_identity.question_sections(text)["Q1"])),
                 "model_challenge_status": "passed", "human_model_approval_status": "approved",
                 "problem_contract_status": "frozen", "semantic_closure_status": "passed", "complexity_sanity_status": "passed",
                 "depends_on": [], "stale_layers": [], "artifacts_stale": False, "artifact_hashes": {}, "validated_artifact_hashes": {}}}}
    save_state(root, state)
    return state


def selector(sheet, keys, column, identities):
    return {"sheet": sheet, "header_row": 1, "row_key": keys, "expected_cardinality": 1,
            "value_type": "scalar", "value_column": column, "identity_columns": identities,
            "unit": {"kind": "column", "column": "单位"}}


def comparison_scope(identity, backend):
    anchor = f"{QUESTION}/" + ("comparison_numerics.py::fit_polynomial" if backend == "python" else "hsk_comparison_fit.m::hsk_comparison_fit")
    questions = []
    for identifier, kind, baseline, candidate, question, claim, evaluation in (
            ("CMP-Q1-01", "model_comparison", "MODEL-Q1-01", "MODEL-Q1-02", "Does a quadratic term affect the declared holdout prediction?", "case_quadratic_term_changes_prediction", "EVAL-Q1-01"),
            ("CMP-Q1-02", "algorithm_comparison", "ALGO-Q1-01", "ALGO-Q1-02", "Do QR and SVD agree for the same quadratic least-squares model?", "case_qr_svd_prediction_agreement", "EVAL-Q1-02")):
        questions.append({"id": identifier, "kind": kind, "baseline_ref": baseline, "candidate_refs": [candidate], "question": question,
                          "target_claim": claim, "evaluation": {"id": evaluation, "metric": "prediction", "unit": "dimensionless", "direction": "lower",
                            "fixed_axes": ["scenario"], "allowed_changes": ["polynomial_degree"] if kind == "model_comparison" else ["decomposition_method"]}})
    return {"protocol_version": "1.0.0", "question": "Q1", "baseline_semantic_identity_hash": semantic_identity.semantic_identity_hash(identity),
            "models": [{"id": "MODEL-Q1-01", "identity": identity, "pure_algorithm_keys": ["solver"]},
                       {"id": "MODEL-Q1-02", "identity": model_identity(1), "pure_algorithm_keys": ["solver"]}],
            "algorithms": [{"id": "ALGO-Q1-01", "model_ref": "MODEL-Q1-01", "definition": {"family": "QR", "update_rule": "economy QR followed by triangular solve", "stop_rule": "direct finite decomposition"}, "implementation_anchor": anchor},
                           {"id": "ALGO-Q1-02", "model_ref": "MODEL-Q1-01", "definition": {"family": "SVD", "update_rule": "economy SVD and singular-coordinate solve", "stop_rule": "direct finite decomposition"}, "implementation_anchor": anchor}],
            "questions": questions}


def activate_comparison(root, state, backend):
    entry = state["subproblems"]["Q1"]
    specs = comparison_scope(model_identity(), backend)
    scope_hash = comparison.scope_sha256(specs)
    reviews = [{"review_id": f"synthetic-scope-{i}", "role": role, "scope_sha256": scope_hash, "verdict": "passed", "method": "separate_pass",
                "conclusion": "Synthetic fixture declaration only; mathematical differences and shared data are explicit.", "blocking_items": [], "unresolved_items": []}
               for i, role in enumerate(("positive_fitness_review", "adversarial_model_challenge"), 1)]
    registry = {"protocol_version": "1.0.0", "scope_ref": "Q1", "scope_sha256": scope_hash,
                "baseline_semantic_identity_hash": entry["semantic_identity_hash"], "approval_binding": {"status": "approved", "owner": "user",
                    "statement": "Repository synthetic approval fixture, not a real user project's approval.", "source_ref": "fixture:analysis-comparison",
                    "approved_scope_sha256": scope_hash, "reviews": reviews}, "checks": []}
    for scoped in specs["questions"]:
        model = scoped["kind"] == "model_comparison"
        sheet = "多模型检验" if model else "同模型多算法检验"
        key = "model-4" if model else "algorithm-4"
        keys = {"检验ID": scoped["id"], "记录键": key, "评价协议ID": scoped["evaluation"]["id"]}
        criterion_id = f"CRIT-{scoped['id']}"
        keys.update({"差异类型": "difference", "判据ID": criterion_id})
        keys.update({"主模型ID": "MODEL-Q1-01", "对照模型ID": "MODEL-Q1-02"} if model else {
            "模型ID": "MODEL-Q1-01", "基准算法ID": "ALGO-Q1-01", "对照算法ID": "ALGO-Q1-02"})
        identities = {"metric": "指标", "scenario": "实例或场景", "model" if model else "algorithm": "对照模型ID" if model else "对照算法ID"}
        evidence = {"id": key, "candidate_ref": scoped["candidate_refs"][0],
                    "baseline": {"source": "primary", "selector": selector("状态明细", {"记录键": "prediction-4"}, "数值", {"metric": "指标", "scenario": "实例或场景"})},
                    "candidate": {"source": "analysis", "selector": selector(sheet, keys, "对照模型数值" if model else "对照数值", identities)},
                    "reported_baseline": {"source": "analysis", "selector": selector(sheet, keys, "主模型数值" if model else "基准数值",
                        identities | {"model" if model else "algorithm": "主模型ID" if model else "基准算法ID"})},
                    "reported_difference": {"source": "analysis", "selector": selector(sheet, keys, "差异", identities)},
                    "operation": {"op": "difference", "comparison_axis": "model" if model else "algorithm"}}
        evidence_refs = [evidence]
        if not model:
            evidence_refs = []
            for index in range(len(PAYLOAD["evaluation_t"])):
                record = deepcopy(evidence)
                record["id"] = f"algorithm-{index}"
                record["baseline"]["selector"]["row_key"]["记录键"] = f"prediction-{index}"
                for reference in ("candidate", "reported_baseline", "reported_difference"):
                    record[reference]["selector"]["row_key"]["记录键"] = record["id"]
                evidence_refs.append(record)
        registry["checks"].append({key: deepcopy(scoped[key]) for key in ("id", "kind", "baseline_ref", "candidate_refs", "question", "target_claim")} | {
            "requirement": "required", "requirement_source": "fixture:declared-two-comparison-contract", "criterion": {
                "id": criterion_id,
                "relation": "abs_ge" if model else "abs_le", "threshold": "1.0" if model else "0.0000000001", "unit": "dimensionless",
                "source": "Predeclared synthetic case criterion", "arithmetic_tolerance": "0.000000000001"}, "evidence_refs": evidence_refs})
    entry.update(analysis_methods=["多模型检验", "同模型多算法检验"], result_analysis_requirement_reason="Check declared model and algorithm dependency in an isolated fixture",
                 result_analysis_status="pending", analysis_comparison=registry)
    (root / "模型论文框架.md").write_text(framework(model_identity(), specs), encoding="utf-8")
    save_state(root, state)
    report = comparison.inspect_plan(entry, question="Q1", specs=specs)
    assert report["issues"] == [], report
    return report["plan_sha256"]


def instantiate(root, backend, stage, state, plan_hash=None):
    python = backend == "python"
    name = ("问题一求解.py" if stage == "primary" else "问题一结果深化分析.py") if python else ("q1_solver.m" if stage == "primary" else "q1_analysis.m")
    helper = "comparison_numerics.py" if python else "hsk_comparison_fit.m"
    folder = "starter" if python else "matlab"
    shutil.copyfile(ROOT / "templates/code" / folder / helper, root / QUESTION / helper)
    relative_helper = f"{QUESTION}/{helper}"
    config = {"stage": stage, "problem_name": "问题一", "solver_backend": backend, "data_paths": ["input.json"],
              "data_sha256": state["subproblems"]["Q1"]["data_hash"], "solver": "explicit_qr" if stage == "primary" else "explicit_qr_svd_comparison",
              "random_seed": 2026, "tolerance": 1e-10, "iteration_or_time_limit": "direct", "run_receipt_protocol_version": "1.1.0",
              "expected_workbook": f"{QUESTION}/问题一{'求解结果' if stage == 'primary' else '结果深化分析'}.xlsx",
              "code_dependencies": [{"path": relative_helper, "sha256": file_hash(root / relative_helper)}]}
    if stage == "primary":
        config["primary_quality_protocol_version"] = "1.0.0"
    else:
        config.update(primary_workbook=f"{QUESTION}/问题一求解结果.xlsx", primary_workbook_sha256=file_hash(root / QUESTION / "问题一求解结果.xlsx"),
                      analysis_comparison_protocol_version="1.0.0", analysis_comparison_plan_sha256=plan_hash)
    source_name = "analysis_comparison.py" if python else "analysis_comparison_example.m"
    source = (ROOT / "templates/code" / folder / source_name).read_text(encoding="utf-8")
    if python:
        source = source.replace("RUN_CONFIG = {}", f"RUN_CONFIG = {config!r}", 1)
    else:
        source = source.replace("analysis_comparison_example()", Path(name).stem + "()", 1)
        literal = json.dumps(config, ensure_ascii=False, separators=(",", ":")).replace("'", "''")
        source = re.sub(r"RUN_CONFIG = jsondecode\('[^\n]*'\);", lambda match: f"RUN_CONFIG = jsondecode('{literal}');", source, count=1)
    code = root / QUESTION / name
    code.write_text(source, encoding="utf-8", newline="\n")
    return code, config


def execute_stage(root, backend, code, stage, matlab_command):
    if backend == "python":
        command = [sys.executable, "-B", str(code)]
    else:
        assert matlab_command, "An actual fresh MATLAB command is required"
        directory = str(code.parent).replace("'", "''")
        command = [matlab_command, "-batch", f"cd('{directory}'); {code.stem}(); disp('HSK comparison {stage} completed');"]
    executed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONUTF8": "1"})
    log = executed.stdout + executed.stderr
    (root / f"{stage}_native.log").write_text(log, encoding="utf-8")
    assert executed.returncode == 0, log
    return executed.returncode


def run_smoke(root, backend, matlab_command=None):
    state = prepare_project(root, backend)
    outcomes, primary_bytes, plan_hash = [], None, None
    for stage in ("primary", "analysis"):
        if stage == "analysis":
            plan_hash = activate_comparison(root, state, backend)
        code, config = instantiate(root, backend, stage, state, plan_hash)
        native = {}
        errors, parsed = delivery.validate_script(root, code, stage, matlab_command=matlab_command,
                                                  require_native=backend == "matlab", native_report=native)
        assert not errors, errors
        delivery.update_state(root, parsed, code)
        state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
        exit_code = execute_stage(root, backend, code, stage, matlab_command)
        workbook = root / config["expected_workbook"]
        entry = state["subproblems"]["Q1"]
        if stage == "analysis":
            for i, check in enumerate(entry["analysis_comparison"]["checks"], 1):
                check["disposition_ref"] = f"E{i}"
            entry["analysis_evidence_dispositions"] = [{"id": f"E{i}", "status": "current", "method_or_source": check["kind"],
                "target_claim": check["target_claim"], "disposition": "support", "key_finding": "Actual decomposition results meet the predeclared case criterion",
                "required_action": "Use only the declared synthetic-case claim", "paper_or_figure_anchor": "fixture:case-only"}
                for i, check in enumerate(entry["analysis_comparison"]["checks"], 1)]
            save_state(root, state)
        result_io.validate_workbook_file(workbook, "solution" if stage == "primary" else "result_analysis",
            capabilities=entry["capabilities"] if stage == "primary" else {}, analysis_methods=entry.get("analysis_methods", []) if stage == "analysis" else [],
            comparison_plan=entry.get("analysis_comparison") if stage == "analysis" else None)
        errors = receipts.validate_one(root, workbook, state, True)
        assert not errors, errors
        execution_key = "primary_execution_status" if stage == "primary" else "analysis_execution_status"
        quality_key = "result_quality_status" if stage == "primary" else "result_analysis_status"
        assert entry[execution_key] == "accepted" and entry[quality_key] == "passed", entry
        metadata, errors = receipts.configuration_map(workbook)
        assert not errors and metadata["solver_backend"] == backend and metadata["run_receipt_version"] == "1.1.0"
        assert metadata["code_bundle_sha256"] == reference_digest(root, [code.relative_to(root).as_posix(), *[row["path"] for row in config["code_dependencies"]]])
        if stage == "primary":
            primary_bytes = workbook.read_bytes()
        else:
            assert (root / QUESTION / "问题一求解结果.xlsx").read_bytes() == primary_bytes
            assert metadata["analysis_comparison_plan_sha256"] == plan_hash
            assert entry["primary_execution_status"] == "accepted" and entry["result_quality_status"] == "passed"
            assert entry["semantic_identity_hash"] == entry["approved_semantic_identity_hash"] == entry["validated_semantic_identity_hash"]
        outcomes.append({"stage": stage, "native_exit": exit_code, "code_analyzer": native, "receipt_accepted": True})
        save_state(root, state)
    book = openpyxl.load_workbook(root / QUESTION / "问题一结果深化分析.xlsx", read_only=True, data_only=True)
    try:
        model_rows = list(book["多模型检验"].iter_rows(values_only=True))
        algorithm_rows = list(book["同模型多算法检验"].iter_rows(values_only=True))
        model = dict(zip(model_rows[0], model_rows[-1]))
        algorithm = dict(zip(algorithm_rows[0], algorithm_rows[-1]))
        assert len(model_rows) == 2 and len(algorithm_rows) == len(PAYLOAD["evaluation_t"]) + 1
        assert model["记录键"] == "model-4" and model["实例或场景"] == "holdout"
        assert abs(model["主模型数值"] - 16) < 1e-10 and abs(model["对照模型数值"] - 11) < 1e-10
        assert abs(model["差异"] + 5) < 1e-10 and abs(algorithm["差异"]) <= 1e-10
    finally:
        book.close()
    report = {"status": "passed", "backend": backend, "synthetic_only": True, "actual_python_execution": backend == "python",
              "actual_matlab_execution": backend == "matlab", "matlab_release": metadata.get("matlab_release") if backend == "matlab" else None,
              "primary_unchanged": True, "model_comparison_verified": True, "algorithm_comparison_verified": True,
              "linear_holdout_prediction": model["对照模型数值"], "quadratic_holdout_prediction": model["主模型数值"], "stages": outcomes}
    if backend == "matlab":
        assert report["matlab_release"] == "2024b", report
    helper = root / QUESTION / ("comparison_numerics.py" if backend == "python" else "hsk_comparison_fit.m")
    original_helper = helper.read_bytes()
    try:
        helper.write_bytes(original_helper + b"\n")
        rejected = receipts.validate_one(root, root / QUESTION / "问题一结果深化分析.xlsx", deepcopy(state), False)
        assert rejected, "A changed actually-used numerical helper must invalidate acceptance"
        assert (root / QUESTION / "问题一求解结果.xlsx").read_bytes() == primary_bytes
        report["helper_drift_rejected"] = True
    finally:
        helper.write_bytes(original_helper)
    (root / "comparison_smoke_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--backend", choices=("python", "matlab"), required=True)
    parser.add_argument("--matlab-command")
    args = parser.parse_args()
    print(json.dumps(run_smoke(args.project.resolve(), args.backend, args.matlab_command), ensure_ascii=False))
