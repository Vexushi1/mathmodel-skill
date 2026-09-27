"""B12 continuous acceptance on real synthetic A2/B1 primary/analysis receipts.

The fixture's model/figure approvals are explicit setup data. Compiler records
are labelled synthetic proof fixtures except in the named real_xelatex variant.
Neither path claims visual or human review. No qualification gate is mocked.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
import shutil
import subprocess
import unittest
import zipfile
import zlib

import openpyxl
import yaml

from tests import test_b2b4_figure_integration as figure_fixture
from tests.claim_fixture import answer_record
from tests.test_model_code_conformance import bytes_in, save
from tests.test_sync_project import capabilities
from artifact_fingerprint import combined_hash, framework_section_hash
import audit_latex_project as latex_audit
import claim_consumption as claims
import claim_evidence
import hsk_pack_submission as packager
import latex_delivery as delivery
import runtime_assurance
import sync_project as synchronizer
import validate_submission_package as package_validator
from project_transaction import LOCK_RELATIVE_PATH


class B2b5ContinuousAcceptanceTests(unittest.TestCase):
    save_projection = figure_fixture.FigureFormalChainIntegrationTests.save_projection
    set_caption = figure_fixture.FigureFormalChainIntegrationTests.set_caption

    @classmethod
    def setUpClass(cls):
        figure_fixture.FigureFormalChainIntegrationTests.setUpClass.__func__(cls)

    @classmethod
    def tearDownClass(cls):
        figure_fixture.FigureFormalChainIntegrationTests.tearDownClass.__func__(cls)

    def setUp(self):
        figure_fixture.FigureFormalChainIntegrationTests.setUp(self)
        self.state.update(
            requirements={'total': 0, 'completed': [], 'pending': []}, decisions={},
            variables={'locked': [], 'source': {}}, risks=[],
            next_gate={'module': 'solve_validate', 'condition': 'Synthetic B12 acceptance'},
        )
        self.state['project'].update(competition='CUMCM', problem='synthetic-B12',
                                     current_phase='solve_validate')
        self.state['artifacts'].update(code=[], results=[], papers=[])
        # The formal LaTeX scope also checks the existing project-level MATLAB
        # handoff. This read-only fixture is supplied, never executed here.
        (self.root / '数据预处理/data_process.m').write_text(
            "function data_process\nreadtable('数据预处理结果.xlsx');\nend\n", encoding='utf-8')
        framework = self.state['paper_framework']
        framework.update(path='模型论文框架.md', version='v0.8-project-memory', mode='compact',
                         sync_status='current', last_sync_scope='results', proposition_count=0,
                         proposition_status='not_assessed', propositions=[])
        entry = self.state['subproblems']['Q1']
        entry.update(classification={'objective': 'explanation', 'structures': []},
                     capabilities=capabilities(requires_equilibrium_residual=True),
                     framework_section='### Q1：第一问',
                     result_summary_anchor='#### 结果摘要', result_summary_status='current')

        # Keep the accepted structured model section intact; all writing records
        # follow a new section, so claim edits cannot change its identity.
        path = self.root / '模型论文框架.md'
        text = path.read_text(encoding='utf-8')
        text = text.replace('# 模型论文框架\n', '# 模型论文框架\n\n'
                            '> 本文件只保留当前有效版本。\n\n'
                            '## 当前有效口径\n\n## 各问模型与结果\n', 1)
        text = text.split('### Paper Fragment Dependency Map', 1)[0]
        text += ('\n## 待办与缺口\n\n'
                 '### Terminology Registry\n\n### Numeric Profile\n\n')
        path.write_text(text, encoding='utf-8')

        auxiliary = answer_record('analysis')
        source = auxiliary['sources'][0]
        source['id'] = 'auxiliary'
        claim = auxiliary['claims'][0]
        claim.update(id='aux_claim', text='The synthetic local stability estimate is 100.',
                     numeric_profile_id='N98')
        claim['evidence'][0]['ref'] = 'source:auxiliary'
        claim['assertion'].update(evidence_ref='source:auxiliary', value='100.0',
                                  display_location='abstract')
        framework['claim_evidence']['sources'].append(source)
        framework['claim_evidence']['claims'].append(claim)
        framework['numeric_profile'][1].update(abstract_decimals=1, body_decimals=2)
        framework['claim_consumption_policy']['required_consumptions'].append({
            'claim_id': 'aux_claim', 'fragment_kinds': ['abstract_claim', 'question_result_text'],
        })
        for identifier, kind, dependency, anchor, filename in (
            ('paper.aux.abstract', 'abstract_claim', 'claim:aux_claim', 'Auxiliary:', 'aux_abstract'),
            ('paper.aux.body', 'question_result_text', 'claim:aux_claim', 'Stability:', 'aux_body'),
            ('paper.aux.context', 'paper_section', 'paper.aux.body', 'Context:', 'aux_context'),
        ):
            self.fragments.append({'id': identifier, 'kind': kind, 'scope': 'Q1',
                                   'depends_on': [dependency], 'anchor': anchor,
                                   'source_file': f'final_latex/{filename}.tex', 'status': 'current'})
        (self.latex / 'aux_abstract.tex').write_text('Auxiliary: 100.0.\n', encoding='utf-8')
        (self.latex / 'aux_body.tex').write_text('Stability: 100.00 at the tested coefficient.\n',
                                               encoding='utf-8')
        (self.latex / 'aux_context.tex').write_text('Context: the local stability assessment.\n',
                                                  encoding='utf-8')
        main = self.latex / 'main.tex'
        main.write_text(main.read_text(encoding='utf-8').replace(
            '\\end{document}', '\\input{aux_abstract}\n\\input{aux_body}\n'
            '\\input{aux_context}\n\\end{document}').replace(
            '\\documentclass{article}', '\\documentclass{article}\n\\usepackage{graphicx}'),
            encoding='utf-8')
        # A minimal valid PNG makes the identical fixture usable by XeLaTeX.
        def chunk(tag, data):
            body = tag + data
            return len(data).to_bytes(4, 'big') + body + zlib.crc32(body).to_bytes(4, 'big')

        image = self.root / 'figures/q1_f01.png'
        image.write_bytes(b'\x89PNG\r\n\x1a\n'
                          + chunk(b'IHDR', (1).to_bytes(4, 'big') * 2 + bytes([8, 2, 0, 0, 0]))
                          + chunk(b'IDAT', zlib.compress(b'\x00\xff\x00\x00'))
                          + chunk(b'IEND', b''))
        entry['artifact_hashes']['figure_bundle'] = combined_hash([image], self.root)
        entry['validated_artifact_hashes']['figure_bundle'] = entry['artifact_hashes']['figure_bundle']
        self.save_projection()
        digest = framework_section_hash(path, entry['framework_section'])
        self.assertIsNotNone(digest)
        entry['artifact_hashes']['framework'] = digest
        entry['validated_artifact_hashes']['framework'] = digest
        save(self.root, self.state)
        baseline = synchronizer.synchronize(self.root, write=True)
        self.assertEqual(baseline['status'], 'passed', baseline['issues'])
        self.assertEqual(baseline['stale_questions'], [])
        self.reload()
        self.assertEqual(claims.formal_figure_gate(self.root)['status'], 'passed')
        self.assert_accepted()
        self.numerical_before = self.numerical_snapshot()

    def tearDown(self):
        figure_fixture.FigureFormalChainIntegrationTests.tearDown(self)

    def reload(self):
        self.state = yaml.safe_load((self.root / 'state/project_state.yaml').read_text(encoding='utf-8'))
        self.fragments = self.state['paper_framework']['paper_fragments']

    def assert_accepted(self):
        verified = set(runtime_assurance.hydrate_project_context(self.root)['verified_artifacts'])
        self.assertTrue({'accepted_solution_workbook', 'accepted_result_analysis_workbook',
                         'locked_model_spec'}.issubset(verified), verified)
        self.assertEqual(claim_evidence.inspect_project(self.root)['status'], 'evidence_checked')

    def numerical_snapshot(self):
        entry = self.state['subproblems']['Q1']
        protected = ('solver_execution', 'implementation_conformance', 'primary_execution_status',
                     'analysis_execution_status', 'validated_artifact_hashes',
                     'human_model_approval_status', 'model_challenge_status',
                     'semantic_revision', 'semantic_identity_hash', 'approved_semantic_revision',
                     'approved_semantic_identity_hash', 'validated_semantic_identity_hash',
                     'data_hash', 'validated_data_hash', 'primary_code_sha256', 'analysis_code_sha256',
                     'artifact_hashes', 'result_quality_status', 'result_analysis_status',
                     'analysis_methods', 'result_analysis_requirement_reason', 'code',
                     'result_analysis_code', 'solution_workbook', 'result_analysis_workbook', 'matlab_script')
        paths = [entry[field] for field in ('code', 'result_analysis_code', 'solution_workbook',
                                           'result_analysis_workbook', 'matlab_script')]
        paths += ['figures/q1_f01.png']
        paths += [self.state['preprocessing'][field] for field in ('code', 'workbook')]
        return {'state': {key: deepcopy(entry.get(key)) for key in protected},
                'files': {path: (self.root / path).read_bytes() for path in paths},
                'preprocessing': deepcopy(self.state['preprocessing']),
                'approved_figures': deepcopy(self.state['artifacts']['approved_figures'])}

    def assert_numerical_unchanged(self):
        self.assertEqual(self.numerical_snapshot(), self.numerical_before)
        self.assertFalse(self.state['subproblems']['Q1']['artifacts_stale'])
        self.assertEqual(self.state['subproblems']['Q1']['stale_layers'], [])
        self.assert_accepted()

    def disposition(self, action, target='aux_claim'):
        auxiliary = target == 'aux_claim'
        self.state['subproblems']['Q1']['analysis_evidence_dispositions'] = [{
            'id': 'E12', 'method_or_source': 'Synthetic accepted coefficient analysis',
            'target_claim': target, 'disposition': action,
            'key_finding': ('One tested coefficient does not justify the auxiliary stability wording.'
                            if auxiliary else 'Synthetic review challenges the registered direct-answer claim.'),
            'required_action': ('rewrite the auxiliary claim to the selected tested coefficient only'
                                if auxiliary else 'rewrite the challenged answer after a dedicated model review'),
            'status': 'current',
        }]
        save(self.root, self.state)

    def statuses(self):
        return {row['id']: row['status'] for row in self.fragments}

    def assert_old_proof_rejected(self, audit, compile_report):
        self.assertTrue(delivery.verify_audit_report(
            project=self.latex, main=self.latex / 'main.tex', report=audit))
        self.assertTrue(delivery.verify_compile_report(
            project=self.latex, main=self.latex / 'main.tex', pdf=self.latex / 'main.pdf',
            report=compile_report))

    def assert_current_proof(self, audit, compile_report):
        self.assertEqual(audit['status'], 'passed', audit)
        self.assertEqual(compile_report['status'], 'passed', compile_report)
        self.assertEqual(delivery.verify_audit_report(
            project=self.latex, main=self.latex / 'main.tex', report=audit), [])
        self.assertEqual(delivery.verify_compile_report(
            project=self.latex, main=self.latex / 'main.tex', pdf=self.latex / 'main.pdf',
            report=compile_report), [])

    def proof(self, *, real_compile=False):
        main = self.latex / 'main.tex'
        findings = latex_audit.audit_project(
            main, framework_path=self.root / '模型论文框架.md', require_framework=True)
        self.assertEqual(findings, [], findings)
        audit = latex_audit.write_audit_report(
            main_file=main, findings=findings, framework_path=self.root / '模型论文框架.md')
        self.assertEqual(audit['status'], 'passed', audit)
        if real_compile:
            command = ['xelatex', '-interaction=nonstopmode', '-halt-on-error',
                       '-file-line-error', '-recorder', 'main.tex']
            for _ in range(2):
                compiled = subprocess.run(command, cwd=self.latex, capture_output=True, text=True,
                                          encoding='utf-8', errors='replace', check=False)
                self.assertEqual(compiled.returncode, 0, compiled.stdout[-3000:] + compiled.stderr[-3000:])
        else:
            (self.latex / 'main.pdf').write_bytes(b'%PDF-synthetic-B12-proof-fixture')
            (self.latex / 'main.log').write_text('This is XeTeX\nSynthetic B12 recorder fixture.\n',
                                                encoding='utf-8')
            gate = claims.formal_figure_gate(self.root)
            options = {'project_root': self.root, 'allowed_external_graphics': {
                row['image_token']: self.root / row['image_path'] for row in gate['figure_graphic_bindings']}}
            inputs = delivery.source_bundle_files(main, **options)
            with (self.latex / 'main.fls').open('w', encoding='utf-8') as recorder:
                recorder.write(f'PWD {self.latex.resolve()}\n')
                for path in inputs:
                    token = path.relative_to(self.latex).as_posix() if path.is_relative_to(self.latex) \
                        else '../' + path.relative_to(self.root).as_posix()
                    recorder.write(f'INPUT {token}\n')
        profile = delivery.current_profile_config('diangong')
        report = delivery.write_compile_report(
            project=self.latex, main=main, profile='diangong', engine='xelatex',
            bibliography='none', sequence=['xelatex', 'xelatex'], profile_config=profile)
        self.assert_current_proof(audit, report)
        return audit, report

    def package(self):
        output = self.root / 'submission/reproducibility.zip'
        output.parent.mkdir(exist_ok=True)
        self.assertTrue((self.root / LOCK_RELATIVE_PATH).is_file())
        files = packager.reproducibility_files(self.root, output)
        self.assertNotIn(self.root / LOCK_RELATIVE_PATH, files)
        manifest = packager.build_manifest(self.root, files, kind='reproducibility', metadata={})
        self.assertNotIn(LOCK_RELATIVE_PATH, {record['path'] for record in manifest['files']})
        with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
            for path in files:
                archive.write(path, path.relative_to(self.root).as_posix())
            archive.writestr('submission_manifest.yaml', yaml.safe_dump(manifest, allow_unicode=True))
        report = package_validator.validate_package(self.root, output)
        self.assertEqual(report['status'], 'passed', report['issues'])
        return output

    def exercise_auxiliary_repair(self, action, *, real_compile=False):
        old_audit, old_compile = self.proof(real_compile=real_compile)
        old_package = self.package()
        initial_statuses = self.statuses()
        self.disposition(action)
        before = bytes_in(self.root)
        observed = claims.inspect_project(self.root)
        affected = ['paper.aux.abstract', 'paper.aux.body', 'paper.aux.context']
        self.assertEqual(observed['suggested_stale_fragment_ids'], affected, observed)
        paths = [
            ('claim:aux_claim', 'paper.aux.abstract'),
            ('claim:aux_claim', 'paper.aux.body'),
            ('claim:aux_claim', 'paper.aux.body', 'paper.aux.context'),
        ]
        self.assertEqual(sorted(observed['stale_reason_paths'], key=lambda row: row['path']), [
            {'disposition_id': 'E12', 'question': 'Q1', 'disposition': action, 'path': list(path)}
            for path in paths
        ])
        rejected = claims.formal_figure_gate(self.root)
        self.assertEqual(rejected['status'], 'failed', rejected)
        self.assertEqual(rejected['figure_image_paths'], [])
        self.assertEqual(bytes_in(self.root), before)

        report = synchronizer.synchronize(self.root, write=True, delivery_scope='latex')
        self.assertTrue(report['state_write_performed'], report['issues'])
        self.assertEqual(report['claim_figure_gate']['status'], 'failed')
        self.assertEqual(report['claim_stale_fragments'], affected)
        self.assertEqual(report['stale_questions'], [])
        self.reload()
        expected = {key: ('stale' if key in affected else value) for key, value in initial_statuses.items()}
        self.assertEqual(self.statuses(), expected)
        claims._check_projection(self.state['paper_framework'],
                                 (self.root / '模型论文框架.md').read_text(encoding='utf-8'))
        self.assert_numerical_unchanged()
        self.assert_old_proof_rejected(old_audit, old_compile)
        self.assertEqual(package_validator.validate_package(self.root, old_package)['status'], 'failed')

        # Resolving a disposition is not a review of stale paper fragments.
        self.state['subproblems']['Q1']['analysis_evidence_dispositions'][0]['status'] = 'resolved'
        save(self.root, self.state)
        resolved = synchronizer.synchronize(self.root, write=True)
        self.assertEqual(resolved['status'], 'passed', resolved['issues'])
        self.reload()
        self.assertEqual(self.statuses(), expected)
        self.assertEqual(claims.formal_figure_gate(self.root)['status'], 'failed')
        self.assert_numerical_unchanged()

        # Explicit synthetic author review: revise the wording and confirm all
        # affected rows. This is fixture input, never an automatic repair API.
        self.state['paper_framework']['claim_evidence']['claims'][1]['text'] = (
            'The synthetic result at the selected coefficient is 100.')
        (self.latex / 'aux_body.tex').write_text('Stability: 100.00 for the selected coefficient only.\n',
                                               encoding='utf-8')
        (self.latex / 'aux_context.tex').write_text('Context: no claim about untested coefficients.\n',
                                                  encoding='utf-8')
        for row in self.fragments:
            if row['id'] in affected:
                row['status'] = 'current'
        self.save_projection()
        repaired = claims.formal_figure_gate(self.root)
        self.assertEqual(repaired['status'], 'passed', repaired)
        self.assertEqual(repaired['human_semantic_coverage'], 'not_assessed')
        self.assert_old_proof_rejected(old_audit, old_compile)
        self.assert_numerical_unchanged()
        audit, compile_report = self.proof(real_compile=real_compile)
        # With all normal observations already materialized, a real formal sync
        # preserves State/Framework bytes and the regenerated proof read set.
        proof_sources = ((self.root / 'state/project_state.yaml').read_bytes(),
                         (self.root / '模型论文框架.md').read_bytes())
        final = synchronizer.synchronize(self.root, write=True, delivery_scope='latex')
        self.assertEqual(final['status'], 'passed', final['issues'])
        self.assertFalse(final['state_write_performed'])
        self.assertEqual(proof_sources, ((self.root / 'state/project_state.yaml').read_bytes(),
                                        (self.root / '模型论文框架.md').read_bytes()))
        self.assertEqual(delivery.verify_compile_report(
            project=self.latex, main=self.latex / 'main.tex', pdf=self.latex / 'main.pdf',
            report=compile_report), [])
        self.assertEqual(package_validator.validate_package(self.root, old_package)['status'], 'failed')
        self.package()

    def test_b12_reject_auxiliary_continuous_repair_retains_accepted_primary(self):
        self.exercise_auxiliary_repair('reject')

    def test_b12_modify_auxiliary_stales_only_dependents_and_retains_accepted_primary(self):
        original = self.statuses()
        affected = ['paper.aux.abstract', 'paper.aux.body', 'paper.aux.context']
        self.disposition('modify')
        before = bytes_in(self.root)
        gate = claims.formal_figure_gate(self.root)
        self.assertEqual(gate['status'], 'failed', gate)
        self.assertEqual(gate['figure_image_paths'], [])
        self.assertEqual(bytes_in(self.root), before)
        report = synchronizer.synchronize(self.root, write=True, delivery_scope='latex')
        self.assertTrue(report['state_write_performed'], report['issues'])
        self.assertEqual(report['claim_figure_gate']['status'], 'failed')
        self.assertEqual(report['invalidated_claim_ids'], ['aux_claim'])
        self.assertEqual(report['claim_stale_fragments'], affected)
        self.assertEqual(report['stale_questions'], [])
        self.reload()
        self.assertEqual(self.statuses(), {
            key: ('stale' if key in affected else value) for key, value in original.items()
        })
        claims._check_projection(self.state['paper_framework'],
                                 (self.root / '模型论文框架.md').read_text(encoding='utf-8'))
        self.assert_numerical_unchanged()

    @unittest.skipUnless(shutil.which('xelatex'), 'XeLaTeX is not installed')
    def test_real_xelatex_b12_repair_regenerates_current_figure_and_package_proof(self):
        self.exercise_auxiliary_repair('reject', real_compile=True)

    def test_core_claim_disposition_stales_its_figure_and_all_registered_consumptions(self):
        """Verify dependency targeting only, not acceptance of a revised core model."""
        original = self.statuses()
        affected = ['paper.abstract.q1', 'paper.figure.q1', 'paper.result.q1']
        self.disposition('reject', target='answer_claim')
        gate = claims.formal_figure_gate(self.root)
        self.assertEqual(gate['status'], 'failed')
        self.assertEqual(gate['figure_image_paths'], [])
        report = synchronizer.synchronize(self.root, write=True)
        self.assertTrue(report['write_performed'], report['issues'])
        self.assertEqual(report['invalidated_claim_ids'], ['answer_claim'])
        self.assertEqual(report['claim_stale_fragments'], affected)
        self.reload()
        self.assertEqual(self.statuses(), {
            key: ('stale' if key in affected else value) for key, value in original.items()
        })
        self.assertEqual(claims.formal_figure_gate(self.root)['status'], 'failed')

    def test_b11_same_selected_value_drift_preserves_original_artifact_invalidation(self):
        self.disposition('reject')
        workbook = self.root / self.state['subproblems']['Q1']['solution_workbook']
        book = openpyxl.load_workbook(workbook)
        selected = book['核心指标']['B2'].value
        note = book.create_sheet('unrelated-note')
        note.append(['note'])
        note.append(['changed bytes, selected value preserved'])
        book.save(workbook)
        book.close()
        with workbook.open('rb') as stream:
            changed = openpyxl.load_workbook(stream, read_only=True, data_only=True)
            self.assertEqual(changed['核心指标']['B2'].value, selected)
            changed.close()
        self.assertNotEqual(hashlib.sha256(workbook.read_bytes()).hexdigest(),
                            self.state['subproblems']['Q1']['validated_artifact_hashes']['solution_workbook'])
        self.assertNotEqual(claim_evidence.inspect_project(self.root)['status'], 'evidence_checked')
        report = synchronizer.synchronize(self.root, write=True)
        self.assertIn('Q1', report['stale_questions'], report['issues'])
        self.reload()
        self.assertTrue(self.state['subproblems']['Q1']['artifacts_stale'])
        self.assertTrue(all(status == 'stale' for status in self.statuses().values()))
        self.assertEqual(claims.formal_figure_gate(self.root)['status'], 'failed')
        self.assertNotIn('accepted_solution_workbook',
                         runtime_assurance.hydrate_project_context(self.root)['verified_artifacts'])


if __name__ == '__main__':
    unittest.main()
