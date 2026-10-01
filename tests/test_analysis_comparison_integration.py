"""Existing coordinators inspect synthetic candidate bytes before numerical acceptance.

These fixture literals test protocol integration; native numerical execution has its
separate Python/MATLAB smoke and is never inferred from these synthetic receipts.
"""
from copy import deepcopy
import hashlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
for directory in (ROOT, ROOT / "scripts"):
    if str(directory) not in sys.path:
        sys.path.insert(0, str(directory))
import analysis_comparison_gate as gate
import execution_protocol as protocol
import runtime_assurance as runtime
import state_transitions as transitions
import sync_project as sync
from tests import test_user_execution_contract as legacy
from tests import analysis_comparison_smoke as smoke
from tests import reading_plan_cases

CONTRACT = yaml.safe_load((ROOT / "core/state_transition_contract.yaml").read_text(encoding="utf-8"))


class ComparisonIntegrationTests(unittest.TestCase):
    def project(self, root):
        fixture = legacy.UserExecutionContractTests()
        primary_code = fixture.make_project(root)
        state = fixture.accept_primary(root, primary_code)
        primary = root / state["subproblems"]["Q1"]["solution_workbook"]
        book = openpyxl.load_workbook(primary)
        detail = book.create_sheet("状态明细")
        detail.append(["记录键", "实例或场景", "指标", "数值", "单位"])
        detail.append(["prediction-4", "holdout", "prediction", 16.0, "dimensionless"])
        book.save(primary)
        self.assertEqual(legacy.RECEIPT.validate_one(root, primary, state, True), [])
        fixture.primary_workbook_sha = hashlib.sha256(primary.read_bytes()).hexdigest()
        identity = smoke.model_identity()
        digest = smoke.semantic_identity.semantic_identity_hash(identity)
        text = smoke.framework(identity)
        entry = state["subproblems"]["Q1"]
        entry.update(semantic_identity_schema_version="1.0.0", semantic_revision=1,
                     validated_semantic_revision=1, approved_semantic_revision=1,
                     semantic_identity_hash=digest, validated_semantic_identity_hash=digest,
                     approved_semantic_identity_hash=digest, model_challenge_status="passed",
                     human_model_approval_status="approved", problem_contract_status="frozen",
                     semantic_closure_status="passed", complexity_sanity_status="passed",
                     semantic_text_hash=smoke.semantic_identity.sha256_text(
                         smoke.semantic_identity.semantic_scope(smoke.semantic_identity.question_sections(text)["Q1"])))
        plan_hash = smoke.activate_comparison(root, state, "python")
        # This coordinator fixture selects only the holdout; the native smoke covers all points.
        for check in entry["analysis_comparison"]["checks"]:
            check["evidence_refs"] = [row for row in check["evidence_refs"] if row["id"].endswith("-4")]
        plan_hash = gate.COMPARISON.plan_sha256(entry["analysis_comparison"])
        fixture.write_state(root, state)
        config = fixture.config("analysis", "问题一结果深化分析.xlsx")
        config.update(analysis_comparison_protocol_version="1.0.0", analysis_comparison_plan_sha256=plan_hash)
        code = root / "问题一求解/问题一结果深化分析.py"
        fixture.write_code(code, config)
        issues, parsed = legacy.CODE.validate_script(root, code, "analysis")
        self.assertEqual(issues, [], issues)
        legacy.CODE.update_state(root, parsed, code)
        state = fixture.read_state(root)
        entry = state["subproblems"]["Q1"]
        for index, check in enumerate(entry["analysis_comparison"]["checks"], 1):
            check["disposition_ref"] = f"E{index}"
        entry["analysis_evidence_dispositions"] = [
            {"id": f"E{index}", "status": "current", "target_claim": check["target_claim"],
             "disposition": "support", "method_or_source": check["kind"],
             "key_finding": "Synthetic exact values meet the predeclared threshold",
             "required_action": "Use only this fixture claim", "paper_or_figure_anchor": "fixture:only"}
            for index, check in enumerate(entry["analysis_comparison"]["checks"], 1)]
        fixture.write_state(root, state)
        workbook = fixture.make_analysis_workbook(root, code)
        book = openpyxl.load_workbook(workbook)
        book["运行配置"].append(["analysis_comparison_protocol_version", "1.0.0"])
        book["运行配置"].append(["analysis_comparison_plan_sha256", plan_hash])
        model = book.create_sheet("多模型检验")
        model.append(["检验ID", "记录键", "主模型ID", "对照模型ID", "评价协议ID", "实例或场景", "指标", "单位", "主模型数值", "对照模型数值", "差异类型", "差异", "判据ID", "判定"])
        model.append(["CMP-Q1-01", "model-4", "MODEL-Q1-01", "MODEL-Q1-02", "EVAL-Q1-01", "holdout", "prediction", "dimensionless", 16.0, 11.0, "difference", -5.0, "CRIT-CMP-Q1-01", True])
        algorithm = book.create_sheet("同模型多算法检验")
        algorithm.append(["检验ID", "记录键", "模型ID", "基准算法ID", "对照算法ID", "评价协议ID", "实例或场景", "重复编号", "指标", "单位", "基准数值", "对照数值", "差异类型", "差异", "判据ID", "判定"])
        algorithm.append(["CMP-Q1-02", "algorithm-4", "MODEL-Q1-01", "ALGO-Q1-01", "ALGO-Q1-02", "EVAL-Q1-02", "holdout", 1, "prediction", "dimensionless", 16.0, 16.0, "difference", 0.0, "CRIT-CMP-Q1-02", True])
        book.save(workbook)
        return fixture, state, code, workbook, primary

    def test_overlay_requires_analysis_source_protocol_and_exact_echo(self):
        good = {"stage": "analysis", "run_receipt_protocol_version": "1.1.0",
                "analysis_comparison_protocol_version": "1.0.0", "analysis_comparison_plan_sha256": "a" * 64}
        receipt = {**good, "run_receipt_version": "1.1.0"}
        self.assertEqual(protocol.comparison_receipt_issues(receipt, good), [])
        for mutation in ({"stage": "primary"}, {"analysis_comparison_protocol_version": "2.0.0"},
                         {"analysis_comparison_plan_sha256": "b" * 64}, {"run_receipt_version": "1.0.0"}):
            self.assertTrue(protocol.comparison_receipt_issues(receipt | mutation, good))
        self.assertTrue(protocol.comparison_config_issues({"analysis_comparison_protocol_version": "1.0.0"}))

    def test_inactive_current_analysis_preserves_native_legacy_artifact_paths(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            reading_plan_cases.build_project(ROOT, root, "current")
            state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
            entry = state["subproblems"]["Q1"]
            self.assertEqual(entry["result_analysis_workbook"],
                             str(Path("问题一求解") / "问题一结果深化分析.xlsx"))
            report = gate.inspect_gate(root, state, "Q1", boundary="current")
            self.assertFalse(report["enabled"])
            self.assertEqual(report["issues"], [], report)
            qualified = runtime.hydrate_project_context(root, "Q1")
            self.assertIn("validated_results", qualified["verified_artifacts"], qualified["artifact_evidence"])

    def test_candidate_accepts_without_accepted_analysis_qualification(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, state, _, workbook, primary = self.project(root)
            baseline = primary.read_bytes()
            self.assertNotEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")
            # A forbidden candidate -> Sources.qualify -> accepted-analysis loop would call hydrate.
            with patch.object(runtime, "hydrate_project_context", side_effect=AssertionError("candidate acceptance loop")):
                self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            self.assertEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")
            self.assertEqual(primary.read_bytes(), baseline)

    def test_copied_baseline_and_missing_required_comparison_cannot_pass(self):
        for case in ("copied_baseline", "missing_algorithm"):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                _, state, _, workbook, _ = self.project(root)
                book = openpyxl.load_workbook(workbook)
                if case == "copied_baseline":
                    book["多模型检验"]["I2"] = 17.0
                    book["多模型检验"]["L2"] = -6.0
                else:
                    del book["同模型多算法检验"]
                book.save(workbook)
                self.assertTrue(legacy.RECEIPT.validate_one(root, workbook, state, True))
                self.assertNotEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")

    def test_current_primary_tamper_during_evidence_check_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, state, _, workbook, primary = self.project(root)
            original = gate.COMPARISON.inspect_evidence
            def tamper(*args, **kwargs):
                report = original(*args, **kwargs)
                primary.write_bytes(primary.read_bytes() + b"changed-after-selection")
                return report
            with patch.object(gate.COMPARISON, "inspect_evidence", side_effect=tamper):
                with self.assertRaises((ValueError, RuntimeError)):
                    legacy.RECEIPT.validate_one(root, workbook, state, True)
            self.assertNotEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")

    def test_current_analysis_replacement_after_artifact_hash_read_cannot_qualify(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, _ = self.project(root)
            self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            fixture.write_state(root, state)
            replacement = io.BytesIO()
            book = openpyxl.load_workbook(workbook)
            book.properties.description = "Synthetic replacement after the accepted artifact hash read"
            book.save(replacement)
            book.close()
            original = runtime._file_evidence

            def replace_after_hash(*args, **kwargs):
                row = original(*args, **kwargs)
                if kwargs.get("artifact") == "accepted_result_analysis_workbook":
                    self.assertEqual(row["status"], "verified")
                    workbook.write_bytes(replacement.getvalue())
                return row

            with patch.object(runtime, "_file_evidence", side_effect=replace_after_hash):
                qualified = runtime.hydrate_project_context(root, "Q1")
            analysis = next(row for row in qualified["artifact_evidence"]
                            if row["artifact"] == "accepted_result_analysis_workbook")
            self.assertEqual(analysis["status"], "not_accepted")
            self.assertIn("analysis bytes differ from the accepted analysis workbook", analysis["reason"])
            self.assertNotIn("validated_results", qualified["verified_artifacts"])
            self.assertIn("accepted_solution_workbook", qualified["verified_artifacts"])

    def test_delivery_primary_replacement_after_prerequisite_read_cannot_commit(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, _, code, _, primary = self.project(root)
            before_state = (root / "state/project_state.yaml").read_bytes()
            _, config = gate.stage_code.parse_stage_config(code)
            original = gate.COMPARISON.inspect_plan
            replaced = False

            def replace_after_plan(*args, **kwargs):
                nonlocal replaced
                report = original(*args, **kwargs)
                if not replaced:
                    primary.write_bytes(primary.read_bytes() + b"changed-after-primary-prerequisite")
                    replaced = True
                return report

            with patch.object(gate.COMPARISON, "inspect_plan", side_effect=replace_after_plan):
                with self.assertRaisesRegex(ValueError, "baseline bytes differ from the accepted primary workbook"):
                    legacy.CODE.update_state(root, config, code)
            self.assertTrue(replaced)
            self.assertEqual((root / "state/project_state.yaml").read_bytes(), before_state)

    def test_current_caller_cannot_register_replacement_by_forging_accepted_hash(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, _ = self.project(root)
            self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            fixture.write_state(root, state)
            captured = deepcopy(state)
            self.assertEqual(gate.inspect_gate(root, state, "Q1", boundary="current")["issues"], [])
            book = openpyxl.load_workbook(workbook)
            book.properties.description = "Unaccepted replacement with caller-forged accepted SHA"
            book.save(workbook)
            book.close()
            state["subproblems"]["Q1"]["validated_artifact_hashes"]["result_analysis_workbook"] = (
                hashlib.sha256(workbook.read_bytes()).hexdigest())
            gate.assert_authorization_context(state, captured, "Q1")
            current = gate.inspect_gate(root, state, "Q1", boundary="current")
            self.assertIn("comparison caller analysis accepted hash differs from captured State", current["issues"])
            self.assertIn("comparison analysis bytes differ from the accepted analysis workbook", current["issues"])
            # Candidate receipt validation retains the coordinator's legitimate batch-output contract.
            self.assertEqual(gate.inspect_gate(root, state, "Q1", boundary="receipt")["issues"], [])

    def test_delivery_caller_cannot_override_captured_primary_accepted_hash(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, state, code, _, primary = self.project(root)
            captured = deepcopy(state)
            primary.write_bytes(primary.read_bytes() + b"unaccepted-primary-replacement")
            state["subproblems"]["Q1"]["validated_artifact_hashes"]["solution_workbook"] = (
                hashlib.sha256(primary.read_bytes()).hexdigest())
            gate.assert_authorization_context(state, captured, "Q1")
            report = gate.inspect_gate(root, state, "Q1", boundary="delivery", code_path=code)
            self.assertIn("comparison caller primary accepted hash differs from captured State", report["issues"])
            self.assertIn("comparison baseline bytes differ from the accepted primary workbook", report["issues"])

    def test_new_sheet_without_registry_is_strict_activation(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, _ = self.project(root)
            entry = state["subproblems"]["Q1"]
            entry.pop("analysis_comparison")
            entry["analysis_methods"] = ["参数敏感性"]
            fixture.write_state(root, state)
            self.assertTrue(gate.workbook_present(workbook))
            self.assertTrue(legacy.RECEIPT.validate_one(root, workbook, state, True))
            self.assertNotEqual(entry["analysis_execution_status"], "accepted")

    def test_decoded_receipt_markers_activate_without_new_tables_or_source_markers(self):
        for encoding in ("numeric_entities", "rich_text"):
            with self.subTest(encoding=encoding), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                fixture = legacy.UserExecutionContractTests()
                fixture.accept_primary(root, fixture.make_project(root))
                fixture.activate_analysis(root)
                code = fixture.make_analysis_code(root)
                issues, config = legacy.CODE.validate_script(root, code, "analysis")
                self.assertEqual(issues, [])
                legacy.CODE.update_state(root, config, code)
                state = fixture.read_state(root)
                workbook = fixture.make_analysis_workbook(root, code)
                self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
                book = openpyxl.load_workbook(workbook)
                for field, value in zip(protocol.COMPARISON_FIELDS, ("1.0.0", "a" * 64)):
                    book["运行配置"].append([field, value])
                book.save(workbook)
                book.close()
                output = io.BytesIO()
                with zipfile.ZipFile(workbook) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
                    for name in source.namelist():
                        raw = source.read(name)
                        if name.startswith("xl/worksheets/"):
                            for field in protocol.COMPARISON_FIELDS:
                                key = field.encode("utf-8")
                                if encoding == "numeric_entities":
                                    raw = raw.replace(key, "".join(f"&#{ord(char)};" for char in field).encode("ascii"))
                                else:
                                    split = len(field) // 2
                                    runs = (f"<r><t>{field[:split]}</t></r><r><t>{field[split:]}</t></r>").encode("ascii")
                                    raw = raw.replace(b"<t>" + key + b"</t>", runs)
                        self.assertFalse(any(field.encode("utf-8") in raw for field in protocol.COMPARISON_FIELDS))
                        target.writestr(name, raw)
                workbook.write_bytes(output.getvalue())
                receipt, issues = legacy.RECEIPT.configuration_map(workbook)
                self.assertEqual(issues, [])
                self.assertEqual(receipt[protocol.COMPARISON_FIELDS[0]], "1.0.0")
                self.assertEqual(receipt[protocol.COMPARISON_FIELDS[1]], "a" * 64)
                digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
                entry = state["subproblems"]["Q1"]
                for registry in ("artifact_hashes", "validated_artifact_hashes"):
                    entry[registry]["result_analysis_workbook"] = digest
                fixture.write_state(root, state)
                self.assertFalse(gate.present(entry, config=config))
                with patch.object(gate.openpyxl, "load_workbook", side_effect=AssertionError("probe must not use openpyxl")):
                    self.assertTrue(gate.workbook_present(workbook))
                report = gate.inspect_gate(root, state, "Q1", boundary="current")
                self.assertTrue(report["enabled"])
                self.assertIn("comparison activation is missing its complete registry/method/protocol binding", report["issues"])

    def test_separate_cell_fragments_do_not_activate_legacy_workbook(self):
        with tempfile.TemporaryDirectory() as temporary:
            workbook = Path(temporary) / "separate-fragments.xlsx"
            book = openpyxl.Workbook()
            for field in protocol.COMPARISON_FIELDS:
                split = len(field) // 2
                book.active.append([field[:split], field[split:]])
            book.save(workbook)
            book.close()
            with patch.object(gate.openpyxl, "load_workbook", side_effect=AssertionError("probe must not use openpyxl")):
                self.assertFalse(gate.workbook_present(workbook))

    def test_required_item_cannot_be_waived_by_not_required_reason(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, code, workbook, _ = self.project(root)
            entry = state["subproblems"]["Q1"]
            entry.update(result_analysis_status="not_required", result_analysis_requirement_reason="A nonempty reason")
            fixture.write_state(root, state)
            before = deepcopy(state)
            self.assertTrue(gate.inspect_gate(root, state, "Q1", boundary="plan")["issues"])
            self.assertTrue(legacy.CODE.validate_script(root, code, "analysis")[0])
            self.assertTrue(legacy.RECEIPT.validate_one(root, workbook, state, True))
            self.assertEqual(state, before)

    def test_read_set_exists_without_optional_a2_b2_or_c(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, state, code, workbook, primary = self.project(root)
            entry = state["subproblems"]["Q1"]
            self.assertNotIn("implementation_conformance_policy", entry)
            self.assertNotIn("claim_consumption_policy", state["paper_framework"])
            self.assertNotIn("review_receipt_policy", state)
            report = gate.inspect_gate(root, state, "Q1", boundary="receipt", workbook=workbook)
            self.assertEqual(report["issues"], [], report)
            observed = report["observed_sources"]
            for path in (code, workbook, primary, root / "data.csv", root / "模型论文框架.md", root / "state/project_state.yaml"):
                self.assertEqual(observed["project"][path.relative_to(root).as_posix()], hashlib.sha256(path.read_bytes()).hexdigest())
            self.assertIn("scripts/analysis_comparison.py", observed["skill"])
            self.assertIn("templates/code/hsk_pipeline/workbook_validation.py", observed["skill"])
            delivery = gate.inspect_gate(root, state, "Q1", boundary="delivery", code_path=code)
            self.assertEqual(delivery["issues"], [], delivery)
            self.assertEqual(delivery["observed_sources"]["project"][primary.relative_to(root).as_posix()],
                             hashlib.sha256(primary.read_bytes()).hexdigest())
            self.assertNotIn(workbook.relative_to(root).as_posix(), delivery["observed_sources"]["project"])

    def test_caller_cannot_forge_captured_scope_approval_or_permissions(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, _ = self.project(root)
            captured = deepcopy(state)
            captured["subproblems"]["Q1"]["human_model_approval_status"] = "awaiting_model_approval"
            fixture.write_state(root, captured)
            before = deepcopy(state)
            issues = legacy.RECEIPT.validate_one(root, workbook, state, True)
            self.assertTrue(any("authorization context differs" in issue for issue in issues), issues)
            self.assertEqual(state, before)
            for field, value in (("analysis_methods", []), ("analysis_comparison", {}), ("analysis_evidence_dispositions", []),
                                 ("approved_semantic_revision", 2), ("approved_semantic_identity_hash", "b" * 64)):
                forged = deepcopy(state)
                forged["subproblems"]["Q1"][field] = value
                with self.subTest(field=field), self.assertRaises(ValueError):
                    gate.assert_authorization_context(forged, state, "Q1")
            for field, value in (("execution", {"solver_backend": "matlab"}), ("review_receipt_policy", {}), ("review_receipts", {})):
                forged = deepcopy(state)
                forged[field] = value
                with self.subTest(field=field), self.assertRaises(ValueError):
                    gate.assert_authorization_context(forged, state, "Q1")

    def test_primary_then_analysis_receipt_batch_allows_execution_updates(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, primary = self.project(root)
            entry = state["subproblems"]["Q1"]
            entry.update(primary_execution_status="awaiting_user_execution", result_quality_status="pending", status="designed")
            entry["validated_artifact_hashes"].pop("solution_workbook", None)
            fixture.write_state(root, state)
            self.assertEqual(legacy.RECEIPT.validate_one(root, primary, state, True), [])
            # The disk still has pending primary execution; only coordinator facts changed in memory.
            self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            self.assertEqual(entry["analysis_execution_status"], "accepted")

    def test_candidate_budget_rejects_before_openpyxl_or_acceptance(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            _, state, _, workbook, _ = self.project(root)
            output = io.BytesIO()
            with zipfile.ZipFile(workbook) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
                for name in source.namelist():
                    raw = source.read(name)
                    if name == "xl/worksheets/sheet1.xml":
                        raw = raw.replace(b'r="A1"', b'r="ZZ1"', 1)
                    destination.writestr(name, raw)
            workbook.write_bytes(output.getvalue())
            with patch.object(gate.openpyxl, "load_workbook", side_effect=AssertionError("unbounded candidate parser")):
                with self.assertRaisesRegex(ValueError, "cell/column budget"):
                    legacy.RECEIPT.validate_one(root, workbook, state, True)
            self.assertNotEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")

    def test_core_rejection_returns_existing_lifecycle_without_b2_policy(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, primary = self.project(root)
            row = state["subproblems"]["Q1"]["analysis_evidence_dispositions"][0]
            row.update(disposition="reject", impact_scope="core_answer", return_stage="solve_validate")
            fixture.write_state(root, state)
            baseline = primary.read_bytes()
            self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            entry = state["subproblems"]["Q1"]
            self.assertEqual(entry["analysis_execution_status"], "redo_required")
            self.assertIn("solution_workbook", entry["stale_layers"])
            self.assertEqual(state["project"]["current_phase"], "solve_validate")
            self.assertEqual(primary.read_bytes(), baseline)
            row["status"] = "resolved"
            self.assertTrue(runtime._current_structured_rejections(state, ["Q1"])["Q1"]["blocks_primary"])

    def test_comparison_modify_does_not_suppress_noncomparison_instability(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture, state, _, workbook, _ = self.project(root)
            state["subproblems"]["Q1"]["analysis_evidence_dispositions"][0].update(
                disposition="modify", impact_scope="auxiliary_wording")
            fixture.write_state(root, state)
            book = openpyxl.load_workbook(workbook)
            book["结论稳定性汇总"].append(["comparison scoped finding", "多模型检验", "fixture", False])
            book.save(workbook)
            self.assertEqual(legacy.RECEIPT.validate_one(root, workbook, state, True), [])
            self.assertEqual(state["subproblems"]["Q1"]["analysis_execution_status"], "accepted")
            self.assertTrue(gate.inspect_gate(root, state, "Q1", boundary="current")["issues"])
            state["subproblems"]["Q1"]["analysis_evidence_dispositions"][0]["status"] = "resolved"
            fixture.write_state(root, state)
            self.assertEqual(gate.inspect_gate(root, state, "Q1", boundary="current")["issues"], [])
            book = openpyxl.load_workbook(workbook)
            book["结论稳定性汇总"]["D2"] = False
            book.save(workbook)
            self.assertTrue(legacy.RECEIPT.validate_one(root, workbook, state, True))
            self.assertEqual(state["subproblems"]["Q1"]["result_analysis_status"], "redo_required")

    def test_action_closure_only_stales_exact_fragments_and_dependents(self):
        state = {"paper_framework": {"sync_status": "current", "paper_fragments": [
            {"id": "paper.q1", "status": "current", "depends_on": ["claim:case_change"], "anchor": "answer"},
            {"id": "paper.abstract", "status": "current", "depends_on": ["paper.q1"], "anchor": "abstract"},
            {"id": "paper.q2", "status": "current", "depends_on": [], "anchor": "other"}]},
                 "subproblems": {"Q1": {"status": "analyzed", "primary_execution_status": "accepted",
                    "result_quality_status": "passed", "result_summary_status": "current", "artifacts_stale": False,
                    "stale_layers": [], "analysis_comparison": {"checks": [{"disposition_ref": "E1"}]},
                    "analysis_evidence_dispositions": [{"id": "E1", "status": "current", "disposition": "modify",
                                                       "target_claim": "case_change", "impact_scope": "auxiliary_wording"}]}}}
        entry_before = deepcopy(state["subproblems"]["Q1"])
        closure = gate.action_fragment_closure(state, state["subproblems"]["Q1"])
        self.assertEqual(closure["fragment_ids"], ["paper.abstract", "paper.q1"])
        transitions.mark_claim_fragments_stale(state, "Q1", closure["fragment_ids"])
        self.assertEqual([row["status"] for row in state["paper_framework"]["paper_fragments"]], ["stale", "stale", "current"])
        self.assertEqual(state["subproblems"]["Q1"], entry_before)
        self.assertEqual(state["paper_framework"]["sync_status"], "current")
        entry = state["subproblems"]["Q1"]
        entry["analysis_evidence_dispositions"][0].update(target_claim="unbound_claim", paper_or_figure_anchor="unknown")
        self.assertEqual(gate.action_fragment_closure(state, entry)["unbound_ids"], ["E1"])

    def test_changed_plan_only_invalidates_analysis_and_downstream(self):
        entry = {"status": "analyzed", "primary_execution_status": "accepted", "result_quality_status": "passed",
                 "analysis_execution_status": "accepted", "result_analysis_status": "passed",
                 "result_analysis_code": "analysis.py", "stale_layers": []}
        state = {"project": {"current_phase": "result_analysis"}, "subproblems": {"Q1": deepcopy(entry)}}
        snapshot = {"artifact_hashes": {}, "analysis_comparison_observed": {"plan_issues": ["changed digest"]}}
        events = sync._snapshot_transition_events(entry, snapshot)
        self.assertIn("analysis_comparison_plan_changed", events)
        transitions.apply_transition(state, event="analysis_comparison_plan_changed", source_question="Q1", contract=CONTRACT)
        changed = state["subproblems"]["Q1"]
        self.assertEqual(changed["primary_execution_status"], "accepted")
        self.assertEqual(changed["result_quality_status"], "passed")
        self.assertNotIn("solution_workbook", changed["stale_layers"])
        self.assertIn("result_analysis_workbook", changed["stale_layers"])


if __name__ == "__main__":
    unittest.main()
