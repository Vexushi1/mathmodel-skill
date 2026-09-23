"""Evidence-bound migration proposals on isolated maintenance projects.

The primary base really executes the existing a*x=b Python micro-fixture. Analysis
receipts below are explicitly synthetic reader fixtures, not claimed native runs;
real MATLAB output is covered by test_solver_backend_end_to_end.verify.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT / "tests")]
import project_solver_backend as preview
import project_transaction as transaction
import stage_code
import validate_user_execution as receipts
from solver_backend_mixed_smoke import python_primary as current_python_primary
from test_solver_backend_end_to_end import CONFIG, file_hash, reference_digest, save_state


def files(root):
    return {p.relative_to(root).as_posix(): file_hash(p) for p in root.rglob("*") if p.is_file()}


def load(root):
    return yaml.safe_load((root / transaction.STATE_RELATIVE_PATH).read_text(encoding="utf-8"))


def python_primary(root, question, upstream):
    """Run the real micro solver, then preserve its evidence as a v9 history fixture."""
    state_path = root / transaction.STATE_RELATIVE_PATH
    if state_path.is_file():
        current = load(root)
        current["execution"] = {
            "solver_backend": "python",
            "solver_backend_selection_reason": "Temporary current policy for real fixture execution",
        }
        for entry in current.get("subproblems", {}).values():
            for selection in (entry.get("solver_execution") or {}).values():
                selection.pop("backend", None)
                selection.pop("selection_reason", None)
        save_state(root, current)
    current_python_primary(root, question, upstream)
    historical = load(root)
    historical.pop("execution", None)
    for entry in historical["subproblems"].values():
        for selection in (entry.get("solver_execution") or {}).values():
            selection.update(backend="python", selection_reason="Historical v9 per-stage selection")
    save_state(root, historical)


def update_receipt(root, relative, updates):
    book = openpyxl.load_workbook(root / relative)
    try:
        sheet = book["运行配置"]
        found = {row[0].value: row[1] for row in sheet.iter_rows(min_row=2)}
        for key, value in updates.items():
            if key in found:
                found[key].value = value
            else:
                sheet.append((key, value))
        book.save(root / relative)
    finally:
        book.close()


def matlab_stage(root, question="Q1", stage="analysis"):
    """Instantiate the actual MATLAB template, never execute or accept it locally."""
    state = load(root)
    number = stage_code.question_number(question)
    problem = stage_code.question_name(question)
    filename = f"q{number}_{'solver' if stage == 'primary' else 'analysis'}.m"
    text = (ROOT / "templates/code/matlab" / ("q1_solver.m" if stage == "primary" else "q1_analysis.m")).read_text(encoding="utf-8")
    text = text.replace("q1_solver()", f"q{number}_solver()", 1).replace("q1_analysis()", f"q{number}_analysis()", 1)
    match = CONFIG.search(text)
    config = json.loads(match.group(1))
    config.update(problem_name=problem, data_paths=["input.json"], data_sha256=reference_digest(root, ["input.json"]),
                  expected_workbook=f"{problem}求解/{problem}{'求解结果' if stage == 'primary' else '结果深化分析'}.xlsx")
    entry = state["subproblems"].setdefault(question, {"status": "designed", "result_analysis_status": "pending"})
    if stage == "analysis":
        config["primary_workbook_sha256"] = entry["validated_artifact_hashes"]["solution_workbook"]
        entry.update(result_analysis_requirement_reason="Synthetic independent sensitivity requirement", analysis_methods=["coefficient"])
    text = text[:match.start(1)] + json.dumps(config, ensure_ascii=False, separators=(",", ":")) + text[match.end(1):]
    relative = f"{problem}求解/{filename}"
    code = root / relative
    code.parent.mkdir(exist_ok=True)
    code.write_text(text, encoding="utf-8")
    codekey, hashkey, _ = preview._STAGE_KEYS[stage]
    entry.update({codekey: relative, hashkey: file_hash(code), f"{stage}_execution_status": "awaiting_user_execution",
                  "data_hash": config["data_sha256"]})
    entry.setdefault("solver_execution", {})[stage] = {
        "backend": "matlab", "selection_reason": "Synthetic static MATLAB fixture", "bundle_sha256": reference_digest(root, [relative])}
    save_state(root, state)
    return relative


def synthetic_python_analysis(root):
    state = load(root)
    entry = state["subproblems"]["Q1"]
    _, config = stage_code.parse_stage_config(root / entry["code"])
    config.update(stage="analysis", expected_workbook="问题一求解/问题一结果深化分析.xlsx",
                  primary_workbook_sha256=file_hash(root / entry["solution_workbook"]))
    config.pop("primary_quality_protocol_version", None)
    relative = "问题一求解/问题一结果深化分析.py"
    code = root / relative
    code.write_text("# Synthetic static source for receipt reader tests, not a user solver.\nRUN_CONFIG = " + repr(config)
                    + "\ndef main():\n    return None\n", encoding="utf-8")
    bundle = reference_digest(root, [relative])
    entry.update(result_analysis_code=relative, analysis_code_sha256=file_hash(code),
                 analysis_execution_status="awaiting_user_execution", result_analysis_requirement_reason="Synthetic sensitivity reader",
                 analysis_methods=["coefficient"])
    entry["solver_execution"]["analysis"] = {"backend": "python", "selection_reason": "Synthetic reader fixture",
                                               "bundle_sha256": bundle}
    rc, issues = receipts.configuration_map(root / entry["solution_workbook"])
    assert not issues, issues
    rc.update(stage="analysis", code_sha256=file_hash(code), code_bundle_sha256=bundle,
              primary_workbook_sha256=config["primary_workbook_sha256"])
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in {
        "运行配置": [("项目", "值"), *rc.items()],
        "分析设计": [("分析项", "设计"), ("coefficient", "synthetic reader fixture")],
        "结论稳定性汇总": [("结论", "是否保持"), ("positive", True)],
        "参数敏感性": [("coefficient", "solution"), (1.5, 4), (2, 3), (3, 2)],
    }.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    wb = root / config["expected_workbook"]
    book.save(wb)
    book.close()
    issues = receipts.validate_one(root, wb, state, False)
    assert not issues, issues
    # This is a pre-migration accepted record. Current receipt writes correctly
    # refuse to grant new acceptance to its historical per-stage selectors.
    entry.update(analysis_execution_status="accepted", result_analysis_status="passed",
                 result_analysis_workbook=wb.relative_to(root).as_posix(), status="analyzed")
    for field in ("artifact_hashes", "validated_artifact_hashes"):
        entry.setdefault(field, {}).update(analysis_code=file_hash(code), result_analysis_workbook=file_hash(wb))
    entry["solver_execution"]["analysis"]["validated_bundle_sha256"] = bundle
    save_state(root, state)


class MigrationPreviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base_tmp = tempfile.TemporaryDirectory()
        cls.base = Path(cls.base_tmp.name)
        shutil.copyfile(ROOT / "tests/fixtures/solver_backends/input.json", cls.base / "input.json")
        python_primary(cls.base, "问题一", False)

    @classmethod
    def tearDownClass(cls):
        cls.base_tmp.cleanup()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "project"
        shutil.copytree(self.base, self.root)

    def change(self, mutate):
        state = load(self.root)
        mutate(state)
        save_state(self.root, state)

    def run_preview(self, target="python", reason="全题统一工程后端"):
        before = files(self.root)
        report = preview.preview_migration(self.root, target_backend=target, reason=reason)
        self.assertEqual(files(self.root), before, "preview changed project bytes or created files")
        self.assertFalse(report["migration_authorized"])
        self.assertFalse(report["write_supported"])
        self.assertFalse(report["execution_authorized"])
        self.assertFalse(report["schema_validated"])
        if report["issues"]:
            self.assertEqual(report["status"], "blocked")
            self.assertIsNone(report["preview_sha256"])
            self.assertEqual(report["state_delta"], [])
        else:
            self.assertEqual(report["status"], "ready_for_review")
            self.assertTrue(report["coverage_complete"])
            self.assertEqual(len(report["preview_sha256"]), 64)
        return report

    def ready(self, **kwargs):
        report = self.run_preview(**kwargs)
        self.assertEqual(report["issues"], [])
        return report

    def blocked(self, text, **kwargs):
        report = self.run_preview(**kwargs)
        self.assertIn(text, "\n".join(report["issues"]))
        return report

    def test_same_language_rechecks_real_primary_without_reaccepting(self):
        report = self.ready()
        primary = report["stages"][0]
        self.assertEqual(primary["evidence_status"], "accepted_rechecked")
        self.assertEqual(primary["action"], "retain_candidate")
        self.assertEqual(report["effects"]["retired_stages"], [])
        self.assertIn("问题一求解/问题一求解结果.xlsx", report["expected_file_hashes"])
        self.assertIsNone(report["expected_file_hashes"]["问题一求解/q1_solver.m"])

    def test_target_change_withdraws_current_binding_but_preserves_original_files(self):
        report = self.ready(target="matlab")
        removed = {tuple(row["path"]) for row in report["state_delta"] if row["operation"] == "remove"}
        for key in ("code", "primary_code_sha256", "solution_workbook", "validated_data_hash"):
            self.assertIn(("subproblems", "Q1", key), removed)
        self.assertIn(("subproblems", "Q1", "solver_execution", "primary", "validated_bundle_sha256"), removed)
        self.assertNotIn(("subproblems", "Q1", "selected_model"), removed)

    def test_deterministic_digest_binds_target_reason_and_raw_state(self):
        first = self.ready()
        self.assertEqual(first, self.ready())
        self.assertNotEqual(first["preview_sha256"], self.ready(target="matlab")["preview_sha256"])
        self.assertNotEqual(first["preview_sha256"], self.ready(reason="different reviewed requirement")["preview_sha256"])
        with (self.root / transaction.STATE_RELATIVE_PATH).open("a", encoding="utf-8") as handle:
            handle.write("\n# same generation, different raw bytes\n")
        self.assertNotEqual(first["preview_sha256"], self.ready()["preview_sha256"])

    def test_independent_mixed_question_is_not_lost_or_majority_voted(self):
        matlab_stage(self.root, "Q3", "primary")
        report = self.ready()
        self.assertEqual(report["kind"], "legacy_mixed")
        self.assertIsNone(report["candidate_backend"])
        self.assertEqual(report["effects"]["retired_stages"], [{"question": "Q3", "stage": "primary"}])
        self.assertEqual(len(report["stages"]), 4)

    def test_analysis_only_retirement_keeps_valid_primary(self):
        matlab_stage(self.root)
        report = self.ready()
        self.assertEqual(report["effects"]["retired_stages"], [{"question": "Q1", "stage": "analysis"}])
        changed = {tuple(row["path"]) for row in report["state_delta"]}
        self.assertNotIn(("subproblems", "Q1", "code"), changed)
        self.assertNotIn(("subproblems", "Q1", "primary_execution_status"), changed)

    def test_primary_retirement_captures_analysis_before_gate_reason_is_cleared(self):
        matlab_stage(self.root)
        report = self.ready(target="matlab")
        self.assertEqual({item["stage"] for item in report["effects"]["retired_stages"]}, {"primary", "analysis"})
        self.assertEqual(report["stages"][1]["action"], "rebuild_dependency")

    def test_both_same_language_accepted_stages_are_rechecked(self):
        synthetic_python_analysis(self.root)
        report = self.ready()
        self.assertEqual([item["evidence_status"] for item in report["stages"]], ["accepted_rechecked"] * 2)
        self.assertEqual(report["effects"]["retired_stages"], [])

    def test_not_required_old_analysis_is_unbound_not_activated(self):
        synthetic_python_analysis(self.root)
        self.change(lambda s: s["subproblems"]["Q1"].update(result_analysis_status="not_required", result_analysis_requirement_reason="No alternative-world claim"))
        report = self.ready()
        self.assertFalse(report["stages"][1]["active"])
        self.assertEqual(report["stages"][1]["action"], "historical_inactive")
        self.assertEqual(report["effects"]["retired_stages"], [])
        self.assertFalse(any(row["path"][-1] == "result_analysis_status" for row in report["state_delta"]))

    def test_unregistered_opposite_code_does_not_choose_or_gain_acceptance(self):
        (self.root / "问题一求解/q1_solver.m").write_text("historical unregistered source", encoding="utf-8")
        report = self.ready()
        self.assertIn("问题一求解/q1_solver.m", report["observed_unbound_files"])
        self.assertEqual(report["effects"]["retired_stages"], [])

    def test_unregistered_analysis_workbook_does_not_activate_gate(self):
        (self.root / "问题一求解/问题一结果深化分析.xlsx").write_bytes(b"historical opaque bytes")
        report = self.ready()
        self.assertEqual(report["stages"][1]["action"], "not_activated")
        self.assertIn("问题一求解/问题一结果深化分析.xlsx", report["observed_unbound_files"])

    def test_empty_stage_selectors_are_diagnostic_not_accepted(self):
        self.change(lambda s: s["subproblems"].update(Q2={"solver_execution": {"primary": {}}}))
        report = self.ready()
        q2 = [row for row in report["stages"] if row["question"] == "Q2"]
        self.assertEqual(q2[0]["action"], "planned")

    def test_accepted_without_binding_is_blocked(self):
        self.change(lambda s: s["subproblems"].update(Q2={"primary_execution_status": "accepted"}))
        self.blocked("accepted stage has no source bindings")

    def test_missing_registered_source_is_blocked(self):
        state = load(self.root)
        (self.root / state["subproblems"]["Q1"]["code"]).unlink()
        self.blocked("declared source is missing")

    def test_same_generation_mid_read_mutation_blocks_proposal(self):
        original = stage_code.parse_stage_config
        def change(*args, **kwargs):
            with (self.root / transaction.STATE_RELATIVE_PATH).open("a", encoding="utf-8") as handle:
                handle.write("\n# concurrent manual edit\n")
            return original(*args, **kwargs)
        with patch.object(stage_code, "parse_stage_config", side_effect=change):
            report = preview.preview_migration(self.root, target_backend="python", reason="test")
        self.assertEqual(report["status"], "blocked")
        self.assertIsNone(report["preview_sha256"])
        self.assertEqual(report["state_delta"], [])
        self.assertTrue(any("ReadSetConflictError" in issue for issue in report["issues"]))

    def test_source_change_after_validation_is_detected(self):
        original = receipts.validate_one
        def change(*args, **kwargs):
            result = original(*args, **kwargs)
            with (self.root / "input.json").open("a", encoding="utf-8") as handle:
                handle.write(" ")
            return result
        with patch.object(receipts, "validate_one", side_effect=change):
            report = preview.preview_migration(self.root, target_backend="python", reason="test")
        self.assertEqual(report["status"], "blocked")
        self.assertIsNone(report["preview_sha256"])

    def test_journal_is_not_implicitly_recovered(self):
        journal = self.root / transaction.JOURNAL_RELATIVE_PATH
        journal.write_text("{}", encoding="utf-8")
        before = files(self.root)
        with self.assertRaises(preview.ProjectStateReadError):
            preview.preview_migration(self.root, target_backend="python", reason="test")
        self.assertEqual(files(self.root), before)

    def test_input_byte_drift_hits_actual_input_check(self):
        with (self.root / "input.json").open("a", encoding="utf-8") as handle:
            handle.write(" ")
        self.blocked("当前实际输入data_sha256")

    def test_entry_byte_drift_hits_original_source_binding(self):
        with (self.root / "问题一求解/问题一求解.py").open("a", encoding="utf-8") as handle:
            handle.write("\n# changed after delivery\n")
        self.blocked("primary入口SHA-256与已交付代码不一致")

    def test_receipt_backend_conflict_cannot_be_fixed_by_changing_target(self):
        update_receipt(self.root, "问题一求解/问题一求解结果.xlsx", {"solver_backend": "matlab"})
        self.blocked("RUN_RECEIPT.solver_backend", target="matlab")

    def test_receipt_bundle_conflict_is_not_accepted(self):
        update_receipt(self.root, "问题一求解/问题一求解结果.xlsx", {"code_bundle_sha256": "e" * 64})
        self.blocked("receipt bundle differs")

    def test_missing_validated_workbook_digest_is_not_preserved(self):
        self.change(lambda s: s["subproblems"]["Q1"]["validated_artifact_hashes"].pop("solution_workbook"))
        self.blocked("lacks its validated raw digest")

    def test_missing_validated_bundle_hits_existing_identity_gate(self):
        self.change(lambda s: s["subproblems"]["Q1"]["solver_execution"]["primary"].pop("validated_bundle_sha256"))
        self.blocked("bundle")

    def test_stale_acceptance_is_retired_not_promoted(self):
        self.change(lambda s: s["subproblems"]["Q1"].update(artifacts_stale=True, stale_layers=["solution_workbook"]))
        report = self.ready()
        self.assertEqual(report["stages"][0]["action"], "rebuild_stale")
        self.assertNotEqual(report["stages"][0]["evidence_status"], "accepted_rechecked")

    def test_figure_only_stale_does_not_withdraw_primary(self):
        self.change(lambda s: s["subproblems"]["Q1"].update(artifacts_stale=True, stale_layers=["figure_bundle"]))
        self.assertEqual(self.ready()["stages"][0]["action"], "retain_candidate")

    def test_workbook_tamper_cannot_hide_behind_matching_state_hash(self):
        wb = self.root / "问题一求解/问题一求解结果.xlsx"
        book = openpyxl.load_workbook(wb)
        sheet = book["主结果质量门"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("是否通过") + 1, False)
        book.save(wb); book.close()
        def resave(s):
            for name in ("artifact_hashes", "validated_artifact_hashes"):
                s["subproblems"]["Q1"][name]["solution_workbook"] = file_hash(wb)
        self.change(resave)
        self.blocked("主结果质量门")

    def test_plain_invalid_xlsx_is_blocking_not_successful_absence(self):
        (self.root / "问题一求解/问题一求解结果.xlsx").write_bytes(b"not a workbook")
        self.blocked("BadZipFile")

    def test_canonical_declaration_adapter_is_private_and_not_schema_admission(self):
        def convert(s):
            s["execution"] = {"solver_backend": "python", "solver_backend_selection_reason": "root choice"}
            record = s["subproblems"]["Q1"]["solver_execution"]["primary"]
            record.pop("backend"); record.pop("selection_reason")
        self.change(convert)
        report = self.ready()
        self.assertEqual(report["kind"], "canonical_declarations")
        self.assertEqual(report["stages"][0]["evidence_status"], "accepted_rechecked")
        self.assertNotIn("backend", load(self.root)["subproblems"]["Q1"]["solver_execution"]["primary"])

    def test_canonical_and_stage_selectors_conflict_without_side_effect(self):
        self.change(lambda s: s.update(execution={"solver_backend": "python", "solver_backend_selection_reason": "root choice"}))
        self.blocked("coexist")

    def test_invalid_hash_and_mapping_shapes_block_before_traversal(self):
        for bad in ([], {"solution_workbook": "wrong"}):
            with self.subTest(bad=bad):
                self.change(lambda s: s["subproblems"]["Q1"].update(validated_artifact_hashes=bad))
                self.blocked("validated_artifact_hashes")

    def test_unsupported_question_is_reported_not_silently_skipped(self):
        self.change(lambda s: s["subproblems"].update(Q11={"status": "designed"}))
        self.blocked("unsupported question identity")

    def test_absent_dependency_source_is_blocking(self):
        self.change(lambda s: s["subproblems"]["Q1"].update(depends_on=[{"question": "Q99", "kind": "data"}]))
        self.blocked("absent or malformed source")

    def test_cross_question_actual_workbook_requires_numeric_dependency(self):
        python_primary(self.root, "问题二", True)
        self.blocked("lacks an explicit numerical")
        self.change(lambda s: s["subproblems"]["Q2"].update(depends_on=[{"question": "Q1", "kind": "model"}]))
        self.blocked("lacks an explicit numerical")
        self.change(lambda s: s["subproblems"]["Q2"].update(depends_on=[{"question": "Q1", "kind": "result"}]))
        self.ready()

    def test_typed_numeric_dependents_and_independent_question_are_distinguished(self):
        matlab_stage(self.root, "Q2", "primary")
        matlab_stage(self.root, "Q3", "primary")
        self.change(lambda s: s["subproblems"]["Q2"].update(depends_on=[{"question": "Q1", "kind": "data"}]))
        report = self.ready(target="matlab")
        retired = {item["question"] for item in report["effects"]["retired_stages"]}
        self.assertEqual(retired, {"Q1", "Q2"})
        q3 = next(row for row in report["stages"] if row["question"] == "Q3" and row["stage"] == "primary")
        self.assertEqual(q3["action"], "retain_candidate")

    def test_model_only_dependency_is_not_numerical_retirement(self):
        matlab_stage(self.root, "Q2", "primary")
        self.change(lambda s: s["subproblems"]["Q2"].update(depends_on=[{"question": "Q1", "kind": "model"}]))
        self.assertEqual(self.ready(target="matlab")["effects"]["retired_stages"], [{"question": "Q1", "stage": "primary"}])

    def test_chained_parameter_data_dependencies_reach_fixed_point(self):
        matlab_stage(self.root, "Q2", "primary")
        matlab_stage(self.root, "Q3", "primary")
        def dependencies(s):
            s["subproblems"]["Q2"]["depends_on"] = [{"question": "Q1", "kind": "result"}]
            s["subproblems"]["Q3"]["depends_on"] = [{"question": "Q2", "kind": "parameter"}]
        self.change(dependencies)
        self.assertEqual({r["question"] for r in self.ready(target="matlab")["effects"]["retired_stages"]}, {"Q1", "Q2", "Q3"})

    def test_cycle_and_diamond_terminate_deterministically(self):
        for q in ("Q2", "Q3", "Q4"):
            matlab_stage(self.root, q, "primary")
        def dependencies(s):
            graph = {"Q1": ["Q4"], "Q2": ["Q1"], "Q3": ["Q1"], "Q4": ["Q2", "Q3"]}
            for q, sources in graph.items():
                s["subproblems"][q]["depends_on"] = [{"question": x, "kind": "parameter"} for x in sources]
        self.change(dependencies)
        first = self.ready(target="matlab")
        self.assertEqual(first, self.ready(target="matlab"))
        self.assertEqual(len(first["effects"]["retired_stages"]), 4)
        self.assertTrue(first["effects"]["dependency_cycles"])

    def test_legacy_untyped_edge_keeps_conservative_model_review(self):
        matlab_stage(self.root, "Q2", "primary")
        self.change(lambda s: s["subproblems"]["Q2"].update(depends_on=["Q1"], human_model_approval_status="approved"))
        report = self.ready(target="matlab")
        self.assertTrue(any(row["path"] == ["subproblems", "Q2", "human_model_approval_status"] and row.get("after") == "stale"
                            for row in report["state_delta"]))

    def test_paper_fragment_closure_preserves_unrelated_mechanism(self):
        self.change(lambda s: s.update(paper_framework={"version": "v0.8-project-memory", "paper_fragments": [
            {"id": "results", "scope": "Q1", "status": "current"},
            {"id": "abstract", "depends_on": ["results"], "status": "current"},
            {"id": "mechanism", "depends_on": [], "status": "current"}]}))
        report = self.ready(target="matlab")
        self.assertEqual(report["effects"]["stale_paper_fragments"], ["abstract", "results"])

    def test_duplicate_fragment_ids_fail_instead_of_silent_overwrite(self):
        self.change(lambda s: s.update(paper_framework={"paper_fragments": [{"id": "a"}, {"id": "a"}]}))
        self.blocked("ids must be unique")

    def test_pure_candidate_does_not_mutate_input_and_does_not_add_a_generation(self):
        state = load(self.root); old = deepcopy(state)
        report = self.ready(target="matlab")
        contract = yaml.safe_load((ROOT / "core/state_transition_contract.yaml").read_text(encoding="utf-8"))
        candidate, _ = preview._retirement_candidate(state, report["stages"], "matlab", "reason", contract)
        self.assertEqual(state, old)
        self.assertNotIn("state_generation", candidate["project"])
        self.assertEqual(candidate["execution"]["solver_backend"], "matlab")
        self.assertEqual(candidate["subproblems"]["Q1"]["primary_execution_status"], "pending")

    def test_preview_does_not_call_any_transaction_writer(self):
        with patch.object(transaction, "commit_project_state", side_effect=AssertionError("writer called")), \
             patch.object(transaction, "prepare_history_archive", side_effect=AssertionError("archive called")):
            self.ready(target="matlab")

    def test_relative_escape_is_rejected(self):
        self.change(lambda s: s["subproblems"]["Q1"].update(code="../evil.py"))
        self.blocked("StageCodeError")

    @unittest.skipIf(os.name == "nt", "Windows symlink creation may require elevated privileges")
    def test_symlink_source_is_rejected(self):
        source = self.root / "问题一求解/问题一求解.py"
        backup = self.root / "entry_backup.py"
        source.rename(backup); source.symlink_to(backup)
        self.blocked("符号链接")

    def test_hardlink_source_is_rejected(self):
        source = self.root / "问题一求解/问题一求解.py"
        try:
            os.link(source, self.root / "entry_backup.py")
        except OSError as exc:
            self.skipTest(f"hard links unavailable: {exc}")
        self.blocked("independent regular file")

    def test_cli_preview_is_explicit_and_read_only(self):
        before = files(self.root)
        process = subprocess.run([sys.executable, str(ROOT / "scripts/project_solver_backend.py"), "inspect",
            "--project-root", str(self.root), "--migration-target", "matlab", "--reason", "统一"],
            capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(process.returncode, 0, process.stderr)
        report = yaml.safe_load(process.stdout)
        self.assertEqual(report["status"], "ready_for_review")
        self.assertEqual(files(self.root), before)

    def test_cli_rejects_ambiguous_or_incomplete_preview_arguments(self):
        base = [sys.executable, str(ROOT / "scripts/project_solver_backend.py"), "inspect", "--project-root", str(self.root)]
        for flags in (["--migration-target", "matlab"], ["--reason", "test"],
                      ["--migration-target", "matlab", "--reason", "test", "--requested-backend", "python"]):
            with self.subTest(flags=flags):
                process = subprocess.run([*base, *flags], capture_output=True, check=False)
                self.assertEqual(process.returncode, 2)

    def test_invalid_explicit_target_or_reason_is_rejected(self):
        for target, reason in (("auto", "a"), (None, "a"), ("python", " \t"), ("python", None)):
            with self.assertRaises(ValueError):
                preview.preview_migration(self.root, target_backend=target, reason=reason)

    def test_delivered_execution_without_bindings_is_not_a_planned_empty_stage(self):
        for status in ("code_delivered", "awaiting_user_execution", "workbook_received", "rejected"):
            with self.subTest(status=status):
                self.change(lambda state: state["subproblems"].update(Q2={"primary_execution_status": status}))
                self.blocked("execution stage has no source bindings")

    def test_missing_planned_path_with_delivered_status_is_a_conflict(self):
        self.change(lambda state: state["subproblems"].update(Q2={
            "code": "问题二求解/问题二求解.py", "primary_execution_status": "code_delivered"}))
        self.blocked("declared source is missing")

    def test_hex_case_does_not_make_validated_data_a_different_identity(self):
        self.change(lambda state: state["subproblems"]["Q1"].update(
            validated_data_hash=state["subproblems"]["Q1"]["validated_data_hash"].upper()))
        self.assertEqual(self.ready()["stages"][0]["evidence_status"], "accepted_rechecked")

    def rebind_reader_fixture(self, *, config_update=None, prefix="", helpers=(), receipt_update=None):
        """Rebind controlled synthetic reader evidence, not a claimed solver rerun."""
        state = load(self.root)
        entry = state["subproblems"]["Q1"]
        path = self.root / entry["code"]
        _, config = stage_code.parse_stage_config(path)
        old = repr(config)
        config.update(config_update or {})
        if helpers:
            config["code_dependencies"] = [{"path": item, "sha256": file_hash(self.root / item)} for item in helpers]
        source = path.read_text(encoding="utf-8")
        self.assertIn("RUN_CONFIG = " + old, source)
        path.write_text(prefix + source.replace("RUN_CONFIG = " + old, "RUN_CONFIG = " + repr(config), 1), encoding="utf-8")
        entry["primary_code_sha256"] = file_hash(path)
        bundle = reference_digest(self.root, [entry["code"], *helpers])
        entry["solver_execution"]["primary"].update(bundle_sha256=bundle, validated_bundle_sha256=bundle)
        update_receipt(self.root, entry["solution_workbook"], {"code_sha256": file_hash(path), "code_bundle_sha256": bundle,
                                                              **(receipt_update or {})})
        for name in ("artifact_hashes", "validated_artifact_hashes"):
            entry.setdefault(name, {}).update(primary_code=file_hash(path), solution_workbook=file_hash(self.root / entry["solution_workbook"]))
        save_state(self.root, state)

    def test_declared_helper_is_part_of_the_bound_read_set(self):
        helper = "问题一求解/helper.py"
        (self.root / helper).write_text("VALUE = 2\n", encoding="utf-8")
        self.rebind_reader_fixture(prefix="import helper\n", helpers=[helper])
        report = self.ready()
        self.assertEqual(report["expected_file_hashes"][helper], file_hash(self.root / helper))

    def test_changed_helper_is_rejected_by_original_fingerprint(self):
        helper = "问题一求解/helper.py"
        (self.root / helper).write_text("VALUE = 2\n", encoding="utf-8")
        self.rebind_reader_fixture(prefix="import helper\n", helpers=[helper])
        (self.root / helper).write_text("VALUE = 3\n", encoding="utf-8")
        self.blocked("源码依赖SHA-256不一致")

    def test_missing_helper_is_not_treated_as_external_package(self):
        helper = "问题一求解/helper.py"
        (self.root / helper).write_text("VALUE = 2\n", encoding="utf-8")
        self.rebind_reader_fixture(prefix="import helper\n", helpers=[helper])
        (self.root / helper).unlink()
        self.blocked("源码文件不存在")

    def test_undeclared_project_helper_reaches_original_closure_gate(self):
        (self.root / "问题一求解/helper.py").write_text("VALUE = 2\n", encoding="utf-8")
        self.rebind_reader_fixture(prefix="import helper\n")
        self.blocked("未声明")

    def test_modern_metadata_cannot_hide_behind_a_legacy_config_marker(self):
        self.rebind_reader_fixture(config_update={"run_receipt_protocol_version": "1.0.0"})
        self.blocked("modern binding cannot downgrade")

    def test_genuine_legacy_reader_fixture_requires_rebuild_even_same_language(self):
        self.rebind_reader_fixture(config_update={"run_receipt_protocol_version": "1.0.0"},
                                   receipt_update={"run_receipt_version": "1.0.0"})
        self.change(lambda s: s["subproblems"]["Q1"].pop("solver_execution"))
        report = self.ready()
        self.assertEqual(report["stages"][0]["action"], "rebuild_legacy")
        self.assertNotEqual(report["stages"][0]["evidence_status"], "accepted_rechecked")

    def test_config_backend_conflict_cannot_be_normalized_into_target(self):
        self.rebind_reader_fixture(config_update={"solver_backend": "matlab"})
        self.blocked("RUN_CONFIG.solver_backend conflicts", target="matlab")

    def test_independent_numerical_recheck_is_not_replaced_by_true_booleans(self):
        wb = self.root / "问题一求解/问题一求解结果.xlsx"
        book = openpyxl.load_workbook(wb)
        sheet = book["均衡残差"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("残差") + 1, 0.1)
        book.save(wb); book.close()
        def bind(s):
            for name in ("artifact_hashes", "validated_artifact_hashes"):
                s["subproblems"]["Q1"][name]["solution_workbook"] = file_hash(wb)
        self.change(bind)
        report = self.blocked("残差")
        self.assertTrue(any("复核" in issue or "容差" in issue or "不通过" in issue for issue in report["issues"]))

    def test_registered_framework_is_in_the_read_set_and_missing_is_blocking(self):
        self.change(lambda s: s.update(paper_framework={"path": "模型论文框架.md"}))
        self.blocked("declared source is missing")
        (self.root / "模型论文框架.md").write_text("# Current project memory\n", encoding="utf-8")
        self.assertEqual(self.ready()["expected_file_hashes"]["模型论文框架.md"], file_hash(self.root / "模型论文框架.md"))

    def test_historical_archive_cannot_become_a_current_solver_path(self):
        self.change(lambda s: s["subproblems"]["Q1"].update(code="state/backend_history/past/问题一求解.py"))
        self.blocked("historical archive cannot be a current source")

    def test_missing_project_state_is_not_an_implicitly_selected_project(self):
        (self.root / transaction.STATE_RELATIVE_PATH).unlink()
        self.blocked("requires an existing project state")

    def test_empty_explicit_project_does_not_gain_execution_eligibility(self):
        save_state(self.root, {"project": {"state_generation": 0}, "subproblems": {}})
        report = self.ready(target="matlab")
        self.assertIsNone(report["candidate_backend"])
        self.assertEqual(report["stages"], [])
        self.assertFalse(report["execution_authorized"])

    def test_preview_digest_is_reproducible_from_its_non_authorizing_payload(self):
        report = self.ready()
        payload = deepcopy(report)
        payload["preview_sha256"] = None
        digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True,
                                          separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()
        self.assertEqual(report["preview_sha256"], digest)


if __name__ == "__main__":
    unittest.main()
