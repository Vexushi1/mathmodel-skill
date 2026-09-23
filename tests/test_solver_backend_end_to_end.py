"""Prepare synthetic fixtures and verify outputs of an actual MATLAB process.

Ordinary unittest discovery checks fixture instantiation only. Native-output tests
run only when HSK_MATLAB_SOLVER_OUTPUT points to a completed MATLAB smoke run.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from templates.code.hsk_pipeline import result_io
QUESTION = "问题一求解"
CONFIG = re.compile(r"RUN_CONFIG = jsondecode\('([^\n]*)'\);")


def load_module(name: str, relative: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reference_digest(root: Path, paths: list[str]) -> str:
    """Independent protocol reference used to compare MATLAB's byte-level output."""
    digest = hashlib.sha256()
    for relative in sorted(paths, key=lambda item: item.encode("utf-8")):
        digest.update(relative.encode("utf-8") + b"\0")
        digest.update(bytes.fromhex(file_hash(root / relative)))
    return digest.hexdigest()


def instantiate(root: Path, stage: str, **updates) -> Path:
    name = "q1_solver.m" if stage == "primary" else "q1_analysis.m"
    source = (ROOT / "templates/code/matlab" / name).read_text(encoding="utf-8")
    match = CONFIG.search(source)
    assert match is not None
    config = json.loads(match.group(1).replace("''", "'"))
    config.update(updates)
    if "data_sha256" not in updates:
        config["data_sha256"] = reference_digest(root, config["data_paths"])
    literal = json.dumps(config, ensure_ascii=False, separators=(",", ":")).replace("'", "''")
    source = source[:match.start(1)] + literal + source[match.end(1):]
    path = root / QUESTION / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(source, encoding="utf-8", newline="\n")
    return path


def stage_state(root: Path, stage: str, code: Path) -> dict:
    return {"bundle_sha256": reference_digest(root, [code.relative_to(root).as_posix()])}


def save_state(root: Path, state: dict) -> None:
    (root / "state").mkdir(exist_ok=True)
    (root / "state/project_state.yaml").write_text(
        yaml.safe_dump(state, allow_unicode=True, sort_keys=False), encoding="utf-8")


def prepare(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    if (root / QUESTION / "问题一求解结果.xlsx").exists():
        raise ValueError("Use a fresh output directory; existing numerical results are not overwritten by preparation.")
    shutil.copyfile(ROOT / "tests/fixtures/solver_backends/input.json", root / "input.json")
    (root / "synthetic_fixture.json").write_text(
        json.dumps({"fixture": "solver_backend_native_v1", "purpose": "repository_test"}), encoding="utf-8")
    code = instantiate(root, "primary")
    state = {"project": {"competition": "test", "problem": "synthetic", "current_phase": "solve_validate"},
             "preprocessing": {"decision": "not_needed", "status": "not_applicable", "quality_status": "not_applicable"},
             "execution": {"solver_backend": "matlab",
                           "solver_backend_selection_reason": "Synthetic whole-project native workbook integration"},
             "subproblems": {"Q1": {
                 "status": "designed", "selected_model": "a*x=b", "capabilities": {"requires_equilibrium_residual": True},
                 "code": code.relative_to(root).as_posix(), "primary_code_sha256": file_hash(code),
                 "data_hash": reference_digest(root, ["input.json"]), "result_quality_status": "pending",
                 "result_analysis_status": "pending", "primary_execution_status": "awaiting_user_execution",
                 "solver_execution": {"primary": stage_state(root, "primary", code)},
             }}}
    save_state(root, state)
    writer = (ROOT / "templates/code/matlab/q1_solver.m").read_text(encoding="utf-8")
    writer = writer[writer.index("function write_workbook("):]
    probe = '''function native_writer_probe(filename,mode)
cells = {'记录键','数值','标记'; '0001',1.5,true; '9007199254740993',"",false};
if mode == "duplicate_header", cells{1,2} = '记录键'; end
if mode == "duplicate_key", cells{3,1} = '0001'; end
if mode == "infinite", cells{2,2} = Inf; end
sheets = {'类型往返',cells};
if mode == "duplicate_sheet", sheets = {'Probe',cells;'probe',cells}; end
if mode == "empty_sheet", sheets = {'',cells}; end
write_workbook(filename,sheets);
end

'''
    (root / "native_writer_probe.m").write_text(probe + writer, encoding="utf-8", newline="\n")
    source = (ROOT / "templates/code/matlab/q1_solver.m").read_text(encoding="utf-8")
    hashes = source[source.index("function path = project_path("):source.index("function receipt = make_receipt(")]
    probe = '''function native_hash_probe(root)
root = fullfile(string(root),"hash-probe");
result = files_digest(root,["\ue000.txt";"\U0001f600.txt";"a.txt"]);
fid = fopen(fullfile(root,"digest.txt"),'w');
assert(fid >= 0);
cleanup = onCleanup(@() fclose(fid));
fwrite(fid,result,'char');
end

'''
    (root / "native_hash_probe.m").write_text(probe + hashes, encoding="utf-8", newline="\n")
    (root / "hash-probe").mkdir()
    for name in ("\ue000.txt", "\U0001f600.txt", "a.txt"):
        (root / "hash-probe" / name).write_text(name, encoding="utf-8")
    for case in ("input_hash", "missing_input", "nonfinite_result"):
        destination = root / "negative" / case
        destination.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(root / "input.json", destination / "input.json")
        if case == "nonfinite_result":
            payload = {"coefficient": 1e-308, "right_hand_side": 1e308, "sensitivity_coefficients": [1.5, 2, 3]}
            (destination / "input.json").write_text(json.dumps(payload), encoding="utf-8")
        instantiate(destination, "primary")
        if case == "input_hash":
            with (destination / "input.json").open("a", encoding="utf-8") as handle:
                handle.write(" ")
        elif case == "missing_input":
            (destination / "input.json").unlink()


def prepare_preprocessed(root: Path) -> None:
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    fixture = load_module("native_preprocessing_fixture", "tests/test_solver_backend_preprocessing.py")
    state, digest = fixture.prepare_native_preprocessing(root)
    (root / "synthetic_fixture.json").write_text(json.dumps({
        "fixture": "solver_backend_native_v1", "purpose": "repository_test"}), encoding="utf-8")
    (root / "preprocessing_native_fixture.json").write_text(json.dumps({
        "preprocessing": state["preprocessing"], "workbook_sha256": digest,
        "upstream_is_synthetic_receipt_fixture": True, "python_preprocessing_executed": False,
    }, ensure_ascii=False), encoding="utf-8")
    code = instantiate(root, "primary", data_paths=["数据预处理/数据预处理结果.xlsx"],
                       data_sha256=digest, data_identity_mode="preprocessing_workbook")
    state["project"]["current_phase"] = "solve_validate"
    state["execution"] = {"solver_backend": "matlab",
                          "solver_backend_selection_reason": "Synthetic whole-project native workbook integration"}
    state["subproblems"]["Q1"].update(
        selected_model="a*x=b", capabilities={"requires_equilibrium_residual": True},
        code=code.relative_to(root).as_posix(), primary_code_sha256=file_hash(code), data_hash=digest,
        result_quality_status="pending", result_analysis_status="pending",
        primary_execution_status="awaiting_user_execution",
        solver_execution={"primary": stage_state(root, "primary", code)})
    save_state(root, state)


def verify(root: Path, stage: str) -> None:
    runtime_report = json.loads((root / f"{stage}_matlab_report.json").read_text(encoding="utf-8"))
    assert runtime_report["actual_matlab_execution"] is True and runtime_report["status"] == "passed"
    assert runtime_report["matlab_release"]
    receipt = load_module("native_smoke_receipt", "scripts/validate_user_execution.py")
    io = result_io
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    workbook = root / QUESTION / ("问题一求解结果.xlsx" if stage == "primary" else "问题一结果深化分析.xlsx")
    code = root / entry["code" if stage == "primary" else "result_analysis_code"]
    config, issues = receipt.configuration_map(workbook)
    if issues:
        raise AssertionError(issues)
    assert config["solver_backend"] == "matlab"
    assert config["run_receipt_version"] == "1.1.0"
    assert config["code_sha256"] == file_hash(code)
    assert config["code_bundle_sha256"] == reference_digest(root, [code.relative_to(root).as_posix()])
    io.validate_workbook_file(workbook, "solution" if stage == "primary" else "result_analysis",
                              capabilities=entry["capabilities"] if stage == "primary" else {})
    issues = receipt.validate_one(root, workbook, state, True)
    if issues:
        raise AssertionError(issues)
    book = openpyxl.load_workbook(workbook, data_only=True)
    try:
        if stage == "primary":
            values = {row[0]: row[1] for row in book["核心指标"].iter_rows(min_row=2, values_only=True)}
            assert abs(values["解"] - 3) < 1e-12
            assert values["绝对方程残差"] <= 1e-10
            assert book["状态明细"]["A2"].value == "0001"
        else:
            actual = [row[3] for row in book["参数敏感性"].iter_rows(min_row=2, values_only=True)]
            assert actual == [4, 3, 2]
    finally:
        book.close()
    if stage == "primary" and runtime_report.get("type_probe") == "text_keys_logicals_missing_rows_and_rejections_passed":
        digest_root = root / "hash-probe"
        assert (digest_root / "digest.txt").read_text(encoding="utf-8") == reference_digest(
            digest_root, ["\ue000.txt", "\U0001f600.txt", "a.txt"])
        book = openpyxl.load_workbook(root / "type-probe" / QUESTION / "问题一求解结果.xlsx", data_only=True)
        try:
            rows = list(book["类型往返"].iter_rows(values_only=True))
            assert rows[1] == ("0001", 1.5, True)
            assert rows[2] == ("9007199254740993", None, False)
        finally:
            book.close()
    if (root / "preprocessing_native_fixture.json").exists():
        baseline = json.loads((root / "preprocessing_native_fixture.json").read_text(encoding="utf-8"))
        assert state["preprocessing"] == baseline["preprocessing"]
        assert config["data_sha256"] == baseline["workbook_sha256"]
        assert file_hash(root / "数据预处理/数据预处理结果.xlsx") == baseline["workbook_sha256"]
        assert runtime_report["preprocessing_workbook_unchanged"]
    save_state(root, state)
    if stage == "analysis":
        # Exercise read-only migration preflight against the actual native receipts.
        # This is not execution/confirmation of a backend migration.
        inspector = load_module("native_migration_preview", "scripts/project_solver_backend.py")
        before = {p.relative_to(root).as_posix(): file_hash(p) for p in root.rglob("*") if p.is_file()}
        same = inspector.preview_migration(root, target_backend="matlab", reason="Native same-backend retention check")
        assert same["status"] == "ready_for_review", same["issues"]
        assert all(row["evidence_status"] == "accepted_rechecked" for row in same["stages"]), same["stages"]
        assert same["effects"]["retired_stages"] == [], same["effects"]
        other = inspector.preview_migration(root, target_backend="python", reason="Native retirement proposal check")
        assert other["status"] == "ready_for_review", other["issues"]
        assert {(row["question"], row["stage"]) for row in other["effects"]["retired_stages"]} == {
            ("Q1", "primary"), ("Q1", "analysis")}, other["effects"]
        assert not other["migration_authorized"] and not other["write_supported"]
        assert {p.relative_to(root).as_posix(): file_hash(p) for p in root.rglob("*") if p.is_file()} == before


def prepare_analysis(root: Path) -> None:
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    assert entry["primary_execution_status"] == "accepted"
    primary = root / QUESTION / "问题一求解结果.xlsx"
    updates = {"primary_workbook_sha256": file_hash(primary)}
    if state["preprocessing"]["decision"] == "project_level":
        updates.update(data_identity_mode="preprocessing_workbook", data_paths=["数据预处理/数据预处理结果.xlsx"],
                       data_sha256=file_hash(root / "数据预处理/数据预处理结果.xlsx"))
    code = instantiate(root, "analysis", **updates)
    entry.update(result_analysis_requirement_reason="Test the declared coefficient range", analysis_methods=["参数敏感性"],
                 result_analysis_code=code.relative_to(root).as_posix(), analysis_code_sha256=file_hash(code),
                 analysis_execution_status="awaiting_user_execution")
    entry["solver_execution"]["analysis"] = stage_state(root, "analysis", code)
    save_state(root, state)


def deliver(root: Path, stage: str, matlab_command: str) -> None:
    marker = json.loads((root / "synthetic_fixture.json").read_text(encoding="utf-8"))
    assert marker["fixture"] == "solver_backend_native_v1" and marker["purpose"] == "repository_test"
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    assert entry.get(f"{stage}_execution_status") != "accepted"
    fields = ("code", "primary_code_sha256") if stage == "primary" else ("result_analysis_code", "analysis_code_sha256")
    for field in fields:
        entry.pop(field, None)
    selection = entry["solver_execution"][stage]
    selection.pop("bundle_sha256", None)
    selection.pop("validated_bundle_sha256", None)
    entry[f"{stage}_execution_status"] = "pending"
    save_state(root, state)
    script = f"{QUESTION}/q1_{'solver' if stage == 'primary' else 'analysis'}.m"
    process = subprocess.run([sys.executable, str(ROOT / "scripts/validate_code_delivery.py"), str(root),
                              "--script", script, "--stage", stage, "--write", "--strict",
                              "--matlab-command", matlab_command], capture_output=True, encoding="utf-8",
                             env={**os.environ, "PYTHONUTF8": "1"})
    (root / f"{stage}_delivery_report.yaml").write_text(process.stdout + process.stderr, encoding="utf-8")
    assert process.returncode == 0, process.stdout + process.stderr
    report = yaml.safe_load(process.stdout)
    assert report["status"] == "passed" and report["task_code_executed"] is False
    assert report["code_quality_metrics"][script]["native_analysis_status"] == "passed"
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    assert entry[f"{stage}_execution_status"] == "awaiting_user_execution"
    assert entry["solver_execution"][stage]["bundle_sha256"] == reference_digest(root, [script])
    assert "validated_bundle_sha256" not in entry["solver_execution"][stage]


class SolverBackendFixtureTests(unittest.TestCase):
    def test_instantiation_binds_real_input_bytes_without_running_matlab(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            prepare(root)
            source = (root / QUESTION / "q1_solver.m").read_text(encoding="utf-8")
            config = json.loads(CONFIG.search(source).group(1))
            self.assertEqual(config["data_sha256"], reference_digest(root, ["input.json"]))
            self.assertEqual(config["solver_backend"], "matlab")
            self.assertFalse(list(root.rglob("*.xlsx")))
            self.assertTrue((root / "negative/nonfinite_result" / QUESTION / "q1_solver.m").is_file())

    def test_analysis_template_does_not_call_primary_or_write_its_workbook(self):
        source = (ROOT / "templates/code/matlab/q1_analysis.m").read_text(encoding="utf-8")
        self.assertNotIn("q1_solver(", source)
        self.assertIn("primary_workbook_sha256", source)
        self.assertIn("workbook ~= primary", source)


@unittest.skipUnless(os.environ.get("HSK_MATLAB_SOLVER_OUTPUT"), "real MATLAB outputs were not supplied")
class NativeMatlabOutputTests(unittest.TestCase):
    def test_native_workbooks_pass_common_python_acceptance(self):
        root = Path(os.environ["HSK_MATLAB_SOLVER_OUTPUT"])
        verify(root, "primary")
        verify(root, "analysis")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in {"prepare", "prepare-preprocessed", "verify-primary", "prepare-analysis", "verify-analysis", "deliver-primary", "deliver-analysis"}:
        parser = argparse.ArgumentParser()
        parser.add_argument("action")
        parser.add_argument("--project", type=Path, required=True)
        parser.add_argument("--matlab-command", default="matlab")
        args = parser.parse_args()
        if args.action == "prepare":
            prepare(args.project.resolve())
        elif args.action == "prepare-preprocessed":
            prepare_preprocessed(args.project.resolve())
        elif args.action == "prepare-analysis":
            prepare_analysis(args.project.resolve())
        elif args.action.startswith("deliver-"):
            deliver(args.project.resolve(), args.action.removeprefix("deliver-"), args.matlab_command)
        else:
            verify(args.project.resolve(), args.action.removeprefix("verify-"))
    else:
        unittest.main()
