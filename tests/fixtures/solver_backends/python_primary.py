"""Synthetic native Python primary; instantiated only by the mixed smoke harness."""
from pathlib import Path
import hashlib
import json
import math
import platform
import random

import openpyxl

RUN_CONFIG = {}


def digest(root, paths):
    result = hashlib.sha256()
    for relative in sorted(paths, key=lambda name: name.encode("utf-8")):
        result.update(relative.encode("utf-8") + b"\0")
        result.update(hashlib.sha256((root / relative).read_bytes()).digest())
    return result.hexdigest()


def main():
    root = Path(__file__).resolve().parents[1]
    code = Path(__file__).resolve()
    config = RUN_CONFIG
    random.seed(config["random_seed"])
    assert digest(root, config["data_paths"]) == config["data_sha256"]
    code_sha = hashlib.sha256(code.read_bytes()).hexdigest()
    bundle_sha = digest(root, [code.relative_to(root).as_posix()])
    if config["input_mode"] == "upstream":
        upstream = openpyxl.load_workbook(root / config["data_paths"][0], data_only=True)
        try:
            values = dict(upstream["核心指标"].iter_rows(min_row=2, values_only=True))
            coefficient, right_hand_side = 2, values["解"]
        finally:
            upstream.close()
    else:
        payload = json.loads((root / config["data_paths"][0]).read_text(encoding="utf-8"))
        coefficient, right_hand_side = payload["coefficient"], payload["right_hand_side"]
    assert coefficient != 0 and all(math.isfinite(v) for v in (coefficient, right_hand_side))
    solution = right_hand_side / coefficient
    residual = abs(coefficient * solution - right_hand_side)
    passed = math.isfinite(solution) and residual <= config["tolerance"]
    receipt = dict(run_receipt_version="1.1.0", primary_quality_protocol_version="1.0.0",
                   execution_owner="user", execution_profile="full_fidelity", stage="primary",
                   problem_name=config["problem_name"], solver_backend="python", code_sha256=code_sha,
                   code_bundle_sha256=bundle_sha, data_sha256=config["data_sha256"],
                   solver=config["solver"], solver_version=platform.python_version(),
                   tolerance=config["tolerance"], iteration_or_time_limit=config["iteration_or_time_limit"],
                   actual_stop_reason="direct_solve_completed", random_seed=config["random_seed"],
                   repetitions_or_scenarios=1, grid_or_time_range="all declared inputs",
                   fallback_used=False, platform=platform.platform())
    receipt.update({name: False for name in (
        "allow_reduced_data", "allow_coarser_grid", "allow_shorter_horizon", "allow_fewer_repetitions",
        "allow_relaxed_tolerance", "allow_silent_solver_fallback")})
    book = openpyxl.Workbook()
    book.remove(book.active)
    sheets = {
        "运行配置": [("项目", "值"), *receipt.items()],
        "核心指标": [("指标", "数值"), ("解", solution), ("绝对方程残差", residual)],
        "数据审计": [("等级", "检查项", "信息", "处理方式"), ("info", "有限输入", "已验证", "只读检查")],
        "主结果质量门": [("Verification ID", "检查项", "是否通过", "证据", "判定关系", "阈值或容差",
                       "实际值", "证据工作表", "阈值来源"),
                    ("PQ-Q1-01", "方程残差", passed, "原系数回代", "abs<=", config["tolerance"],
                     residual, "均衡残差", "solver_tolerance")],
        "均衡残差": [("主体或均衡", "残差", "容差", "是否满足"),
                  ("a*x=b", residual, config["tolerance"], passed)],
        "状态明细": [("记录键", "状态", "数值", "单位"), ("0001", "解", solution, "无量纲")],
    }
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    assert digest(root, config["data_paths"]) == config["data_sha256"]
    assert hashlib.sha256(code.read_bytes()).hexdigest() == code_sha
    assert passed
    book.save(root / config["expected_workbook"])
    book.close()
    print(json.dumps({"actual_python_execution": True, "python_version": platform.python_version(),
                      "answer": solution, "code_sha256": code_sha, "status": "passed"}))


if __name__ == "__main__":
    main()
