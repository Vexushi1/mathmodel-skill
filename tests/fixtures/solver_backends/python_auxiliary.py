"""Synthetic receipt-1.2 fixture; never used to solve a user's contest problem."""
from pathlib import Path, PurePosixPath
import hashlib
import json
import math
import platform
import random

import openpyxl

RUN_CONFIG = {}


def file_sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def bounded(root, relative):
    parts = PurePosixPath(relative).parts
    assert parts and not relative.startswith("/") and ":" not in relative and "\\" not in relative
    assert not any(part in (".", "..") for part in relative.split("/"))
    path = root
    for part in parts:
        path = path / part
        assert not path.is_symlink()
    assert path.resolve().is_relative_to(root) and path.is_file()
    return path


def digest(root, paths):
    result = hashlib.sha256()
    for relative in sorted(paths, key=lambda value: value.encode("utf-8")):
        result.update(relative.encode("utf-8") + b"\0")
        result.update(bytes.fromhex(file_sha(bounded(root, relative))))
    return result.hexdigest()


def observe(root, code, config):
    assert config["run_receipt_protocol_version"] == "1.2.0"
    assert config["data_identity_mode"] == "preprocessing_workbook" and config["solver_backend"] == "python"
    assert len(config["data_paths"]) == 1 and config["auxiliary_data_paths"]
    paths = [*config["data_paths"], *config["auxiliary_data_paths"]]
    assert len({p.casefold() for p in paths}) == len(paths)
    files = [bounded(root, p) for p in paths]
    assert not any(p.samefile(other) for i, p in enumerate(files) for other in files[:i])
    primary_hash = file_sha(files[0])
    auxiliary_hash = digest(root, config["auxiliary_data_paths"])
    assert primary_hash == config["data_sha256"].lower()
    assert auxiliary_hash == config["auxiliary_data_sha256"].lower()
    return dict(data_sha256=primary_hash, auxiliary_data_sha256=auxiliary_hash,
                code_sha256=file_sha(code), code_bundle_sha256=digest(root, [code.relative_to(root).as_posix()]))


def model_input(root, config):
    book = openpyxl.load_workbook(bounded(root, config["data_paths"][0]), data_only=True)
    try:
        rows = list(book["模型输入"].iter_rows(values_only=True))
        payload = dict(zip(rows[0], rows[1]))
    finally:
        book.close()
    offset = sum(json.loads(bounded(root, path).read_text(encoding="utf-8"))["right_hand_side_offset"]
                 for path in config["auxiliary_data_paths"])
    a, b = payload["coefficient"], payload["right_hand_side"] + offset
    assert a != 0 and all(math.isfinite(value) for value in (a, b))
    return a, b, json.loads(payload["sensitivity_coefficients"])


def receipt(config, identity):
    result = {key: config[key] for key in ("stage", "problem_name", "solver_backend", "solver", "random_seed",
                                          "tolerance", "iteration_or_time_limit", "data_identity_mode")}
    result.update(identity)
    result.update(run_receipt_version="1.2.0", execution_owner="user", execution_profile="full_fidelity",
                  solver_version=platform.python_version(), actual_stop_reason="declared_calculation_completed",
                  repetitions_or_scenarios=1, grid_or_time_range="all declared inputs", fallback_used=False,
                  platform=platform.platform(), auxiliary_data_paths=json.dumps(config["auxiliary_data_paths"]))
    result.update({name: False for name in ("allow_reduced_data", "allow_coarser_grid", "allow_shorter_horizon",
                   "allow_fewer_repetitions", "allow_relaxed_tolerance", "allow_silent_solver_fallback")})
    return result


def result_sheets(root, config, a, b, coefficients, metadata):
    if config["stage"] == "analysis":
        primary = bounded(root, "问题一求解/问题一求解结果.xlsx")
        assert file_sha(primary) == config["primary_workbook_sha256"]
        metadata["primary_workbook_sha256"] = file_sha(primary)
        values = [b / coefficient for coefficient in coefficients]
        assert all(math.isfinite(value) and value > 0 for value in values)
        metadata["repetitions_or_scenarios"] = len(values)
        return {"分析设计": [("风险来源", "分析问题", "方法", "指标", "通过标准"),
                            ("系数变化", "解是否为正", "参数敏感性", "x", "所有场景正值")],
                "参数敏感性": [("参数", "基准值", "变化值", "结果指标"),
                            *[("a", a, coefficient, value) for coefficient, value in zip(coefficients, values)]],
                "结论稳定性汇总": [("核心结论", "分析方法", "稳定范围", "是否保持"),
                                  ("解为正值", "参数敏感性", "全部声明场景", True)]}
    solution = b / a
    residual = abs(a * solution - b)
    passed = math.isfinite(solution) and residual <= config["tolerance"]
    assert passed
    metadata["primary_quality_protocol_version"] = config["primary_quality_protocol_version"]
    return {"核心指标": [("指标", "数值"), ("解", solution), ("绝对方程残差", residual)],
            "数据审计": [("等级", "检查项", "信息", "处理方式"), ("info", "有限输入", "已验证", "只读检查")],
            "主结果质量门": [("Verification ID", "检查项", "是否通过", "证据", "判定关系", "阈值或容差",
                            "实际值", "证据工作表", "阈值来源"),
                         ("PQ-Q1-01", "方程残差", passed, "原系数回代", "abs<=", config["tolerance"],
                          residual, "均衡残差", "solver_tolerance")],
            "均衡残差": [("主体或均衡", "残差", "容差", "是否满足"), ("a*x=b", residual, config["tolerance"], passed)],
            "状态明细": [("记录键", "状态", "数值", "单位"), ("0001", "解", solution, "无量纲")]}


def main():
    code = Path(__file__).resolve()
    root, config = code.parents[1], RUN_CONFIG
    identity = observe(root, code, config)
    random.seed(config["random_seed"])
    a, b, coefficients = model_input(root, config)
    metadata = receipt(config, identity)
    sheets = result_sheets(root, config, a, b, coefficients, metadata)
    sheets["运行配置"] = [("项目", "值"), *metadata.items()]
    assert observe(root, code, config) == identity
    if config["stage"] == "analysis":
        assert file_sha(bounded(root, "问题一求解/问题一求解结果.xlsx")) == config["primary_workbook_sha256"]
    book = openpyxl.Workbook()
    book.remove(book.active)
    try:
        for name, rows in sheets.items():
            sheet = book.create_sheet(name)
            for row in rows:
                sheet.append(row)
        book.save(root / config["expected_workbook"])
    finally:
        book.close()
    print(json.dumps({"synthetic_native_python": True, "stage": config["stage"], "identity": identity}))


if __name__ == "__main__":
    main()
