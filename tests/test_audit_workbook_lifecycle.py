from __future__ import annotations

import importlib.util
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "audit_workbook_lifecycle", ROOT / "scripts/validate_user_execution.py"
)
RECEIPT = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(RECEIPT)


class WorkbookLifecycleTests(unittest.TestCase):
    def check_reader(self, reader, sheets, *, expected=None, issue=None, read_error=False):
        """Retain real workbooks so neither the assertion nor rename relies on GC."""
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "验收工作簿.xlsx"
            book = openpyxl.Workbook()
            book.remove(book.active)
            for name, rows in sheets.items():
                sheet = book.create_sheet(name)
                for row in rows:
                    sheet.append(row)
            book.save(path)
            book.close()
            opened = []
            load_workbook = openpyxl.load_workbook

            def retain_workbook(*args, **kwargs):
                loaded = load_workbook(*args, **kwargs)
                opened.append(loaded)
                return loaded

            try:
                with ExitStack() as stack:
                    stack.enter_context(patch.object(openpyxl, "load_workbook", retain_workbook))
                    if read_error:
                        stack.enter_context(patch.object(
                            ReadOnlyWorksheet, "iter_rows", side_effect=ValueError("fixture read failure")
                        ))
                        with self.assertRaisesRegex(ValueError, "fixture read failure"):
                            reader(path)
                    else:
                        result = reader(path)
                        if expected is not None:
                            self.assertEqual(result, expected)
                        if issue is not None:
                            self.assertTrue(any(issue in text for text in result[-1]), result)
                self.assertTrue(opened)
                # Native Windows rename/replace must succeed before test cleanup closes anything.
                renamed = path.with_name("已验收.xlsx")
                path.rename(renamed)
                renamed.replace(path)
                self.assertTrue(path.is_file())
                # POSIX permits renaming open files, so also verify actual archive ownership there.
                for loaded in opened:
                    self.assertIsNone(loaded._archive.fp)
            finally:
                for loaded in opened:
                    loaded.close()

    def config_sheets(self):
        config = {
            "execution_owner": "user", "execution_profile": "full_fidelity", "stage": "primary",
            "problem_name": "问题一", "code_sha256": "a" * 64, "data_sha256": "b" * 64,
            "solver": "fixture", "solver_version": "1", "tolerance": 1e-8,
            "iteration_or_time_limit": "full", "actual_stop_reason": "optimal", "random_seed": 1,
            "repetitions_or_scenarios": 1, "grid_or_time_range": "full", "fallback_used": False,
            "platform": "fixture", **{name: False for name in RECEIPT.FALSE_FLAGS},
        }
        return {"运行配置": [["项目", "值"], *config.items()]}, config

    def preprocessing_sheets(self):
        sheets = {
            name: [columns, ["fixture"] * len(columns)]
            for name, columns in RECEIPT.PREPROCESSING_EVIDENCE_SHEETS.items()
        }
        sheets["绘图数据索引"] = [["图组", "图作用", "源工作表"], ["F1", "核对", "底层数据"]]
        sheets["预处理质量门"] = [["检查项", "是否通过", "证据"], ["完整", True, "fixture"]]
        sheets["底层数据"] = [["对象", "值"], ["A", 1]]
        return sheets

    def analysis_sheets(self):
        return {"分析设计": [["方法"], ["fixture"]],
                "结论稳定性汇总": [["核心结论", "是否保持"], ["结论A", True]]}

    def test_configuration_success_releases_file(self):
        sheets, config = self.config_sheets()
        self.check_reader(RECEIPT.configuration_map, sheets, expected=(config, []))

    def test_configuration_early_returns_release_file(self):
        for sheets, issue in [
            ({"其他": [["值"]]}, "缺少运行配置"),
            ({"运行配置": []}, "运行配置表头"),
            ({"运行配置": [["错误", "表头"]]}, "运行配置表头"),
        ]:
            with self.subTest(sheets=sheets):
                self.check_reader(RECEIPT.configuration_map, sheets, issue=issue)

    def test_boolean_gate_success_and_rejection_release_file(self):
        for passed in (True, False):
            with self.subTest(passed=passed):
                issues = [] if passed else ["主结果质量门未通过: 数值检查"]
                self.check_reader(RECEIPT.quality_passed,
                    {"主结果质量门": [["检查项", "是否通过"], ["数值检查", passed]]},
                    expected=(passed, issues))

    def test_boolean_gate_early_returns_release_file(self):
        for sheets, issue in [
            ({"其他": [["值"]]}, "缺少主结果质量门"),
            ({"主结果质量门": []}, "主结果质量门为空"),
            ({"主结果质量门": [["检查项", "错误列"]]}, "缺少是否通过列"),
        ]:
            with self.subTest(sheets=sheets):
                self.check_reader(RECEIPT.quality_passed, sheets, issue=issue)

    def test_preprocessing_success_releases_both_workbooks(self):
        self.check_reader(RECEIPT.preprocessing_passed, self.preprocessing_sheets(), expected=(True, []))

    def test_preprocessing_early_returns_release_file(self):
        for mutation, issue in [("missing", "缺少工作表"), ("empty", "实质证据"), ("header", "缺少列")]:
            with self.subTest(mutation=mutation):
                sheets = self.preprocessing_sheets()
                if mutation == "missing":
                    del sheets["数据审计"]
                elif mutation == "empty":
                    sheets["数据审计"] = []
                else:
                    sheets["数据审计"] = [["错误列"], ["fixture"]]
                self.check_reader(RECEIPT.preprocessing_passed, sheets, issue=issue)

    def test_preprocessing_invalid_reference_releases_both_workbooks(self):
        sheets = self.preprocessing_sheets()
        del sheets["底层数据"]
        self.check_reader(RECEIPT.preprocessing_passed, sheets, issue="引用不存在的工作表")

    def test_analysis_success_and_redo_release_file(self):
        for stable in (True, False):
            with self.subTest(stable=stable):
                sheets = self.analysis_sheets()
                sheets["结论稳定性汇总"][1][1] = stable
                self.check_reader(RECEIPT.analysis_passed, sheets, expected=(
                    stable, "passed" if stable else "redo_required", [] if stable else ["存在核心结论未保持"]
                ))

    def test_analysis_early_returns_release_file(self):
        for mutation, issue in [("missing", "缺少工作表"), ("empty", "无实质数据"), ("header", "缺少是否保持列")]:
            with self.subTest(mutation=mutation):
                sheets = self.analysis_sheets()
                if mutation == "missing":
                    del sheets["分析设计"]
                elif mutation == "empty":
                    sheets["结论稳定性汇总"] = []
                else:
                    sheets["结论稳定性汇总"] = [["错误列"], ["fixture"]]
                self.check_reader(RECEIPT.analysis_passed, sheets, issue=issue)

    def test_read_exceptions_propagate_after_releasing_file(self):
        for reader, sheets in [
            (RECEIPT.configuration_map, self.config_sheets()[0]),
            (RECEIPT.quality_passed, {"主结果质量门": [["检查项", "是否通过"], ["检查", True]]}),
            (RECEIPT.preprocessing_passed, self.preprocessing_sheets()),
            (RECEIPT.analysis_passed, self.analysis_sheets()),
        ]:
            with self.subTest(reader=reader.__name__):
                self.check_reader(reader, sheets, read_error=True)


if __name__ == "__main__":
    unittest.main()
