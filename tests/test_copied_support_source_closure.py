"""B23/T18: real copied workbook helpers through the modern source/receipt chain.

Only isolated repository-maintenance scalar fixtures are executed here. These
checks do not authorize competition code, certify legacy pipeline state writers,
or enable the future project-wide backend selection/migration interface.
"""
from __future__ import annotations

import ast
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import stage_code as STAGE
import validate_code_delivery as DELIVERY
import validate_user_execution as RECEIPT

QUESTION = "问题一求解"
HELPERS = ("__init__.py", "main_pipeline.py", "result_io.py", "workbook_validation.py")
DYNAMIC_LOADER = '''def _load_workbook_validation():
    import importlib.util
    import sys
    path = Path(__file__).resolve().with_name("workbook_validation.py")
    spec = importlib.util.spec_from_file_location("hsk_pipeline_workbook_validation", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module

WORKBOOK_VALIDATION = _load_workbook_validation()
'''


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reference_bundle(root: Path, paths: list[str]) -> str:
    result = hashlib.sha256()
    for relative in sorted(paths, key=lambda value: value.encode("utf-8")):
        result.update(relative.encode("utf-8") + b"\0")
        result.update(bytes.fromhex(digest(root / relative)))
    return result.hexdigest()


def replace_once(text: str, old: str, new: str) -> str:
    if text.count(old) != 1:
        raise AssertionError(f"Maintenance fixture anchor is not unique: {old!r}")
    return text.replace(old, new, 1)


def instantiate(root: Path, *, minimal: bool = False) -> tuple[Path, dict]:
    folder = root / QUESTION
    support = folder / "hsk_pipeline"
    support.mkdir(parents=True)
    for name in HELPERS:
        if minimal and name == "main_pipeline.py":
            continue
        if minimal and name == "__init__.py":
            (support / name).write_text('"""Project-local workbook I/O only."""\n', encoding="utf-8")
        else:
            shutil.copyfile(ROOT / "templates/code/hsk_pipeline" / name, support / name)
    shutil.copyfile(ROOT / "tests/fixtures/solver_backends/input.json", root / "input.json")
    code = folder / "问题一求解.py"
    config = {
        "stage": "primary", "problem_name": "问题一", "solver_backend": "python",
        "data_paths": ["input.json"], "data_sha256": reference_bundle(root, ["input.json"]),
        "solver": "scalar_division", "random_seed": 2026, "tolerance": 1e-10,
        "iteration_or_time_limit": "direct", "expected_workbook": f"{QUESTION}/问题一求解结果.xlsx",
        "run_receipt_protocol_version": "1.1.0", "primary_quality_protocol_version": "1.0.0",
        "code_dependencies": [], "input_mode": "json",
    }
    refresh_dependencies(root, config)
    source = (ROOT / "tests/fixtures/solver_backends/python_primary.py").read_text(encoding="utf-8")
    source = replace_once(source, "import openpyxl\n", "import openpyxl\nimport pandas as pd\n")
    source = replace_once(source, "RUN_CONFIG = {}", "RUN_CONFIG = " + repr(config))
    source = replace_once(source, '    bundle_sha = digest(root, [code.relative_to(root).as_posix()])', '''    for dependency in config["code_dependencies"]:
        assert hashlib.sha256((root / dependency["path"]).read_bytes()).hexdigest() == dependency["sha256"]
    sources = [code.relative_to(root).as_posix(), *[item["path"] for item in config["code_dependencies"]]]
    bundle_sha = digest(root, sources)
    # Import copied helpers only after checking their declared byte identities.
    from hsk_pipeline.result_io import write_workbook''')
    source = replace_once(source, "    book = openpyxl.Workbook()\n    book.remove(book.active)\n", "")
    source = replace_once(source, '''    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
''', '''    tables = {name: pd.DataFrame(rows[1:], columns=rows[0]) for name, rows in sheets.items()}
''')
    source = replace_once(source, '''    book.save(root / config["expected_workbook"])
    book.close()''', '''    assert digest(root, sources) == bundle_sha
    write_workbook(root / config["expected_workbook"], tables, workbook_kind="solution",
                   capabilities={"requires_equilibrium_residual": True})''')
    ast.parse(source)
    code.write_text(source, encoding="utf-8", newline="\n")
    (root / "state").mkdir()
    state = {
        "project": {"competition": "test", "problem": "synthetic", "current_phase": "solve_validate",
                    "state_generation": 0},
        "execution": {"solver_backend": "python",
                      "solver_backend_selection_reason": "全题维护微例与支撑依赖已审视"},
        "preprocessing": {"decision": "not_needed", "status": "not_applicable", "quality_status": "not_applicable"},
        "subproblems": {"Q1": {
            "status": "designed", "selected_model": "a*x=b",
            "capabilities": {"requires_equilibrium_residual": True},
            "primary_execution_status": "pending", "result_quality_status": "pending",
            "result_analysis_status": "pending",
            "solver_execution": {"primary": {}},
        }},
    }
    (root / "state/project_state.yaml").write_text(yaml.safe_dump(state, allow_unicode=True), encoding="utf-8")
    (root / "requirements.txt").write_text("numpy\npandas>=2.0\nopenpyxl>=3.1\nPyYAML>=6.0\n", encoding="utf-8")
    return code, config


def refresh_dependencies(root: Path, config: dict) -> None:
    config["code_dependencies"] = [
        {"path": path.relative_to(root).as_posix(), "sha256": digest(path)}
        for path in sorted((root / QUESTION / "hsk_pipeline").glob("*.py"))
    ]


def rewrite_config(code: Path, config: dict) -> None:
    source = code.read_text(encoding="utf-8")
    assignment = next(node for node in ast.parse(source).body if isinstance(node, ast.Assign)
                      and any(isinstance(target, ast.Name) and target.id == "RUN_CONFIG" for target in node.targets))
    lines = source.splitlines(keepends=True)
    lines[assignment.lineno - 1:assignment.end_lineno] = ["RUN_CONFIG = " + repr(config) + "\n"]
    code.write_text("".join(lines), encoding="utf-8", newline="\n")


class CopiedSupportSourceClosureTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.code, self.config = instantiate(self.root)
        self.state_path = self.root / "state/project_state.yaml"
        self.workbook = self.root / self.config["expected_workbook"]

    def state(self):
        return yaml.safe_load(self.state_path.read_text(encoding="utf-8"))

    def invoke(self, script: Path, *arguments: str):
        env = {key: value for key, value in os.environ.items()
               if key not in {"PYTHONPATH", "HSK_WORKBOOK_SCHEMA"}}
        env["PYTHONUTF8"] = "1"
        return subprocess.run([sys.executable, "-B", str(script), *arguments], cwd=self.root,
                              env=env, capture_output=True, encoding="utf-8", timeout=45)

    def deliver(self):
        process = self.invoke(ROOT / "scripts/validate_code_delivery.py", str(self.root),
                              "--script", self.code.relative_to(self.root).as_posix(),
                              "--stage", "primary", "--write", "--strict")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        report = yaml.safe_load(process.stdout)
        self.assertEqual(report["status"], "passed", report)
        self.assertFalse(report["task_code_executed"])
        self.assertFalse(self.workbook.exists())
        binding = self.state()["subproblems"]["Q1"]["solver_execution"]["primary"]
        self.assertEqual(binding["bundle_sha256"], reference_bundle(self.root, [
            self.code.relative_to(self.root).as_posix(), *[item["path"] for item in self.config["code_dependencies"]]]))
        self.assertNotIn("validated_bundle_sha256", binding)
        return binding

    def execute(self):
        before = self.state_path.read_bytes()
        input_bytes = (self.root / "input.json").read_bytes()
        process = self.invoke(self.code)
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        self.assertEqual(json.loads(process.stdout)["answer"], 3)
        self.assertEqual(self.state_path.read_bytes(), before, "The maintenance numerical entry must not write project state")
        self.assertEqual((self.root / "input.json").read_bytes(), input_bytes)
        self.assertTrue(self.workbook.is_file())

    def test_complete_copied_package_passes_static_delivery_without_execution(self):
        before = self.state_path.read_bytes()
        issues, config = DELIVERY.validate_script(self.root, self.code, "primary")
        self.assertEqual(issues, [])
        fingerprint = STAGE.stage_code_fingerprint(self.root, self.code, config["code_dependencies"])
        self.assertEqual(len(fingerprint["files"]), 5)
        self.assertFalse(self.workbook.exists())
        self.assertEqual(self.state_path.read_bytes(), before)

    def test_real_delivery_execution_receipt_acceptance_and_primary_identity(self):
        binding = self.deliver()
        self.execute()
        receipt, issues = RECEIPT.configuration_map(self.workbook)
        self.assertEqual(issues, [])
        self.assertEqual(receipt["run_receipt_version"], "1.1.0")
        self.assertEqual(receipt["code_bundle_sha256"], binding["bundle_sha256"])
        process = self.invoke(ROOT / "scripts/validate_user_execution.py", str(self.root),
                              "--workbook", self.workbook.relative_to(self.root).as_posix(), "--write", "--strict")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        self.assertEqual(yaml.safe_load(process.stdout)["status"], "passed")
        entry = self.state()["subproblems"]["Q1"]
        self.assertEqual(entry["primary_execution_status"], "accepted")
        self.assertEqual(entry["solver_execution"]["primary"]["validated_bundle_sha256"], binding["bundle_sha256"])
        workbook = openpyxl.load_workbook(self.workbook, data_only=True)
        try:
            self.assertEqual(dict(workbook["核心指标"].iter_rows(min_row=2, values_only=True))["解"], 3)
            self.assertEqual(workbook["状态明细"]["A2"].value, "0001")
        finally:
            workbook.close()

    def test_minimal_io_package_needs_no_legacy_pipeline(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            code, config = instantiate(root, minimal=True)
            self.assertEqual(DELIVERY.validate_script(root, code, "primary")[0], [])
            self.assertEqual(len(config["code_dependencies"]), 3)
            self.assertFalse((root / QUESTION / "hsk_pipeline/main_pipeline.py").exists())
            before = (root / "state/project_state.yaml").read_bytes()
            process = self.invoke(code)
            self.assertEqual(process.returncode, 0, process.stderr)
            self.assertEqual(json.loads(process.stdout)["answer"], 3)
            self.assertEqual((root / "state/project_state.yaml").read_bytes(), before)

    def test_each_undeclared_transitive_helper_is_rejected(self):
        before = self.state_path.read_bytes()
        for dependency in self.config["code_dependencies"]:
            with self.subTest(path=dependency["path"]):
                bad = deepcopy(self.config)
                bad["code_dependencies"] = [item for item in bad["code_dependencies"] if item != dependency]
                rewrite_config(self.code, bad)
                issues, _ = DELIVERY.validate_script(self.root, self.code, "primary")
                self.assertTrue(any("未声明" in issue and dependency["path"] in issue for issue in issues), issues)
        self.assertEqual(self.state_path.read_bytes(), before)
        self.assertFalse(self.workbook.exists())

    def test_declared_helper_hash_mismatch_is_rejected(self):
        for dependency in self.config["code_dependencies"]:
            with self.subTest(path=dependency["path"]):
                bad = deepcopy(self.config)
                next(item for item in bad["code_dependencies"] if item["path"] == dependency["path"])["sha256"] = "0" * 64
                rewrite_config(self.code, bad)
                self.assertTrue(any("SHA-256不一致" in issue for issue in DELIVERY.validate_script(self.root, self.code)[0]))

    def test_missing_helper_is_rejected(self):
        helper = self.root / QUESTION / "hsk_pipeline/workbook_validation.py"
        helper.unlink()
        self.assertTrue(any("源码文件不存在" in issue for issue in DELIVERY.validate_script(self.root, self.code)[0]))

    def test_restored_dynamic_loader_still_rejected_with_current_hashes(self):
        helper = self.root / QUESTION / "hsk_pipeline/result_io.py"
        text = helper.read_text(encoding="utf-8")
        # Reintroduce the historical loader without hiding it behind missing hashes.
        text += "\n" + DYNAMIC_LOADER
        helper.write_text(text, encoding="utf-8")
        refresh_dependencies(self.root, self.config)
        rewrite_config(self.code, self.config)
        issues, _ = DELIVERY.validate_script(self.root, self.code, "primary")
        self.assertIn("新源码闭包不支持动态Python代码加载", issues)
        self.assertIn("新源码闭包不支持执行命名空间的间接传递/反射", issues)
        self.assertFalse(any("SHA-256不一致" in issue for issue in issues), issues)

    def test_changed_helper_stops_fixture_before_execution_and_state_write(self):
        self.deliver()
        before = self.state_path.read_bytes()
        helper = self.root / QUESTION / "hsk_pipeline/workbook_validation.py"
        helper.write_bytes(helper.read_bytes() + b"\n# changed after delivery\n")
        process = self.invoke(self.code)
        self.assertNotEqual(process.returncode, 0)
        self.assertFalse(self.workbook.exists())
        self.assertEqual(self.state_path.read_bytes(), before)

    def test_old_workbook_cannot_be_accepted_after_helper_change(self):
        self.deliver()
        self.execute()
        old_workbook = self.workbook.read_bytes()
        helper = self.root / QUESTION / "hsk_pipeline/workbook_validation.py"
        helper.write_bytes(helper.read_bytes() + b"\n# changed after execution\n")
        state = self.state()
        old_binding = deepcopy(state["subproblems"]["Q1"]["solver_execution"]["primary"])
        issues = RECEIPT.validate_one(self.root, self.workbook, state, True)
        self.assertTrue(issues)
        self.assertEqual(state["subproblems"]["Q1"]["solver_execution"]["primary"], old_binding)
        self.assertNotEqual(state["subproblems"]["Q1"].get("primary_execution_status"), "accepted")
        self.assertEqual(self.workbook.read_bytes(), old_workbook)

    def test_redelivered_new_bundle_does_not_validate_old_receipt(self):
        self.deliver()
        self.execute()
        old_bytes = self.workbook.read_bytes()
        old_receipt, _ = RECEIPT.configuration_map(self.workbook)
        helper = self.root / QUESTION / "hsk_pipeline/workbook_validation.py"
        helper.write_bytes(helper.read_bytes() + b"\n# new delivered source revision\n")
        refresh_dependencies(self.root, self.config)
        rewrite_config(self.code, self.config)
        self.assertEqual(DELIVERY.validate_script(self.root, self.code)[0], [])
        process = self.invoke(ROOT / "scripts/validate_code_delivery.py", str(self.root),
                              "--script", self.code.relative_to(self.root).as_posix(),
                              "--stage", "primary", "--write", "--strict")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        self.assertEqual(yaml.safe_load(process.stdout)["status"], "passed")
        state = self.state()
        binding = deepcopy(state["subproblems"]["Q1"]["solver_execution"]["primary"])
        self.assertNotEqual(binding["bundle_sha256"], old_receipt["code_bundle_sha256"])
        self.assertNotIn("validated_bundle_sha256", binding)
        issues = RECEIPT.validate_one(self.root, self.workbook, state, True)
        self.assertTrue(any("bundle" in issue or "code_sha256" in issue for issue in issues), issues)
        self.assertEqual(state["subproblems"]["Q1"]["solver_execution"]["primary"], binding)
        self.assertNotEqual(state["subproblems"]["Q1"].get("primary_execution_status"), "accepted")
        self.assertEqual(self.workbook.read_bytes(), old_bytes)

    def test_legacy_flat_io_import_remains_runnable_without_skill_path(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            for name in ("result_io.py", "workbook_validation.py"):
                shutil.copyfile(ROOT / "templates/code/hsk_pipeline" / name, root / name)
            probe = root / "probe.py"
            probe.write_text('''from pathlib import Path
import result_io
assert result_io.load_workbook_schema() is result_io._FALLBACK_SCHEMA
path = Path(__file__).with_name("roundtrip.xlsx")
result_io.write_workbook(path, {"维护微例": [{"记录键": "0001", "数值": 3}]})
assert result_io.read_workbook_tables(path)["维护微例"]["数值"].iloc[0] == 3
''', encoding="utf-8")
            process = self.invoke(probe)
            self.assertEqual(process.returncode, 0, process.stderr)


if __name__ == "__main__":
    unittest.main()
