"""Historical byte preservation and durable recovery; isolated maintenance fixtures only."""
from __future__ import annotations

from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import yaml

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))
import project_transaction as TX


class HistoryArchiveTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / "state").mkdir()
        self.state = {"project": {"state_generation": 0, "current_phase": "model_design"}, "subproblems": {}}
        self.state_path = self.root / TX.STATE_RELATIVE_PATH
        self.state_path.write_bytes(("# 原始记录\r\n" + yaml.safe_dump(self.state, allow_unicode=True)).encode("utf-8"))
        self.code = "问题一求解/问题一求解.py"
        (self.root / self.code).parent.mkdir()
        (self.root / self.code).write_bytes(b"# untouched historical source\r\nvalue = 3\r\n")
        self.workbook = "问题一求解/问题一求解结果.xlsx"
        book = Workbook()
        book.active.title = "结果"
        book.active.append(["键", "结果"])
        book.active.append(["0001", 3])
        book.create_sheet("运行回执").append(["historical", "python"])
        book.save(self.root / self.workbook)
        book.close()
        (self.root / "模型论文框架.md").write_text("原始框架\n", encoding="utf-8")
        (self.root / "raw.bin").write_bytes(bytes(range(256)) * 5000)
        self.expected = self.hashes(TX.STATE_RELATIVE_PATH, self.code, self.workbook,
                                    "模型论文框架.md", "raw.bin", "absent.txt")
        self.old_bytes = {name: (self.root / name).read_bytes() for name, digest in self.expected.items() if digest}
        self.directory = "state/backend_history/migration-test"

    def hashes(self, *names):
        return {name: TX.sha256_file(self.root / name) if (self.root / name).is_file() else None for name in names}

    def prepare(self, **options):
        return TX.prepare_history_archive(self.root, options.pop("archive_relative", self.directory),
            expected_generation=options.pop("expected_generation", 0),
            expected_file_hashes=options.pop("expected_file_hashes", self.expected), **options)

    def stored(self, relative):
        return self.root / self.directory / "files" / relative

    def commit(self, ref, **options):
        expected = {**self.expected, "history_reference.yaml": None}
        return TX.commit_project_state(self.root, deepcopy(self.state), expected_generation=0,
            expected_file_hashes=options.pop("expected_file_hashes", expected),
            preserved_archives=options.pop("preserved_archives", [ref]),
            writes_after_state=options.pop("writes_after_state", [("history_reference.yaml", yaml.safe_dump(ref))]), **options)

    def assert_live_original(self):
        for relative, content in self.old_bytes.items():
            self.assertEqual((self.root / relative).read_bytes(), content, relative)
        self.assertFalse((self.root / "absent.txt").exists())

    def assert_no_journal(self):
        self.assertFalse((self.root / TX.JOURNAL_RELATIVE_PATH).exists())
        self.assertFalse(list(self.root.rglob("*.stage")))
        self.assertFalse(list(self.root.rglob("*.bak")))

    def rewrite_manifest(self, ref, mutate):
        path = self.root / ref["manifest"]
        data = json.loads(path.read_text(encoding="utf-8"))
        mutate(data)
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return {"manifest": ref["manifest"], "sha256": TX.sha256_file(path)}

    def test_real_workbook_source_and_binary_bytes_are_independent_copies(self):
        ref = self.prepare()
        report = TX.verify_history_archive(self.root, ref)
        self.assertEqual(report["source_hashes"], self.expected)
        self.assertEqual(report["file_count"], len(self.old_bytes))
        self.assertEqual(report["total_bytes"], sum(map(len, self.old_bytes.values())))
        for relative, content in self.old_bytes.items():
            self.assertEqual(self.stored(relative).read_bytes(), content)
            self.assertFalse(os.path.samefile(self.stored(relative), self.root / relative))
        book = load_workbook(self.stored(self.workbook), read_only=True, data_only=True)
        try:
            self.assertEqual(list(book["结果"].values)[1], ("0001", 3))
        finally:
            book.close()
        self.assert_live_original()
        self.assert_no_journal()

    def test_verification_is_read_only_and_survives_original_replacement(self):
        ref = self.prepare()
        (self.root / self.code).unlink()
        (self.root / self.workbook).write_bytes(b"later output")
        self.state_path.write_text("project: {state_generation: 4}\n")
        before = {p: p.read_bytes() for p in (self.root / self.directory).rglob("*") if p.is_file()}
        self.assertEqual(TX.verify_history_archive(self.root, ref)["base_generation"], 0)
        self.assertEqual(before, {p: p.read_bytes() for p in before})
        self.assertEqual(self.state_path.read_text(), "project: {state_generation: 4}\n")

    @unittest.skipIf(os.name == "nt", "symlink creation can require Windows privileges")
    def test_historical_names_do_not_resolve_retired_original_symlinks(self):
        ref = self.prepare()
        original = self.root / self.code
        original.unlink()
        original.symlink_to(self.root.parent / "missing-target")
        TX.verify_history_archive(self.root, ref)

    def test_uppercase_source_and_reference_digests(self):
        expected = {p: h.upper() if h else None for p, h in self.expected.items()}
        ref = self.prepare(expected_file_hashes=expected)
        ref["sha256"] = ref["sha256"].upper()
        self.assertEqual(TX.verify_history_archive(self.root, ref)["source_hashes"], self.expected)

    def test_existing_destination_is_not_modified(self):
        ref = self.prepare()
        before = (self.root / ref["manifest"]).read_bytes()
        with self.assertRaises(TX.HistoryArchiveError) as error:
            self.prepare()
        self.assertEqual(error.exception.cleanup_status, "not_created")
        self.assertEqual((self.root / ref["manifest"]).read_bytes(), before)
        self.assert_live_original()

    def test_bad_paths_and_source_overlap_are_rejected(self):
        for name in ("../outside", "/tmp/archive", "state/./history", "state/../history", "state//history",
                     "C:/history", "state\\history", "state", TX.STATE_RELATIVE_PATH, self.code + "/nested"):
            with self.subTest(name=name), self.assertRaises(TX.ProjectTransactionError):
                self.prepare(archive_relative=name)
        self.assert_live_original()

    def test_archive_requires_original_state_and_valid_snapshot_types(self):
        for expected in ({}, {TX.STATE_RELATIVE_PATH: None}, {TX.STATE_RELATIVE_PATH: "bad"}, []):
            with self.subTest(expected=expected), self.assertRaises(TX.ProjectTransactionError):
                self.prepare(expected_file_hashes=expected)
        for generation in (True, "0", -1, 0.0):
            with self.subTest(generation=generation), self.assertRaises(TX.ProjectTransactionError):
                self.prepare(expected_generation=generation)
        self.assert_live_original()

    def test_absent_source_cannot_silently_appear(self):
        (self.root / "absent.txt").write_text("unexpected")
        with self.assertRaises(TX.ReadSetConflictError):
            self.prepare()
        self.assertFalse((self.root / self.directory).exists())

    def test_same_generation_byte_edit_rejects_old_confirmation(self):
        self.state_path.write_bytes(self.old_bytes[TX.STATE_RELATIVE_PATH] + b"# changed\n")
        with self.assertRaises(TX.ReadSetConflictError):
            self.prepare()
        self.assertFalse((self.root / self.directory).exists())

    def test_generation_conflict_rejects_matching_bytes(self):
        with self.assertRaises(TX.GenerationConflictError):
            self.prepare(expected_generation=1)
        self.assertFalse((self.root / self.directory).exists())

    def test_low_disk_space_fails_before_archive_creation(self):
        usage = type("Usage", (), {"free": 0})()
        with patch.object(TX.shutil, "disk_usage", return_value=usage), self.assertRaises(TX.HistoryArchiveError):
            self.prepare()
        self.assertFalse((self.root / self.directory).exists())
        self.assert_live_original()

    def test_prepare_never_recovers_an_existing_journal(self):
        journal = self.root / TX.JOURNAL_RELATIVE_PATH
        journal.write_text("status: prepared\nversion: 1\n", encoding="utf-8")
        with self.assertRaises(TX.TransactionRecoveryError):
            self.prepare()
        self.assertEqual(journal.read_text(), "status: prepared\nversion: 1\n")
        self.assert_live_original()

    def test_partial_copy_failure_reports_retained_directory(self):
        def stop(point):
            if point == "after_archive_copy:raw.bin":
                raise OSError("simulated disk I/O failure")
        with self.assertRaises(TX.HistoryArchiveError) as error:
            self.prepare(failure_hook=stop)
        self.assertEqual(error.exception.archive_relative, self.directory)
        self.assertEqual(error.exception.cleanup_status, "retained")
        self.assertFalse(error.exception.state_commit_performed)
        self.assertFalse((self.root / self.directory / "manifest.json").exists())
        self.assertTrue(self.stored("raw.bin").is_file())
        self.assert_live_original()
        self.assert_no_journal()

    def test_failure_after_manifest_does_not_delete_unreferenced_history(self):
        def stop(point):
            if point == "after_archive_manifest":
                raise RuntimeError("interrupted before return")
        with self.assertRaises(TX.HistoryArchiveError) as error:
            self.prepare(failure_hook=stop)
        self.assertEqual(error.exception.cleanup_status, "retained")
        manifest = self.root / self.directory / "manifest.json"
        TX.verify_history_archive(self.root, {"manifest": manifest.relative_to(self.root).as_posix(),
                                             "sha256": TX.sha256_file(manifest)})
        self.assert_live_original()

    def test_source_change_during_copy_rejects_and_preserves_external_edit(self):
        def change(point):
            if point == f"after_archive_copy:{self.code}":
                (self.root / self.code).write_bytes(b"external edit")
        with self.assertRaises(TX.HistoryArchiveError) as error:
            self.prepare(failure_hook=change)
        self.assertIsInstance(error.exception.__cause__, TX.ReadSetConflictError)
        self.assertEqual((self.root / self.code).read_bytes(), b"external edit")
        self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])

    def test_caller_cannot_rewrite_captured_expected_hashes(self):
        expected = dict(self.expected)
        def change(point):
            if point == "before_archive_manifest":
                (self.root / "raw.bin").write_bytes(b"new")
                expected["raw.bin"] = TX.sha256_file(self.root / "raw.bin")
        with self.assertRaises(TX.HistoryArchiveError):
            self.prepare(expected_file_hashes=expected, failure_hook=change)
        self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])

    def test_missing_tampered_and_unlisted_archive_members_reject(self):
        ref = self.prepare()
        original = self.stored(self.workbook).read_bytes()
        for value in (b"corrupt", None):
            with self.subTest(value=value):
                if value is None:
                    self.stored(self.workbook).unlink()
                else:
                    self.stored(self.workbook).write_bytes(value)
                with self.assertRaises((TX.ProjectTransactionError, OSError)):
                    TX.verify_history_archive(self.root, ref)
                self.stored(self.workbook).write_bytes(original)
        extra = self.root / self.directory / "unlisted.bin"
        extra.write_bytes(b"unknown")
        with self.assertRaises(TX.ProjectTransactionError):
            TX.verify_history_archive(self.root, ref)
        extra.unlink()
        extra.mkdir()
        with self.assertRaises(TX.ProjectTransactionError):
            TX.verify_history_archive(self.root, ref)

    def test_manifest_digest_is_external_not_self_attested(self):
        ref = self.prepare()
        manifest = self.root / ref["manifest"]
        manifest.write_bytes(manifest.read_bytes() + b" ")
        with self.assertRaisesRegex(TX.ProjectTransactionError, "digest"):
            TX.verify_history_archive(self.root, ref)

    def test_unknown_manifest_fields_invalid_size_and_generation_reject(self):
        ref = self.prepare()
        path = self.root / ref["manifest"]
        original = path.read_bytes()
        cases = [lambda m: m.update(version=2), lambda m: m.update(version=True),
                 lambda m: m.update(extra="not permitted"), lambda m: m.update(base_generation=1),
                 lambda m: m["files"][self.code].update(size_bytes=-1),
                 lambda m: m["files"][self.code].update(sha256=None),
                 lambda m: m["files"].pop(TX.STATE_RELATIVE_PATH)]
        for index, mutate in enumerate(cases):
            with self.subTest(case=index):
                path.write_bytes(original)
                changed_ref = self.rewrite_manifest(ref, mutate)
                with self.assertRaises(TX.ProjectTransactionError):
                    TX.verify_history_archive(self.root, changed_ref)

    def test_duplicate_json_keys_invalid_encoding_and_oversize_reject(self):
        ref = self.prepare()
        path = self.root / ref["manifest"]
        original = path.read_bytes()
        for raw in (original.replace(b'"version": 1', b'"version": 1, "version": 1'), b"\xff\xfe"):
            path.write_bytes(raw)
            changed_ref = {**ref, "sha256": hashlib.sha256(raw).hexdigest()}
            with self.assertRaises(TX.ProjectTransactionError):
                TX.verify_history_archive(self.root, changed_ref)
        path.write_bytes(original)
        with patch.object(TX, "MAX_ARCHIVE_MANIFEST_BYTES", 10), self.assertRaises(TX.ProjectTransactionError):
            TX.verify_history_archive(self.root, ref)

    @unittest.skipIf(os.name == "nt", "symlink creation can require Windows privileges")
    def test_symlink_source_archive_member_and_destination_are_rejected(self):
        original = self.root / self.code
        old = original.read_bytes()
        original.unlink()
        original.symlink_to(self.root / "raw.bin")
        with self.assertRaises(TX.ProjectTransactionError):
            self.prepare()
        original.unlink(); original.write_bytes(old)
        ref = self.prepare()
        stored = self.stored(self.code)
        stored.unlink(); stored.symlink_to(original)
        with self.assertRaises(TX.ProjectTransactionError):
            TX.verify_history_archive(self.root, ref)
        target = self.root / "state" / "alias"
        target.symlink_to(self.root / "问题一求解", target_is_directory=True)
        with self.assertRaises(TX.ProjectTransactionError):
            self.prepare(archive_relative="state/alias/history")

    def test_hardlinked_archived_member_is_rejected(self):
        ref = self.prepare()
        stored = self.stored(self.code)
        stored.unlink()
        try:
            os.link(self.root / self.code, stored)
        except OSError as exc:
            self.skipTest(f"hard links unavailable: {exc}")
        with self.assertRaises(TX.ProjectTransactionError):
            TX.verify_history_archive(self.root, ref)

    def test_byte_guarded_commit_persists_reference_without_changing_history(self):
        ref = self.prepare()
        def candidate_validator(staged):
            candidate = yaml.safe_load(staged[TX.STATE_RELATIVE_PATH].read_text(encoding="utf-8"))
            self.assertEqual(candidate["project"]["state_generation"], 1)
            self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])
            TX.verify_history_archive(self.root, ref)
        result = self.commit(ref, validators=[candidate_validator])
        self.assertEqual(result["target_generation"], 1)
        self.assertEqual(yaml.safe_load((self.root / "history_reference.yaml").read_text()), ref)
        for relative, content in self.old_bytes.items():
            self.assertEqual(self.stored(relative).read_bytes(), content)
        self.assert_no_journal()

    def test_archived_sources_must_all_match_confirmed_read_set(self):
        ref = self.prepare()
        for expected in ({TX.STATE_RELATIVE_PATH: self.expected[TX.STATE_RELATIVE_PATH], "history_reference.yaml": None},
                         {**self.expected, "raw.bin": "a" * 64, "history_reference.yaml": None}):
            with self.subTest(expected=expected), self.assertRaises(TX.ReadSetConflictError):
                self.commit(ref, expected_file_hashes=expected)
        self.assert_live_original()
        self.assert_no_journal()

    def test_archives_require_byte_protection_and_cannot_be_transaction_targets(self):
        ref = self.prepare()
        with self.assertRaises(TX.ProjectTransactionError):
            TX.commit_project_state(self.root, self.state, expected_generation=0, preserved_archives=[ref])
        with self.assertRaises(TX.ProjectTransactionError):
            self.commit(ref, writes_after_state=[(ref["manifest"], "replace history")])
        with self.assertRaises(TX.ProjectTransactionError):
            self.commit(ref, preserved_archives=[ref, ref])
        self.assert_live_original()
        self.assert_no_journal()

    def test_other_generation_archive_cannot_bind_new_transaction(self):
        ref = self.prepare()
        self.state_path.write_text("project: {state_generation: 1}\n")
        with self.assertRaises(TX.GenerationConflictError):
            TX.commit_project_state(self.root, self.state, expected_generation=1,
                expected_file_hashes=self.hashes(*self.expected), preserved_archives=[ref])

    def test_candidate_failure_keeps_complete_unreferenced_archive(self):
        ref = self.prepare()
        def fail(staged):
            raise ValueError("candidate rejected")
        with self.assertRaisesRegex(ValueError, "candidate rejected"):
            self.commit(ref, validators=[fail])
        self.assert_live_original()
        self.assert_no_journal()
        TX.verify_history_archive(self.root, ref)

    def test_archive_tamper_during_validation_blocks_before_prepare(self):
        ref = self.prepare()
        def corrupt(staged):
            self.stored(self.workbook).write_bytes(b"damaged")
        with self.assertRaises(TX.ProjectTransactionError):
            self.commit(ref, validators=[corrupt])
        self.assert_live_original()
        self.assert_no_journal()

    def test_unlisted_member_during_validation_is_not_ignored(self):
        ref = self.prepare()
        def inject(staged):
            (self.root / self.directory / "injected").write_bytes(b"unexpected")
        with self.assertRaises(TX.TransactionRecoveryError):
            self.commit(ref, validators=[inject])
        self.assert_live_original()
        self.assert_no_journal()

    def interrupt(self, ref, boundary):
        def stop(point):
            if point == boundary:
                raise RuntimeError(boundary)
        with self.assertRaises(RuntimeError):
            self.commit(ref, failure_hook=stop, writes_before_state=[("模型论文框架.md", "新框架\n")])
        journal = yaml.safe_load((self.root / TX.JOURNAL_RELATIVE_PATH).read_text(encoding="utf-8"))
        self.assertEqual(journal["version"], 2)
        self.assertEqual(journal["preserved_archives"], [ref])
        return journal

    def test_each_replace_boundary_recovers_forward_with_archives(self):
        ref = self.prepare()
        boundaries = ["after_journal_prepared", "after_replace:模型论文框架.md",
                      "after_replace:state/project_state.yaml", "after_replace:history_reference.yaml",
                      "after_journal_committed"]
        for boundary in boundaries:
            with self.subTest(boundary=boundary):
                for relative, content in self.old_bytes.items():
                    (self.root / relative).write_bytes(content)
                (self.root / "history_reference.yaml").unlink(missing_ok=True)
                self.interrupt(ref, boundary)
                result = TX.recover_project_transaction(self.root)
                self.assertIn(result["status"], ("rolled_forward", "committed_cleanup"))
                self.assertEqual(yaml.safe_load(self.state_path.read_text())["project"]["state_generation"], 1)
                self.assertEqual((self.root / "模型论文框架.md").read_text(encoding="utf-8"), "新框架\n")
                TX.verify_history_archive(self.root, ref)
                self.assert_no_journal()

    def test_corrupt_archive_blocks_recovery_before_any_remaining_replacement(self):
        ref = self.prepare()
        self.interrupt(ref, "after_replace:模型论文框架.md")
        self.stored("raw.bin").write_bytes(b"corrupted history")
        before = {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        with self.assertRaises(TX.TransactionRecoveryError):
            TX.recover_project_transaction(self.root)
        self.assertEqual(before, {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file()})
        self.stored("raw.bin").write_bytes(self.old_bytes["raw.bin"])
        self.assertEqual(TX.recover_project_transaction(self.root)["status"], "rolled_forward")
        self.assert_no_journal()

    def test_committed_cleanup_also_refuses_damaged_archive(self):
        ref = self.prepare()
        self.interrupt(ref, "after_journal_committed")
        self.stored(self.code).unlink()
        with self.assertRaises(TX.TransactionRecoveryError):
            TX.recover_project_transaction(self.root)
        self.assertTrue((self.root / TX.JOURNAL_RELATIVE_PATH).is_file())

    def test_journal_cleanup_paths_cannot_delete_preserved_history(self):
        ref = self.prepare()
        original = self.interrupt(ref, "after_journal_committed")
        path = self.root / TX.JOURNAL_RELATIVE_PATH
        for field in ("backup", "staged"):
            with self.subTest(field=field):
                journal = deepcopy(original)
                journal["entries"][0][field] = self.stored(self.workbook).relative_to(self.root).as_posix()
                path.write_text(yaml.safe_dump(journal), encoding="utf-8")
                with self.assertRaises(TX.TransactionRecoveryError):
                    TX.recover_project_transaction(self.root)
                self.assertEqual(self.stored(self.workbook).read_bytes(), self.old_bytes[self.workbook])
                self.assertTrue(path.is_file())
        path.write_text(yaml.safe_dump(original), encoding="utf-8")
        TX.recover_project_transaction(self.root)
        TX.verify_history_archive(self.root, ref)
        self.assert_no_journal()

    def test_legacy_loader_cannot_bypass_v2_archive_verification(self):
        ref = self.prepare()
        self.interrupt(ref, "after_journal_prepared")
        self.stored(self.workbook).write_bytes(b"corrupt")
        with self.assertRaises(TX.TransactionRecoveryError):
            TX.load_state_for_update(self.root)
        self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])

    def test_v2_cannot_be_downgraded_or_drop_archive_list(self):
        ref = self.prepare()
        original = self.interrupt(ref, "after_journal_prepared")
        path = self.root / TX.JOURNAL_RELATIVE_PATH
        variants = [{**original, "version": 1}, {**original, "preserved_archives": []},
                    {k: v for k, v in original.items() if k != "preserved_archives"},
                    {**original, "version": 3}, {**original, "base_generation": True}]
        for variant in variants:
            with self.subTest(variant=variant):
                path.write_text(yaml.safe_dump(variant), encoding="utf-8")
                with self.assertRaises(TX.TransactionRecoveryError):
                    TX.recover_project_transaction(self.root)
                self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])

    def test_journal_old_state_hash_must_match_archived_state(self):
        ref = self.prepare()
        journal = self.interrupt(ref, "after_journal_prepared")
        for entry in journal["entries"]:
            if entry["path"] == TX.STATE_RELATIVE_PATH:
                entry["old_sha256"] = "f" * 64
        (self.root / TX.JOURNAL_RELATIVE_PATH).write_text(yaml.safe_dump(journal), encoding="utf-8")
        with self.assertRaises(TX.TransactionRecoveryError):
            TX.recover_project_transaction(self.root)
        self.assertEqual(self.state_path.read_bytes(), self.old_bytes[TX.STATE_RELATIVE_PATH])

    def test_unknown_third_party_live_content_still_blocks_recovery(self):
        ref = self.prepare()
        self.interrupt(ref, "after_journal_prepared")
        self.state_path.write_bytes(b"unknown third-party contents")
        with self.assertRaises(TX.ProjectTransactionError):
            TX.recover_project_transaction(self.root)
        self.assertEqual(self.state_path.read_bytes(), b"unknown third-party contents")
        TX.verify_history_archive(self.root, ref)

    def test_two_archivers_never_overwrite_same_destination(self):
        entered, release = threading.Event(), threading.Event()
        results = []
        def wait(point):
            if point == "before_archive_create":
                entered.set()
                if not release.wait(10):
                    raise RuntimeError("test synchronization timeout")
        def run(hook=None):
            try:
                results.append(self.prepare(failure_hook=hook))
            except Exception as exc:
                results.append(exc)
        one, two = threading.Thread(target=run, args=(wait,)), threading.Thread(target=run)
        one.start()
        try:
            self.assertTrue(entered.wait(10))
            two.start()
        finally:
            release.set(); one.join(10)
            if two.ident is not None:
                two.join(10)
        self.assertFalse(one.is_alive()); self.assertFalse(two.is_alive())
        self.assertEqual(sum(isinstance(x, dict) for x in results), 1)
        self.assertEqual(sum(isinstance(x, TX.HistoryArchiveError) for x in results), 1)
        self.assert_live_original()

    def test_fresh_process_recovers_using_only_persisted_reference(self):
        ref = self.prepare()
        self.interrupt(ref, "after_replace:state/project_state.yaml")
        command = [sys.executable, "-c", "import sys; from pathlib import Path; "
                   "sys.path.insert(0, sys.argv[1]); import project_transaction as tx; "
                   "print(tx.recover_project_transaction(Path(sys.argv[2]))['status'])",
                   str(ROOT / "scripts"), str(self.root)]
        result = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("rolled_forward", result.stdout)
        self.assert_no_journal()
        TX.verify_history_archive(self.root, ref)


    def test_archive_corruption_after_prepare_blocks_before_first_replace(self):
        ref = self.prepare()
        def corrupt(point):
            if point == "after_journal_prepared":
                self.stored("raw.bin").write_bytes(b"corrupt")
        with self.assertRaises(TX.TransactionRecoveryError):
            self.commit(ref, failure_hook=corrupt)
        self.assert_live_original()
        self.assertTrue((self.root / TX.JOURNAL_RELATIVE_PATH).is_file())

    def test_reference_mutation_cannot_rebind_inflight_commit(self):
        ref = self.prepare()
        def mutate(staged):
            path = self.root / ref["manifest"]
            path.write_bytes(path.read_bytes() + b" ")
            ref["sha256"] = TX.sha256_file(path)
        with self.assertRaises(TX.ReadSetConflictError):
            self.commit(ref, validators=[mutate])
        self.assert_live_original()
        self.assert_no_journal()

    def test_destination_race_preserves_other_creators_files(self):
        def collide(point):
            if point == "before_archive_create":
                path = self.root / self.directory
                path.mkdir(parents=True)
                (path / "other-owner").write_bytes(b"do not delete")
        with self.assertRaises(TX.HistoryArchiveError) as error:
            self.prepare(failure_hook=collide)
        self.assertEqual(error.exception.cleanup_status, "not_created")
        self.assertEqual((self.root / self.directory / "other-owner").read_bytes(), b"do not delete")
        self.assert_live_original()

    def test_reserve_and_reference_parameters_fail_closed(self):
        for reserve in (True, -1, "0"):
            with self.subTest(reserve=reserve), self.assertRaises(TX.ProjectTransactionError):
                self.prepare(reserve_bytes=reserve)
        ref = self.prepare()
        for bad in ({}, {**ref, "extra": True}, {**ref, "sha256": "0"},
                    {**ref, "manifest": "../manifest.json"}, None):
            with self.subTest(reference=bad), self.assertRaises(TX.ProjectTransactionError):
                TX.verify_history_archive(self.root, bad)


if __name__ == "__main__":
    unittest.main()
