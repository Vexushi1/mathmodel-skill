from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import stage_code as STAGE
import validate_code_delivery as CODE
import validate_user_execution as RECEIPT
import analysis_prerequisites as PREREQUISITES


class SolverBackendTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        (self.root / "问题一求解").mkdir()

    def config(self, backend="python", stage="primary", **updates):
        result = {"stage": stage, "problem_name": "问题一", "solver_backend": backend,
                  "data_paths": ["data.csv"], "data_sha256": "a" * 64, "solver": "direct",
                  "random_seed": 2026, "tolerance": 1e-8, "iteration_or_time_limit": "direct",
                  "expected_workbook": f"问题一{'求解结果' if stage == 'primary' else '结果深化分析'}.xlsx",
                  "run_receipt_protocol_version": "1.1.0", "code_dependencies": []}
        if stage == "primary":
            result["primary_quality_protocol_version"] = "1.0.0"
        result.update(updates)
        return result

    def source(self, config, extra=""):
        backend, stage = config.get("solver_backend", "python"), config["stage"]
        if backend == "matlab":
            stem = "q1_solver" if stage == "primary" else "q1_analysis"
            path = self.root / "问题一求解" / (stem + ".m")
            encoded = json.dumps(config, ensure_ascii=False).replace("'", "''")
            text = f"function {stem}()\nRUN_CONFIG=jsondecode('{encoded}');\n{extra}\nend\n"
        else:
            name = "问题一求解.py" if stage == "primary" else "问题一结果深化分析.py"
            path = self.root / "问题一求解" / name
            text = f"RUN_CONFIG={config!r}\n{extra}\ndef main():\n    return 0\nif __name__ == '__main__':\n    main()\n"
        path.write_text(text, encoding="utf-8")
        return path

    def entry(self, source, config, *, accepted=False):
        stage = config["stage"]
        fingerprint = STAGE.stage_code_fingerprint(self.root, source, config.get("code_dependencies", []))
        record = {"backend": config["solver_backend"], "selection_reason": "fixture explicit backend",
                  "bundle_sha256": fingerprint["bundle_sha256"]}
        if accepted:
            record["validated_bundle_sha256"] = fingerprint["bundle_sha256"]
        return {"code" if stage == "primary" else "result_analysis_code": source.relative_to(self.root).as_posix(),
                "primary_code_sha256" if stage == "primary" else "analysis_code_sha256": fingerprint["entry_sha256"],
                "solver_execution": {stage: record}, "data_hash": config["data_sha256"], "capabilities": {}}

    def state(self, entry=None):
        state = {"project": {"current_phase": "solve_validate"}, "preprocessing": {"decision": "not_needed"},
                 "subproblems": {"Q1": entry or {"status": "designed"}}, "data": {}}
        directory = self.root / "state"
        directory.mkdir(exist_ok=True)
        (directory / "project_state.yaml").write_text(yaml.safe_dump(state, allow_unicode=True), encoding="utf-8")
        return state

    def test_python_and_matlab_identity_and_plot_rejection(self):
        for backend in ("python", "matlab"):
            source = self.source(self.config(backend))
            self.assertEqual(CODE.script_identity(source), ("问题一", "primary"))
            self.assertEqual(STAGE.script_identity(source).backend, backend)
        with self.assertRaises(ValueError):
            CODE.script_identity(self.root / "问题一求解/q1_plot.m")

    def test_q10_uses_shared_contract_mapping(self):
        folder = self.root / "问题十求解"
        folder.mkdir()
        source = folder / "q10_solver.m"
        source.write_text("function q10_solver()\nend")
        self.assertEqual(STAGE.resolve_stage_code(self.root, "Q10", "primary", "matlab").path, source)

    def test_unselected_matlab_and_ambiguous_candidates_are_rejected(self):
        self.source(self.config("matlab"))
        with self.assertRaisesRegex(ValueError, "显式"):
            STAGE.resolve_stage_code(self.root, "Q1", "primary")
        self.source(self.config("python"))
        with self.assertRaisesRegex(ValueError, "多个后端"):
            STAGE.resolve_stage_code(self.root, "Q1", "primary")
        self.assertEqual(STAGE.resolve_stage_code(self.root, "Q1", "primary", "matlab").backend, "matlab")

    def test_chosen_missing_backend_never_falls_back(self):
        self.source(self.config("python"))
        with self.assertRaisesRegex(STAGE.StageCodeMissingError, "fallback"):
            STAGE.resolve_stage_code(self.root, "Q1", "primary", "matlab")
        with self.assertRaises(STAGE.StageCodeError) as caught:
            STAGE.resolve_stage_code(self.root, "Q1", "primary", "matlab",
                                     entry={"code": "问题一求解/q1_solver.m"})
        self.assertNotIsInstance(caught.exception, STAGE.StageCodeMissingError)

    def test_declared_path_backend_mismatch_is_rejected(self):
        config = self.config("python")
        source = self.source(config)
        entry = self.entry(source, config)
        entry["solver_execution"]["primary"]["backend"] = "matlab"
        with self.assertRaisesRegex(ValueError, "冲突"):
            STAGE.resolve_stage_code(self.root, "Q1", "primary", entry=entry)

    def test_legacy_root_python_discovery_remains_readable(self):
        path = self.root / "问题一求解.py"
        path.write_text("x=1\n")
        code = STAGE.resolve_stage_code(self.root, "Q1", "primary")
        self.assertTrue(code.legacy)
        self.assertEqual(STAGE.validate_stage_binding(self.root, {"code": path.name}, "primary"), [])

    def test_fingerprint_is_order_independent_and_uses_raw_bytes(self):
        source = self.source(self.config())
        dependencies = []
        for name, content in (("甲.py", b"x=1\r\n"), ("z.py", b"x=2\n")):
            (self.root / name).write_bytes(content)
            dependencies.append({"path": name, "sha256": hashlib.sha256(content).hexdigest()})
        a = STAGE.stage_code_fingerprint(self.root, source, dependencies)
        b = STAGE.stage_code_fingerprint(self.root, source, list(reversed(dependencies)))
        self.assertEqual(a, b)
        expected = hashlib.sha256()
        for row in a["files"]:
            expected.update(row["path"].encode("utf-8") + b"\0" + bytes.fromhex(row["sha256"]))
        self.assertEqual(a["bundle_sha256"], expected.hexdigest())
        self.assertEqual(a["entry_sha256"], hashlib.sha256(source.read_bytes()).hexdigest())

    def test_fingerprint_rejects_bad_paths_missing_duplicates_and_self(self):
        source = self.source(self.config())
        cases = ["../outside.py", "C:/outside.py", "missing.py", source.relative_to(self.root).as_posix(), "./x.py"]
        for path in cases:
            with self.subTest(path=path), self.assertRaises(ValueError):
                STAGE.stage_code_fingerprint(self.root, source, [{"path": path, "sha256": "a" * 64}])

    def test_wrong_helper_hash_fails_but_snapshot_can_observe_current(self):
        source = self.source(self.config())
        helper = self.root / "helper.py"
        helper.write_text("value=2\n")
        dependencies = [{"path": "helper.py", "sha256": "a" * 64}]
        with self.assertRaisesRegex(ValueError, "不一致"):
            STAGE.stage_code_fingerprint(self.root, source, dependencies)
        self.assertEqual(len(STAGE.stage_code_fingerprint(self.root, source, dependencies, check_declared_hashes=False)["files"]), 2)

    def test_declared_helpers_are_checked_without_sync(self):
        helper = self.root / "helper.py"
        helper.write_text("value=2\n")
        config = self.config(code_dependencies=[{"path": "helper.py", "sha256": hashlib.sha256(helper.read_bytes()).hexdigest()}])
        source = self.source(config, "import helper")
        entry = self.entry(source, config, accepted=True)
        self.assertEqual(STAGE.validate_stage_binding(self.root, entry, "primary", require_validated=True), [])
        helper.write_text("value=3\n")
        self.assertTrue(STAGE.validate_stage_binding(self.root, entry, "primary", require_validated=True))
        self.assertTrue(any("依赖SHA" in issue for issue in PREREQUISITES.primary_issues(self.root, {}, entry)))

    def test_local_python_import_requires_explicit_dependency(self):
        (self.root / "helper.py").write_text("value=2")
        config = self.config()
        source = self.source(config, "import helper")
        self.assertTrue(any("未声明" in issue for issue in STAGE.dependency_reference_issues(self.root, source, config)))

    def test_dynamic_python_import_is_not_assumed_closed(self):
        config = self.config()
        source = self.source(config, "module=__import__('helper')")
        self.assertTrue(STAGE.dependency_reference_issues(self.root, source, config))

    def test_modern_python_and_matlab_static_delivery(self):
        for backend in ("python", "matlab"):
            config = self.config(backend)
            source = self.source(config)
            self.assertEqual(CODE.validate_script(self.root, source)[0], [])

    def test_matlab_must_have_new_protocol(self):
        config = self.config("matlab", run_receipt_protocol_version="1.0.0")
        self.assertTrue(any("1.1.0" in issue for issue in CODE.validate_script(self.root, self.source(config))[0]))

    def test_preprocessing_cannot_opt_into_new_bundle_protocol(self):
        folder = self.root / "数据预处理"
        folder.mkdir()
        config = self.config(stage="preprocessing", problem_name="数据预处理", expected_workbook="数据预处理结果.xlsx")
        path = folder / "数据预处理.py"
        path.write_text(f"RUN_CONFIG={config!r}\nif __name__ == '__main__':\n    pass\n", encoding="utf-8")
        self.assertTrue(any("不支持当前阶段" in issue for issue in CODE.validate_script(self.root, path)[0]))

    def test_config_mutation_and_bundle_self_reference_are_rejected(self):
        config = self.config()
        source = self.source(config, "RUN_CONFIG['tolerance']=1.0")
        self.assertTrue(any("覆盖" in issue for issue in CODE.validate_script(self.root, source)[0]))
        config["code_bundle_sha256"] = "a" * 64
        self.assertTrue(any("入口自身" in issue for issue in CODE.validate_script(self.root, self.source(config))[0]))
        source = self.source(self.config(), "def shadow():\n    RUN_CONFIG = {}")
        self.assertTrue(any("局部重定义" in issue for issue in CODE.validate_script(self.root, source)[0]))

    def test_nested_tolerance_is_not_silently_stringified(self):
        config = self.config(tolerance={"AbsTol": 1e-6})
        self.assertTrue(any("标量" in issue for issue in CODE.validate_script(self.root, self.source(config))[0]))

    def test_delivery_persists_delivered_but_not_validated_bundle(self):
        config = self.config()
        source = self.source(config)
        self.state()
        CODE.update_state(self.root, config, source)
        entry = yaml.safe_load((self.root / "state/project_state.yaml").read_text(encoding="utf-8"))["subproblems"]["Q1"]
        self.assertEqual(entry["primary_execution_status"], "awaiting_user_execution")
        self.assertIn("bundle_sha256", entry["solver_execution"]["primary"])
        self.assertNotIn("validated_bundle_sha256", entry["solver_execution"]["primary"])

    def receipt(self, config, bundle="b" * 64):
        result = {field: config[field] for field in RECEIPT.RUN_RECEIPT_ECHO_FIELDS}
        result.update(run_receipt_version="1.1.0", solver_backend=config["solver_backend"], code_bundle_sha256=bundle)
        return result

    def test_receipt_11_echo_and_downgrade_checks(self):
        config = self.config("matlab")
        receipt = self.receipt(config)
        self.assertEqual(RECEIPT.validate_run_receipt_binding(receipt, config), [])
        for key, value in (("solver_backend", "python"), ("run_receipt_version", "1.0.0"), ("code_bundle_sha256", ""), ("tolerance", 1.0)):
            with self.subTest(key=key):
                changed = {**receipt, key: value}
                self.assertTrue(RECEIPT.validate_run_receipt_binding(changed, config))

    def test_preprocessing_receipt_10_and_old_10_read(self):
        config = self.config(run_receipt_protocol_version="1.0.0")
        config.pop("solver_backend")
        receipt = {field: config[field] for field in RECEIPT.RUN_RECEIPT_ECHO_FIELDS}
        receipt["run_receipt_version"] = "1.0.0"
        self.assertEqual(RECEIPT.validate_run_receipt_binding(receipt, config), [])
        receipt["stage"] = config["stage"] = "preprocessing"
        self.assertEqual(RECEIPT.validate_run_receipt_binding(receipt, config), [])
        receipt["run_receipt_version"] = config["run_receipt_protocol_version"] = "1.1.0"
        self.assertTrue(RECEIPT.validate_run_receipt_binding(receipt, config))

    def test_duplicate_receipt_keys_are_rejected(self):
        path = self.root / "duplicate.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "运行配置"
        sheet.append(["项目", "值"])
        sheet.append(["solver_backend", "matlab"])
        sheet.append(["solver_backend", "python"])
        book.save(path)
        self.assertTrue(any("重复" in issue for issue in RECEIPT.configuration_map(path)[1]))

    def primary_workbook(self, source, config, bundle):
        path = self.root / "问题一求解/问题一求解结果.xlsx"
        book = openpyxl.Workbook()
        runtime = book.active
        runtime.title = "运行配置"
        runtime.append(["项目", "值"])
        receipt = self.receipt(config, bundle)
        receipt.update(execution_owner="user", execution_profile="full_fidelity",
                       code_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), solver_version="fixture",
                       actual_stop_reason="direct", repetitions_or_scenarios=1, grid_or_time_range="direct",
                       fallback_used=False, platform="fixture", primary_quality_protocol_version="1.0.0")
        receipt.update({flag: False for flag in RECEIPT.FALSE_FLAGS})
        for key, value in receipt.items():
            runtime.append([key, value])
        quality = book.create_sheet("主结果质量门")
        quality.append(["Verification ID", "检查项", "是否通过", "证据", "判定关系", "阈值或容差", "实际值", "证据工作表", "阈值来源"])
        quality.append(["PQ-Q1-01", "完整计算", True, "fixture", "bool_true", True, True, "基础数值证据", "solver_tolerance"])
        evidence = book.create_sheet("基础数值证据")
        evidence.append(["检查项", "数值"])
        evidence.append(["full", 1])
        book.save(path)
        return path

    def test_receipt_acceptance_sets_only_validated_stage_bundle(self):
        for backend in ("python", "matlab"):
            with self.subTest(backend=backend):
                config = self.config(backend)
                source = self.source(config)
                entry = self.entry(source, config)
                state = self.state(entry)
                bundle = entry["solver_execution"]["primary"]["bundle_sha256"]
                workbook = self.primary_workbook(source, config, bundle)
                self.assertEqual(RECEIPT.validate_one(self.root, workbook, state, True), [])
                self.assertEqual(entry["primary_execution_status"], "accepted")
                self.assertEqual(entry["solver_execution"]["primary"]["validated_bundle_sha256"], bundle)
                self.assertNotIn("analysis", entry["solver_execution"])

    def test_wrong_bundle_receipt_does_not_create_validated_identity(self):
        config = self.config("matlab")
        source = self.source(config)
        entry = self.entry(source, config)
        state = self.state(entry)
        workbook = self.primary_workbook(source, config, "f" * 64)
        issues = RECEIPT.validate_one(self.root, workbook, state, True)
        self.assertTrue(any("bundle" in issue for issue in issues))
        self.assertEqual(entry["primary_execution_status"], "rejected")
        self.assertNotIn("validated_bundle_sha256", entry["solver_execution"]["primary"])

    def test_accepted_primary_freeze_includes_dependency_changes(self):
        helper = self.root / "helper.py"
        helper.write_text("value=1\n")
        config = self.config(code_dependencies=[{"path": "helper.py", "sha256": hashlib.sha256(helper.read_bytes()).hexdigest()}])
        source = self.source(config)
        entry = self.entry(source, config, accepted=True)
        entry["primary_execution_status"] = "accepted"
        state = self.state(entry)
        state["project"]["current_phase"] = "figure_evidence"
        (self.root / "state/project_state.yaml").write_text(yaml.safe_dump(state, allow_unicode=True), encoding="utf-8")
        helper.write_text("value=2\n")
        config["code_dependencies"][0]["sha256"] = hashlib.sha256(helper.read_bytes()).hexdigest()
        self.source(config)
        with self.assertRaisesRegex(ValueError, "冻结"):
            CODE.update_state(self.root, config, source)

    def test_selection_without_entry_sha_cannot_claim_new_binding(self):
        config = self.config()
        source = self.source(config)
        entry = self.entry(source, config)
        entry.pop("primary_code_sha256")
        self.assertTrue(any("入口SHA" in issue for issue in STAGE.validate_stage_binding(self.root, entry, "primary")))

    def test_uppercase_bundle_hashes_keep_identical_binding(self):
        config = self.config()
        source = self.source(config)
        entry = self.entry(source, config, accepted=True)
        selection = entry["solver_execution"]["primary"]
        for field in ("bundle_sha256", "validated_bundle_sha256"):
            selection[field] = selection[field].upper()
        self.assertEqual(STAGE.validate_stage_binding(self.root, entry, "primary", require_validated=True), [])
        state = self.state(entry)
        workbook = self.primary_workbook(source, config, selection["bundle_sha256"].lower())
        self.assertEqual(RECEIPT.validate_one(self.root, workbook, state, True), [])

    def test_new_python_cannot_become_legacy_by_removing_selection(self):
        config = self.config()
        source = self.source(config)
        entry = self.entry(source, config)
        entry.pop("solver_execution")
        self.assertTrue(STAGE.requires_bundle_binding(self.root, entry, "primary"))
        self.assertTrue(STAGE.validate_stage_binding(self.root, entry, "primary"))
        config["run_receipt_protocol_version"] = "1.0.0"
        self.source(config)
        self.assertFalse(STAGE.requires_bundle_binding(self.root, entry, "primary"))
        source.write_text("value=1\n", encoding="utf-8")
        self.assertFalse(STAGE.requires_bundle_binding(self.root, entry, "primary"))
        for selection in ([], [1], {"primary": []}, {"primary": {"backend": "python"}}):
            self.assertTrue(STAGE.requires_bundle_binding(self.root, {"solver_execution": selection}, "primary"))

    def test_invalid_selection_payload_returns_issues_not_attribute_error(self):
        for payload in ([], [1], "matlab", {"primary": []}, {"primary": ["matlab"]}, {"primary": "matlab"}, {"primary": {"backend": []}}):
            with self.subTest(payload=payload):
                self.assertTrue(STAGE.validate_stage_binding(self.root, {"solver_execution": payload}, "primary"))
                with self.assertRaises(ValueError):
                    STAGE.resolve_stage_code(self.root, "Q1", "primary", entry={"solver_execution": payload})

    def test_delivery_rejects_source_changed_after_static_check(self):
        config = self.config()
        source = self.source(config)
        self.state()
        state_path = self.root / "state/project_state.yaml"
        before = state_path.read_bytes()
        checked_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        source.write_text(source.read_text(encoding="utf-8") + "\n# changed after validation\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "检查后改变"):
            CODE.update_state(self.root, config, source, expected_source_sha256=checked_hash)
        self.assertEqual(state_path.read_bytes(), before)

    def test_native_analysis_cannot_rebind_a_changed_source(self):
        config = self.config("matlab")
        source = self.source(config)
        self.state()

        def analyze(path, executable):
            path.write_text(path.read_text(encoding="utf-8") + "\n% changed before native analysis\n", encoding="utf-8")
            return {"status": "passed", "issues": [], "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}

        with patch.object(CODE.MATLAB_CHECKS, "native_code_analysis", side_effect=analyze):
            issues, _ = CODE.validate_script(self.root, source, require_native=True)
        self.assertTrue(any("原生检查源码" in issue for issue in issues))
        self.assertTrue(any("期间改变" in issue for issue in issues))

    def test_delivery_invalid_solver_selection_is_reported(self):
        config = self.config()
        source = self.source(config)
        for payload in ([], {"primary": []}, {"primary": "matlab"}):
            self.state({"solver_execution": payload})
            issues, _ = CODE.validate_script(self.root, source)
            self.assertTrue(any("必须为映射" in issue for issue in issues))
            with self.assertRaisesRegex(ValueError, "必须为映射"):
                CODE.update_state(self.root, config, source)

    def test_project_preprocessing_mode_binds_the_one_accepted_workbook(self):
        pre = self.root / "数据预处理/数据预处理结果.xlsx"
        pre.parent.mkdir()
        book = openpyxl.Workbook()
        book.active.append(["synthetic input"])
        book.save(pre)
        digest = hashlib.sha256(pre.read_bytes()).hexdigest()
        relative = pre.relative_to(self.root).as_posix()
        state = self.state()
        state["preprocessing"] = {"decision": "project_level", "status": "accepted", "quality_status": "passed",
                                  "workbook": relative, "workbook_sha256": digest, "covered_raw_sources": ["raw.csv"]}
        (self.root / "state/project_state.yaml").write_text(yaml.safe_dump(state, allow_unicode=True), encoding="utf-8")
        for backend in ("python", "matlab"):
            for stage in ("primary", "analysis"):
                with self.subTest(backend=backend, stage=stage):
                    check = lambda mode, paths, expected=digest: CODE._decision_gate_issues(
                        self.root, stage, expected, paths, backend=backend, data_identity_mode=mode, modern=True)
                    self.assertEqual(check("preprocessing_workbook", [relative]), [])
                    self.assertTrue(check("combined", [relative]))
                    self.assertTrue(check("preprocessing_workbook", [relative, "raw.csv"]))
                    self.assertTrue(check("preprocessing_workbook", ["other.xlsx"]))
                    self.assertTrue(check("preprocessing_workbook", [relative], "f" * 64))
        pre.write_bytes(pre.read_bytes() + b"modified bytes")
        self.assertTrue(CODE._decision_gate_issues(self.root, "primary", digest, [relative],
                                                 data_identity_mode="preprocessing_workbook", modern=True))

    def test_preprocessing_identity_mode_is_not_a_raw_data_bypass(self):
        self.state()
        config = self.config(data_identity_mode="preprocessing_workbook")
        source = self.source(config)
        self.assertTrue(any("只允许" in issue for issue in CODE.validate_script(self.root, source)[0]))
        for value in (None, "raw", [], {}):
            config["data_identity_mode"] = value
            self.assertTrue(any("data_identity_mode" in issue for issue in CODE.validate_script(self.root, self.source(config))[0]))
        config = self.config(data_identity_mode="combined", run_receipt_protocol_version="1.0.0")
        self.assertTrue(any("data_identity_mode" in issue for issue in CODE.validate_script(self.root, self.source(config))[0]))
        receipt = self.receipt(self.config(), "f" * 64)
        receipt["data_identity_mode"] = "preprocessing_workbook"
        self.assertTrue(any("data_identity_mode" in issue for issue in RECEIPT.validate_run_receipt_binding(receipt, self.config())))

    def helper_source(self, body="y=x;", *, backend="matlab"):
        helper = self.root / "问题一求解/helper.m"
        helper.write_text("function y=helper(x)\n" + body + "\nend\n", encoding="utf-8")
        config = self.config(backend, code_dependencies=[{
            "path": helper.relative_to(self.root).as_posix(), "sha256": hashlib.sha256(helper.read_bytes()).hexdigest()}])
        return self.source(config), helper, config

    def test_declared_matlab_helpers_cannot_hide_syntax_or_dynamic_execution(self):
        self.state()
        for body in ("y=(;", "mlock; y=x;", "eval('y=x');", "system('python other.py');", "y=py.solver(x);"):
            source, helper, config = self.helper_source(body)
            issues, _ = CODE.validate_script(self.root, source)
            self.assertTrue(issues, body)
        source, _, _ = self.helper_source("persistent cache; y=x;")
        self.assertEqual(CODE.validate_script(self.root, source)[0], [])

    def test_every_matlab_helper_is_natively_checked_and_bound(self):
        self.state()
        source, helper, config = self.helper_source()
        reports = {}
        visited = []

        def analyze(path, executable):
            visited.append(path)
            return {"status": "passed", "issues": [], "warnings": [], "release": "test-only-mock",
                    "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}

        with patch.object(CODE.MATLAB_CHECKS, "native_code_analysis", side_effect=analyze):
            self.assertEqual(CODE.validate_script(self.root, source, require_native=True, native_report=reports)[0], [])
        self.assertEqual(visited, [source, helper])
        self.assertEqual(reports["dependency_code_quality"][helper.relative_to(self.root).as_posix()]["native_analysis_status"], "passed")
        before = (self.root / "state/project_state.yaml").read_bytes()
        with self.assertRaisesRegex(ValueError, "bundle在代码检查后改变"):
            CODE.update_state(self.root, config, source, expected_bundle_sha256="f" * 64)
        self.assertEqual((self.root / "state/project_state.yaml").read_bytes(), before)

        def mutate_helper(path, executable):
            result = analyze(path, executable)
            if path == helper:
                helper.write_text(helper.read_text(encoding="utf-8") + "% changed after analysis\n", encoding="utf-8")
            return result

        with patch.object(CODE.MATLAB_CHECKS, "native_code_analysis", side_effect=mutate_helper):
            issues, _ = CODE.validate_script(self.root, source, require_native=True)
        self.assertTrue(any("bundle在工程/原生检查期间改变" in item for item in issues))

    def test_helper_native_failure_or_unavailability_blocks_formal_delivery(self):
        self.state()
        source, helper, _ = self.helper_source()
        for status in ("failed", "unverified"):
            def analyze(path, executable):
                return {"status": status if path == helper else "passed", "issues": ["helper not verified"] if path == helper else [],
                        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
            with patch.object(CODE.MATLAB_CHECKS, "native_code_analysis", side_effect=analyze):
                self.assertTrue(any("helper not verified" in item for item in CODE.validate_script(self.root, source, require_native=True)[0]))

    def test_python_shell_or_matlab_launch_does_not_form_a_single_solver(self):
        for body in ("import subprocess as sp\nsp.run(['matlab'])", "from os import system as launch\nlaunch('matlab')",
                     "import matlab.engine\nmatlab.engine.start_matlab()"):
            config = self.config()
            source = self.source(config, extra=body)
            self.assertTrue(STAGE.dependency_reference_issues(self.root, source, config))

    def test_aliased_dynamic_loading_is_not_a_declared_source_closure(self):
        for body in ("from importlib import import_module as load\nload('extra')",
                     "import importlib as loader\nloader.import_module('extra')",
                     "from runpy import run_path as load\nload('extra.py')",
                     "from importlib.util import spec_from_file_location as load\nload('extra','extra.py')"):
            source = self.source(self.config(), extra=body)
            self.assertTrue(any("动态Python" in issue for issue in STAGE.dependency_reference_issues(self.root, source, self.config())))

    def test_analysis_receipt_requires_the_declared_accepted_primary_identity(self):
        config = self.config("matlab", "analysis", primary_workbook_sha256="a" * 64)
        receipt = self.receipt(config, "f" * 64)
        receipt["primary_workbook_sha256"] = "A" * 64
        self.assertEqual(RECEIPT.validate_run_receipt_binding(receipt, config), [])
        for value in (None, "", "f" * 64):
            receipt["primary_workbook_sha256"] = value
            self.assertTrue(any("primary_workbook_sha256" in issue for issue in RECEIPT.validate_run_receipt_binding(receipt, config)))
        receipt.pop("primary_workbook_sha256")
        receipt["run_receipt_version"] = "1.0.0"
        config["run_receipt_protocol_version"] = "1.0.0"
        self.assertFalse(any("primary_workbook_sha256" in issue for issue in RECEIPT.validate_run_receipt_binding(receipt, config)))


if __name__ == "__main__":
    unittest.main()
