"""Real Q1/Q2 numerical smoke in separate Python and MATLAB projects.

The MATLAB step runs between prepare and verify. Python prepares and accepts
the native MATLAB project; it never substitutes for MATLAB numerical execution.
"""
from __future__ import annotations

import argparse
import ast
from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess
import sys

import openpyxl
import yaml

from test_solver_backend_end_to_end import (
    ROOT, QUESTION, CONFIG, file_hash, load_module,
    reference_digest, save_state, result_io,
)


def accept(root: Path, question: str, backend: str, expected: float) -> None:
    receipt = load_module("mixed_smoke_receipt", "scripts/validate_user_execution.py")
    state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
    workbook = root / f"{question}求解" / f"{question}求解结果.xlsx"
    result_io.validate_workbook_file(workbook, "solution", capabilities={"requires_equilibrium_residual": True})
    config, issues = receipt.configuration_map(workbook)
    assert issues == [], issues
    assert config["solver_backend"] == backend and config["run_receipt_version"] == "1.1.0"
    issues = receipt.validate_one(root, workbook, state, True)
    assert issues == [], issues
    book = openpyxl.load_workbook(workbook, data_only=True)
    try:
        values = dict(book["核心指标"].iter_rows(min_row=2, values_only=True))
        assert values["解"] == expected and values["绝对方程残差"] <= 1e-10
    finally:
        book.close()
    save_state(root, state)


def register(root: Path, question: str, code: Path, backend: str, paths: list[str]) -> None:
    path = root / "state/project_state.yaml"
    state = yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {
        "project": {"competition": "test", "problem": "synthetic", "current_phase": "solve_validate"},
        "preprocessing": {"decision": "not_needed", "status": "not_applicable", "quality_status": "not_applicable"},
        "subproblems": {},
    }
    execution = state.setdefault("execution", {})
    selected = execution.get("solver_backend")
    if selected is not None and selected != backend:
        raise ValueError(f"{question} backend {backend} conflicts with project backend {selected}")
    execution.setdefault("solver_backend", backend)
    execution.setdefault("solver_backend_selection_reason", "Whole-project numerical smoke requirements reviewed")
    key = "Q1" if question == "问题一" else "Q2"
    relative = code.relative_to(root).as_posix()
    state["subproblems"][key] = {
        "status": "designed", "selected_model": "a*x=b", "capabilities": {"requires_equilibrium_residual": True},
        "code": relative, "primary_code_sha256": file_hash(code), "data_hash": reference_digest(root, paths),
        "result_quality_status": "pending", "result_analysis_status": "pending",
        "primary_execution_status": "awaiting_user_execution",
        "solver_execution": {"primary": {"bundle_sha256": reference_digest(root, [relative])}},
    }
    save_state(root, state)


def python_primary(root: Path, question: str, upstream: bool) -> None:
    paths = [f"{QUESTION}/问题一求解结果.xlsx"] if upstream else ["input.json"]
    folder = root / f"{question}求解"
    folder.mkdir(parents=True, exist_ok=True)
    config = dict(stage="primary", problem_name=question, solver_backend="python", data_paths=paths,
                  data_sha256=reference_digest(root, paths), solver="scalar_division", random_seed=2026,
                  tolerance=1e-10, iteration_or_time_limit="direct", expected_workbook=f"{folder.name}/{question}求解结果.xlsx",
                  run_receipt_protocol_version="1.1.0", primary_quality_protocol_version="1.0.0",
                  code_dependencies=[], input_mode="upstream" if upstream else "json")
    source = (ROOT / "tests/fixtures/solver_backends/python_primary.py").read_text(encoding="utf-8")
    source = source.replace("RUN_CONFIG = {}", "RUN_CONFIG = " + repr(config), 1)
    ast.parse(source)
    code = folder / f"{question}求解.py"
    code.write_text(source, encoding="utf-8", newline="\n")
    (root / "requirements.txt").write_text("openpyxl>=3.1\n", encoding="utf-8")
    register(root, question, code, "python", paths)
    process = subprocess.run([sys.executable, str(code)], cwd=root, capture_output=True, text=True, check=True)
    report = json.loads(process.stdout)
    assert report["actual_python_execution"] and report["status"] == "passed"
    (root / f"{question}_python_report.json").write_text(json.dumps(report), encoding="utf-8")
    accept(root, question, "python", 1.5 if upstream else 3)


PYTHON_ANALYSIS = '''"""Independent native Python parameter sensitivity on an accepted primary."""
from pathlib import Path
import hashlib
import json
import math
import platform

import openpyxl

RUN_CONFIG = {}


def digest(root, paths):
    result = hashlib.sha256()
    for relative in sorted(paths, key=lambda name: name.encode("utf-8")):
        result.update(relative.encode("utf-8") + b"\\0")
        result.update(hashlib.sha256((root / relative).read_bytes()).digest())
    return result.hexdigest()


def main():
    root = Path(__file__).resolve().parents[1]
    code = Path(__file__).resolve()
    config = RUN_CONFIG
    primary = root / config["primary_workbook"]
    assert digest(root, config["data_paths"]) == config["data_sha256"]
    assert hashlib.sha256(primary.read_bytes()).hexdigest() == config["primary_workbook_sha256"]
    code_sha = hashlib.sha256(code.read_bytes()).hexdigest()
    bundle_sha = digest(root, [code.relative_to(root).as_posix()])
    upstream = openpyxl.load_workbook(primary, data_only=True)
    try:
        baseline = dict(upstream["核心指标"].iter_rows(min_row=2, values_only=True))["解"]
    finally:
        upstream.close()
    payload = json.loads((root / config["data_paths"][0]).read_text(encoding="utf-8"))
    coefficients = payload["sensitivity_coefficients"]
    assert len(coefficients) == config["iteration_or_time_limit"]
    values = [payload["right_hand_side"] / coefficient for coefficient in coefficients]
    assert all(math.isfinite(value) and value > 0 for value in values) and baseline > 0
    receipt = dict(run_receipt_version="1.1.0", execution_owner="user",
                   execution_profile="full_fidelity", stage="analysis",
                   problem_name=config["problem_name"], solver_backend="python",
                   code_sha256=code_sha, code_bundle_sha256=bundle_sha,
                   data_sha256=config["data_sha256"],
                   primary_workbook_sha256=config["primary_workbook_sha256"],
                   solver=config["solver"], solver_version=platform.python_version(),
                   tolerance=config["tolerance"],
                   iteration_or_time_limit=config["iteration_or_time_limit"],
                   actual_stop_reason="all_scenarios_completed",
                   random_seed=config["random_seed"], repetitions_or_scenarios=len(values),
                   grid_or_time_range="all declared inputs", fallback_used=False,
                   platform=platform.platform())
    receipt.update({name: False for name in (
        "allow_reduced_data", "allow_coarser_grid", "allow_shorter_horizon",
        "allow_fewer_repetitions", "allow_relaxed_tolerance",
        "allow_silent_solver_fallback")})
    book = openpyxl.Workbook()
    book.remove(book.active)
    sheets = {
        "运行配置": [("项目", "值"), *receipt.items()],
        "分析设计": [("风险来源", "分析问题", "方法", "指标", "通过标准"),
                     ("系数变化", "解是否保持正值", "参数敏感性", "x", "全部场景x>0")],
        "参数敏感性": [("参数", "基准值", "变化值", "结果指标"),
                       *[("a", payload["coefficient"], coefficient, value)
                         for coefficient, value in zip(coefficients, values)]],
        "结论稳定性汇总": [("核心结论", "分析方法", "稳定范围", "是否保持"),
                           ("解为正值", "参数敏感性", "a in [1.5, 3]", True)],
    }
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    assert digest(root, config["data_paths"]) == config["data_sha256"]
    assert hashlib.sha256(primary.read_bytes()).hexdigest() == config["primary_workbook_sha256"]
    assert hashlib.sha256(code.read_bytes()).hexdigest() == code_sha
    book.save(root / config["expected_workbook"])
    book.close()
    print(json.dumps({"actual_python_execution": True, "status": "passed",
                      "sensitivity_values": values, "code_sha256": code_sha}))


if __name__ == "__main__":
    main()
'''


def python_analysis(root: Path) -> None:
    primary = root / QUESTION / "问题一求解结果.xlsx"
    state_path = root / "state/project_state.yaml"
    state = yaml.safe_load(state_path.read_text(encoding="utf-8"))
    entry = state["subproblems"]["Q1"]
    assert entry["primary_execution_status"] == "accepted"
    entry.update(result_analysis_requirement_reason="Check the declared coefficient range",
                 analysis_methods=["参数敏感性"])
    save_state(root, state)
    paths = ["input.json"]
    config = dict(stage="analysis", problem_name="问题一", solver_backend="python",
                  data_paths=paths, data_sha256=reference_digest(root, paths),
                  solver="analytic_parameter_sweep", random_seed=2026, tolerance=1e-10,
                  iteration_or_time_limit=3,
                  expected_workbook=f"{QUESTION}/问题一结果深化分析.xlsx",
                  run_receipt_protocol_version="1.1.0", code_dependencies=[],
                  primary_workbook=f"{QUESTION}/问题一求解结果.xlsx",
                  primary_workbook_sha256=file_hash(primary))
    source = PYTHON_ANALYSIS.replace("RUN_CONFIG = {}", "RUN_CONFIG = " + repr(config), 1)
    ast.parse(source)
    code = root / QUESTION / "问题一结果深化分析.py"
    code.write_text(source, encoding="utf-8", newline="\n")
    delivery = load_module("same_backend_analysis_delivery", "scripts/validate_code_delivery.py")
    issues, parsed = delivery.validate_script(root, code, "analysis")
    assert issues == [], issues
    delivery.update_state(root, parsed, code)
    process = subprocess.run([sys.executable, str(code)], cwd=root, capture_output=True, text=True, check=True)
    report = json.loads(process.stdout)
    assert report["actual_python_execution"] and report["status"] == "passed"
    assert report["sensitivity_values"] == [4, 3, 2]
    (root / "analysis_python_report.json").write_text(json.dumps(report), encoding="utf-8")
    workbook = root / QUESTION / "问题一结果深化分析.xlsx"
    result_io.validate_workbook_file(workbook, "result_analysis", capabilities={})
    receipt = load_module("same_backend_analysis_receipt", "scripts/validate_user_execution.py")
    state = yaml.safe_load(state_path.read_text(encoding="utf-8"))
    issues = receipt.validate_one(root, workbook, state, True)
    assert issues == [], issues
    save_state(root, state)


def matlab_downstream(root: Path, *, register_code: bool = True) -> Path:
    source = (ROOT / "templates/code/matlab/q1_solver.m").read_text(encoding="utf-8")
    source = source.replace("q1_solver()", "q2_solver()", 1)
    match = CONFIG.search(source)
    config = json.loads(match.group(1))
    paths = ["input.json", f"{QUESTION}/问题一求解结果.xlsx"]
    config.update(problem_name="问题二", data_paths=paths, data_sha256=reference_digest(root, paths),
                  expected_workbook="问题二求解/问题二求解结果.xlsx")
    literal = json.dumps(config, ensure_ascii=False, separators=(",", ":"))
    source = source[:match.start(1)] + literal + source[match.end(1):]
    needle = "validateattributes(input.coefficient,"
    native_read = '''upstream = readcell(project_path(root, RUN_CONFIG.data_paths{2}), 'Sheet', '核心指标');
row = find(string(upstream(:, 1)) == "解");
assert(isscalar(row), 'HSK:Upstream', 'One upstream accepted solution is required.');
input.right_hand_side = upstream{row, 2};
'''
    source = source.replace(needle, native_read + needle, 1)
    folder = root / "问题二求解"
    folder.mkdir(exist_ok=True)
    code = folder / "q2_solver.m"
    code.write_text(source, encoding="utf-8", newline="\n")
    if register_code:
        register(root, "问题二", code, "matlab", paths)
    return code


def prepare(source_root: Path, project: Path) -> None:
    if project.exists():
        raise ValueError("Same-backend smoke requires a fresh project directory.")
    source_state = yaml.safe_load((source_root / "state/project_state.yaml").read_text(encoding="utf-8"))
    assert source_state["subproblems"]["Q1"]["primary_execution_status"] == "accepted"
    assert source_state["subproblems"]["Q1"]["analysis_execution_status"] == "accepted"
    assert source_state["execution"]["solver_backend"] == "matlab"
    for stage in ("primary", "analysis"):
        native = json.loads((source_root / f"{stage}_matlab_report.json").read_text(encoding="utf-8"))
        assert native["actual_matlab_execution"] and native["status"] == "passed"
    matlab_project = project / "matlab-project"
    matlab_project.mkdir(parents=True)
    shutil.copytree(source_root / QUESTION, matlab_project / QUESTION)
    shutil.copyfile(source_root / "input.json", matlab_project / "input.json")
    save_state(matlab_project, source_state)
    matlab_downstream(matlab_project)
    python_project = project / "python-project"
    python_project.mkdir()
    shutil.copyfile(source_root / "input.json", python_project / "input.json")
    python_primary(python_project, "问题一", False)
    accepted_python_sha = file_hash(python_project / QUESTION / "问题一求解结果.xlsx")
    python_primary(python_project, "问题二", True)
    python_analysis(python_project)
    assert file_hash(python_project / QUESTION / "问题一求解结果.xlsx") == accepted_python_sha
    (project / "synthetic_fixture.json").write_text(json.dumps({
        "fixture": "same_backend_solver_native_v1",
        "source_matlab_sha256": file_hash(source_root / QUESTION / "问题一求解结果.xlsx"),
        "accepted_python_sha256": accepted_python_sha,
        "purpose": "repository_test"}), encoding="utf-8")


def verify(project: Path) -> None:
    marker = json.loads((project / "synthetic_fixture.json").read_text(encoding="utf-8"))
    assert marker["fixture"] == "same_backend_solver_native_v1" and marker["purpose"] == "repository_test"
    report = json.loads((project / "same_backend_matlab_report.json").read_text(encoding="utf-8"))
    assert report["actual_matlab_execution"] and report["status"] == "passed"
    matlab_project = project / "matlab-project"
    python_project = project / "python-project"
    assert file_hash(matlab_project / QUESTION / "问题一求解结果.xlsx") == marker["source_matlab_sha256"]
    assert file_hash(python_project / QUESTION / "问题一求解结果.xlsx") == marker["accepted_python_sha256"]
    assert report["upstream_primary_unchanged"]
    accept(matlab_project, "问题二", "matlab", 1.5)
    receipt = load_module("same_backend_verify_receipt", "scripts/validate_user_execution.py")
    for root, backend in ((python_project, "python"), (matlab_project, "matlab")):
        state = yaml.safe_load((root / "state/project_state.yaml").read_text(encoding="utf-8"))
        assert state["execution"]["solver_backend"] == backend
        assert all(state["subproblems"][key]["primary_execution_status"] == "accepted" for key in ("Q1", "Q2"))
        assert state["subproblems"]["Q1"]["analysis_execution_status"] == "accepted"
        for name in ("问题一求解结果.xlsx", "问题一结果深化分析.xlsx"):
            workbook = root / QUESTION / name
            assert receipt.validate_one(root, workbook, state, False) == []
        workbook = root / "问题二求解" / "问题二求解结果.xlsx"
        assert receipt.validate_one(root, workbook, state, False) == []
        analysis = openpyxl.load_workbook(root / QUESTION / "问题一结果深化分析.xlsx", data_only=True)
        try:
            values = [row[3] for row in analysis["参数敏感性"].iter_rows(min_row=2, values_only=True)]
            assert values == [4, 3, 2], values
        finally:
            analysis.close()
    assert json.loads((python_project / "analysis_python_report.json").read_text(encoding="utf-8"))["actual_python_execution"]
    negative = project / "mixed-rejected"
    negative.mkdir()
    shutil.copytree(python_project / QUESTION, negative / QUESTION)
    shutil.copyfile(python_project / "input.json", negative / "input.json")
    negative_state = deepcopy(yaml.safe_load((python_project / "state/project_state.yaml").read_text(encoding="utf-8")))
    negative_state["subproblems"].pop("Q2")
    save_state(negative, negative_state)
    opposite = matlab_downstream(negative, register_code=False)
    delivery = load_module("same_backend_rejection_delivery", "scripts/validate_code_delivery.py")
    before_state = (negative / "state/project_state.yaml").read_bytes()
    before_primary = file_hash(negative / QUESTION / "问题一求解结果.xlsx")
    issues, config = delivery.validate_script(negative, opposite, "primary")
    assert any("项目后端与源码/RUN_CONFIG/协议不一致" in issue for issue in issues), issues
    assert not any("缺少" in issue or "不存在" in issue for issue in issues), issues
    try:
        delivery.update_state(negative, config, opposite)
    except ValueError as exc:
        assert "项目后端与源码/RUN_CONFIG/协议不一致" in str(exc), str(exc)
    else:
        raise AssertionError("Opposite-backend Q2 delivery was accepted")
    assert (negative / "state/project_state.yaml").read_bytes() == before_state
    assert file_hash(negative / QUESTION / "问题一求解结果.xlsx") == before_primary
    negative_state["subproblems"]["Q2"] = {"solver_execution": {
        "primary": {"backend": "matlab", "selection_reason": "Retired mixed-stage history"}}}
    save_state(negative, negative_state)
    q1_issues, _ = delivery.validate_script(negative, negative / QUESTION / "问题一求解.py", "primary")
    assert any("legacy stage selectors" in issue or "历史" in issue for issue in q1_issues), q1_issues
    (project / "same_backend_acceptance_report.json").write_text(json.dumps({"status": "passed", "cases": [
        "Python project: accepted Q1 -> real Q2 = 1.5 and required Q1 analysis = [4,3,2]",
        "MATLAB project: accepted Q1 -> native Q2 = 1.5 and native Q1 analysis = [4,3,2]",
        "Python project rejects opposite MATLAB Q2 delivery and mixed Q2 history even when checking Q1"],
        "common_receipt_version": "1.1.0", "matlab_release": report["matlab_release"]}), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["prepare", "verify"])
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--source", type=Path)
    args = parser.parse_args()
    if args.action == "prepare":
        assert args.source is not None
        prepare(args.source.resolve(), args.project.resolve())
    else:
        verify(args.project.resolve())
